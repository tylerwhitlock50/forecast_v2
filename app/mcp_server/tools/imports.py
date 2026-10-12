"""
Bulk import tools for building a model from files: customers, products,
machines, labor rates (forecast_import_records), routings, BOMs, prices and
deduction rules, plus the template catalog and a file preview.

Every import calls the /import/* endpoint the Data Import screen uses (same
header aliases, validation, dry run and 1-based row errors). Rows come either
as JSON objects (an assistant that has read a CSV / XLSX passes them) or from
a file the API process can read (``file_path``): the file is parsed here with
the csv module / openpyxl and its rows are sent to the same endpoint.

File paths: a relative path is read from the data folder (``db_manager.
data_dir``; ``/data`` in Docker, which is ``./data`` on the host), and a
leading ``data/`` or ``./data/`` is accepted so host-style paths work. No
path may leave the data folder (symlinks are resolved first): there is no
authentication, so an MCP caller must not be able to list or read the rest
of the API's machine. An absolute path inside the data folder is fine. An
absolute or '~' path elsewhere is refused unless the API runs with
``FORECAST_ALLOW_ABSOLUTE_PATHS=1`` (a local, single-user run).
"""
import csv
import io
import os
import re
from datetime import date, datetime, time
from typing import Annotated, Any, Dict, List, Literal, Optional, Tuple

import anyio
from mcp.server.mcpserver.exceptions import ToolError
from pydantic import Field

import db.database as dbmod
from utils.period_normalizer import month_from_name, normalize_month
from mcp_server.core import DESTRUCTIVE, READ, UPSERT, api, call
from mcp_server.tools._common import cap, out, put_capped, register

FILE_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm")
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000

# Template kind -> the MCP tool that imports it.
KIND_TOOLS = {
    "customers": "forecast_import_records(table='customers')",
    "products": "forecast_import_records(table='units')",
    "machines": "forecast_import_records(table='machines')",
    "labor_rates": "forecast_import_records(table='labor_rates')",
    "routings": "forecast_import_routings",
    "boms": "forecast_import_boms",
    "prices": "forecast_import_prices",
    "deductions": "forecast_import_deductions",
    "sales": "forecast_import_sales",
}

Rows = Annotated[Optional[List[Dict[str, Any]]], Field(
    description="Rows as objects of column: value, one per spreadsheet row (header names are matched "
                "ignoring case / spaces / underscores, plus aliases: see forecast_list_import_templates). "
                "Give rows OR file_path.")]
FilePath = Annotated[Optional[str], Field(
    description="Instead of rows: a .csv/.xlsx in the data folder, relative to it (Docker: ./data on the host), "
                "e.g. 'imports/boms.xlsx'. Paths outside the data folder are refused unless the API runs with "
                "FORECAST_ALLOW_ABSOLUTE_PATHS=1.")]
Sheet = Annotated[Optional[str], Field(description="XLSX sheet (default: the active one)")]
HeaderRow = Annotated[Optional[int], Field(
    description="1-based header row (default: detected - the first row whose cells are known column names or "
                "aliases, so a report title above the header is skipped; else the first row with 2+ filled cells)",
    ge=1)]
DryRun = Annotated[bool, Field(description="Validate and report without writing anything. Do this first.")]
Detail = Annotated[bool, Field(description="Full per-row lists instead of counts and a sample")]


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------

def _data_dir() -> str:
    return os.path.realpath(dbmod.db_manager.data_dir)


def _inside(path: str, folder: str) -> bool:
    return path == folder or path.startswith(folder.rstrip(os.sep) + os.sep)


def _docker_hint() -> str:
    data_dir = _data_dir()
    if data_dir == "/data":
        return (" The API runs in Docker and only sees the data folder: copy the file into ./data on the host "
                "(e.g. ./data/imports/boms.xlsx) and pass 'imports/boms.xlsx'.")
    return f" Relative paths are read from the data folder {data_dir}."


#: Environment switch for reading files outside the data folder (absolute and '~' paths). Off by default.
ALLOW_ABSOLUTE_PATHS_ENV = "FORECAST_ALLOW_ABSOLUTE_PATHS"


def absolute_paths_allowed() -> bool:
    return os.getenv(ALLOW_ABSOLUTE_PATHS_ENV, "").strip().lower() in ("1", "true", "yes", "on")


def resolve_path(file_path: Optional[str], allow_dir: bool = False) -> str:
    """Absolute path of an importable file (or, with allow_dir, a folder); ToolError otherwise."""
    raw = (file_path or "").strip()
    data_dir = _data_dir()
    if not raw:
        if allow_dir:
            return data_dir
        raise ToolError("file_path is empty")
    outside_hint = (f" Copy the file into the data folder ({data_dir}) and pass a path relative to it"
                    f"{' (in Docker: ./data on the host)' if data_dir == '/data' else ''}; for a local single-user "
                    f"run the API can be started with {ALLOW_ABSOLUTE_PATHS_ENV}=1 to allow other folders.")
    if raw.startswith("~") or os.path.isabs(raw):
        if raw.startswith("~") and not absolute_paths_allowed():
            raise ToolError(f"'{raw}': home-folder paths are not read.{outside_hint}")
        path = os.path.realpath(os.path.expanduser(raw))
        if not _inside(path, data_dir) and not absolute_paths_allowed():
            raise ToolError(f"'{raw}' is outside the data folder.{outside_hint}")
        if not os.path.exists(path):
            raise ToolError(f"File not found: {raw}.{_docker_hint()}")
    else:
        norm = raw.replace("\\", "/")
        candidates = [norm]
        for prefix in ("./data/", "data/"):
            if norm.startswith(prefix):
                candidates.append(norm[len(prefix):])
        path = ""
        for candidate in candidates:
            resolved = os.path.realpath(os.path.join(data_dir, candidate))
            if not _inside(resolved, data_dir):
                raise ToolError(f"'{raw}' points outside the data folder ({data_dir}). Relative paths must stay "
                                f"inside it.{outside_hint}")
            if os.path.exists(resolved):
                path = resolved
                break
        if not path:
            raise ToolError(f"File not found in the data folder: {raw}.{_docker_hint()} "
                            f"forecast_preview_file with no file_path lists the folder.")
    if os.path.isdir(path):
        if allow_dir:
            return path
        raise ToolError(f"{raw} is a folder; give a file (forecast_preview_file lists a folder's files)")
    if not os.access(path, os.R_OK):
        raise ToolError(f"{raw} is not readable by the API process")
    if not path.lower().endswith(FILE_EXTENSIONS):
        raise ToolError(f"{raw}: only {', '.join(FILE_EXTENSIONS)} files can be read (save .xls as .xlsx or CSV)")
    if os.path.getsize(path) > MAX_FILE_BYTES:
        raise ToolError(f"{raw} is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    return path


def _cell(value: Any, header: bool = False) -> Any:
    if isinstance(value, datetime):
        if header:
            return value.strftime("%Y-%m")
        return value.strftime("%Y-%m-%d") if value.time() == time(0, 0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.strftime("%Y-%m") if header else value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _read_csv(path: str) -> List[List[Any]]:
    with open(path, "rb") as fh:
        data = fh.read()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")  # Excel "CSV" on Windows
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.lower().endswith(".tsv") else csv.excel
    grid: List[List[Any]] = []
    for row in csv.reader(io.StringIO(text), dialect):
        grid.append(row)
        if len(grid) > MAX_ROWS + 50:
            raise ToolError(f"File has more than {MAX_ROWS:,} rows; split it")
    return grid


def _read_xlsx(path: str, sheet: Optional[str]) -> Tuple[List[List[Any]], str, List[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # zipfile.BadZipFile, InvalidFileException, ...
        raise ToolError(f"Cannot open {os.path.basename(path)} as an Excel workbook: {exc}")
    try:
        names = list(wb.sheetnames)
        if sheet:
            match = next((n for n in names if n == sheet), None) or \
                next((n for n in names if n.strip().lower() == sheet.strip().lower()), None)
            if match is None:
                raise ToolError(f"Sheet '{sheet}' not found; sheets: {', '.join(names)}")
            ws = wb[match]
        else:
            ws = wb.active if wb.active is not None else wb.worksheets[0]
        grid: List[List[Any]] = []
        for row in ws.iter_rows(values_only=True):
            grid.append(list(row))
            if len(grid) > MAX_ROWS + 50:
                raise ToolError(f"Sheet has more than {MAX_ROWS:,} rows; split it")
        return grid, ws.title, names
    finally:
        wb.close()


HEADER_SCAN_ROWS = 30


def _detect_header_row(grid: List[List[Any]]) -> Optional[int]:
    """
    0-based header row: the import API's alias-aware guess_header_row, so a
    report title above the header ('Costed Bill of Material', 'Run date:
    10/12/2026') is skipped even when it fills two cells. Header cells are
    shaped as read_table_file names them first (an Excel date heading is
    'YYYY-MM', which counts as a month column).
    """
    from api.import_routes import guess_header_row  # import_routes holds the header aliases of every import

    sample = [[_cell(v, header=True) for v in (row or [])] for row in grid[:HEADER_SCAN_ROWS]]
    return guess_header_row(sample, scan=HEADER_SCAN_ROWS)


def _header_names(header_cells: List[Any], body: List[List[Any]]) -> Tuple[List[Optional[str]], List[str]]:
    """
    Column name per position (None = dropped) and the generated names.

    A column whose header cell is blank but holds data below is kept as
    'column_<n>' (n = 1-based spreadsheet column), so the import reports it
    as an ignored column instead of its values vanishing; a blank column
    without data is dropped. Repeated names get ' (2)', ' (3)'.
    """
    width = max([len(header_cells)] + [len(r) for r in body if r])
    names = [_cell(header_cells[i], header=True) if i < len(header_cells) else None for i in range(width)]
    pending = {i for i, name in enumerate(names) if _blank(name)}
    has_data = set()
    for raw in body:
        if not pending:
            break
        if not raw:
            continue
        for i in [i for i in pending if i < len(raw) and not _blank(_cell(raw[i]))]:
            has_data.add(i)
            pending.discard(i)

    headers: List[Optional[str]] = []
    unnamed: List[str] = []
    seen: Dict[str, int] = {}
    for i, name in enumerate(names):
        generated = _blank(name)
        if generated:
            if i not in has_data:
                headers.append(None)
                continue
            name = f"column_{i + 1}"
        name = str(name).strip()
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        if generated:
            unnamed.append(name)
        headers.append(name)
    return headers, unnamed


def read_table_file(file_path: str, sheet: Optional[str] = None,
                    header_row: Optional[int] = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Parse a CSV / XLSX file into row objects keyed by the header row.

    Returns (rows, info) with info = {file, sheet, sheets, header_row, columns,
    row_count, unnamed_columns?}. The header row is ``header_row`` when given,
    else detected (_detect_header_row: report titles above it are skipped).
    Columns with a blank header that hold data are kept as 'column_<n>'
    (listed in unnamed_columns). Excel dates become 'YYYY-MM-DD' (header
    cells 'YYYY-MM', so a one-column-per-month sheet keeps its months); blank
    rows are skipped.
    """
    path = resolve_path(file_path)
    sheets: List[str] = []
    sheet_name = None
    if path.lower().endswith((".xlsx", ".xlsm")):
        grid, sheet_name, sheets = _read_xlsx(path, sheet)
    else:
        if sheet:
            raise ToolError("sheet applies to .xlsx files only")
        grid = _read_csv(path)

    if header_row is not None:
        if header_row > len(grid):
            raise ToolError(f"header_row {header_row} is past the end of the file ({len(grid)} rows)")
        header_index = header_row - 1
    else:
        header_index = _detect_header_row(grid)
        if header_index is None:
            raise ToolError(f"{file_path} is empty")

    headers, unnamed = _header_names(grid[header_index] or [], grid[header_index + 1:])
    if not any(h for h in headers if h and h not in unnamed):
        raise ToolError(f"Row {header_index + 1} of {file_path} has no column names; set header_row")

    rows: List[Dict[str, Any]] = []
    for raw in grid[header_index + 1:]:
        record = {}
        for i, name in enumerate(headers):
            if name is None or i >= len(raw):
                continue
            record[name] = _cell(raw[i])
        if all(_blank(v) for v in record.values()):
            continue
        rows.append(record)
    if len(rows) > MAX_ROWS:
        raise ToolError(f"File has more than {MAX_ROWS:,} data rows; split it")
    info = {
        "file": path,
        "sheet": sheet_name,
        "sheets": sheets or None,
        "header_row": header_index + 1,
        "columns": [h for h in headers if h],
        "row_count": len(rows),
    }
    if unnamed:
        info["unnamed_columns"] = unnamed
    return rows, info


def load_rows(rows: Optional[List[Dict[str, Any]]], file_path: Optional[str], sheet: Optional[str] = None,
              header_row: Optional[int] = None) -> Tuple[List[Dict[str, Any]], Optional[Dict[str, Any]]]:
    """Rows from the argument or from a file (exactly one of the two)."""
    if rows and file_path:
        raise ToolError("Give rows or file_path, not both")
    if file_path:
        parsed, info = read_table_file(file_path, sheet, header_row)
        if not parsed:
            raise ToolError(f"{file_path} has a header row but no data rows")
        return parsed, {k: v for k, v in info.items()
                        if k in ("file", "sheet", "header_row", "row_count", "unnamed_columns")}
    if sheet or header_row:
        raise ToolError("sheet / header_row apply to file_path only")
    if not rows:
        raise ToolError("Give rows (a non-empty list of objects) or file_path")
    if not all(isinstance(r, dict) for r in rows):
        raise ToolError("every entry of rows must be an object of column: value")
    return rows, None


async def load_rows_async(rows: Optional[List[Dict[str, Any]]], file_path: Optional[str], sheet: Optional[str],
                header_row: Optional[int]) -> Tuple[List[Dict[str, Any]], Optional[Dict[str, Any]]]:
    """load_rows off the event loop (parsing a large workbook must not stall the API). For any import tool."""
    return await anyio.to_thread.run_sync(load_rows, rows, file_path, sheet, header_row)


# ---------------------------------------------------------------------------
# Result shaping
# ---------------------------------------------------------------------------

def _summarize(body: Dict[str, Any], source: Optional[Dict[str, Any]], detail: bool) -> Dict[str, Any]:
    """The import endpoint's counts, errors and warnings, capped for the model's context."""
    result: Dict[str, Any] = {"message": body.get("message")}
    for key in ("kind", "table", "dry_run", "mode", "rows", "created", "updated", "unchanged", "replaced",
                "boms_replaced", "lines_deleted", "routings_replaced", "operations_deleted", "skipped"):
        if body.get(key) not in (None, ""):
            result[key] = body[key]
    if source:
        result["source"] = source
    errors = body.get("errors") or []
    result["error_count"] = len(errors)
    put_capped(result, "errors", errors, 500 if detail else 25)
    put_capped(result, "warnings", body.get("warnings") or [], 200 if detail else 20)
    if body.get("summary"):
        result["summary"] = body["summary"]
    written = (body.get("created") or 0) + (body.get("updated") or 0) + (body.get("replaced") or 0)
    if errors:
        result["status"] = "partial" if written else "failed"
    else:
        result["status"] = "ok"
    if body.get("dry_run"):
        result["next"] = ("Dry run only: nothing was written. Fix any errors (row numbers are 1-based data rows), "
                          "then call again with dry_run=false.")
    return result


def _records_summary(records: List[Dict[str, Any]], key: str, detail: bool) -> Dict[str, Any]:
    """
    Ids by action (created_ids / updated_ids / skipped_ids / unchanged_ids),
    capped; updates that name their changed fields go in updated_records
    (never in 'updated', which stays the endpoint's count). The full list
    with detail.
    """
    if detail:
        return {"records": records}
    summary: Dict[str, Any] = {}
    for action in ("created", "updated", "skipped", "unchanged"):
        items = [r for r in records if r.get("action") == action]
        if action == "updated" and any(r.get("fields") for r in items):
            put_capped(summary, "updated_records", [{key: r.get(key), "fields": r.get("fields")} for r in items], 25)
        else:
            put_capped(summary, f"{action}_ids", [r.get(key) for r in items], 50)
    return summary


# ---------------------------------------------------------------------------
# Tools
# ---------------------------------------------------------------------------

@register("forecast_list_import_templates", READ)
async def list_import_templates(
    kind: Annotated[Optional[str], Field(
        description="One template: customers, products, machines, labor_rates, routings, boms, prices, deductions "
                    "or sales. Omit for all (setup order).")] = None,
) -> str:
    """Import templates in setup order (customers -> products -> machines ->
    labor rates -> routings -> BOMs -> prices -> deductions -> sales): each
    template's columns (required flag, meaning, accepted header aliases) and
    the MCP tool that imports it. With kind, also the example CSV."""
    body = await call("GET", "/import/templates")
    templates = body.get("templates") or []
    if kind:
        wanted = kind.strip().lower().replace(" ", "_")
        wanted = {"units": "products", "product": "products", "labor": "labor_rates", "bom": "boms",
                  "routing": "routings", "price": "prices", "deduction": "deductions"}.get(wanted, wanted)
        templates = [t for t in templates if t.get("kind") == wanted]
        if not templates:
            raise ToolError(f"Unknown template '{kind}'; one of {', '.join(KIND_TOOLS)}")
    items = []
    for t in templates:
        columns = []
        for c in t.get("columns") or []:
            col = {"name": c.get("name"), "required": bool(c.get("required")), "description": c.get("description")}
            if c.get("aliases"):
                col["aliases"] = c["aliases"]
            columns.append(col)
        item = {"step": t.get("step"), "kind": t.get("kind"), "tool": KIND_TOOLS.get(t.get("kind"), t.get("endpoint")),
                "description": t.get("description"), "columns": columns}
        if kind:
            csv_body = await api("GET", f"/import/templates/{t.get('kind')}")
            item["example_csv"] = csv_body.get("text") if isinstance(csv_body, dict) else None
        items.append(item)
    return out({"templates": items, "notes": body.get("notes")})


@register("forecast_preview_file", READ)
async def preview_file(
    file_path: Annotated[Optional[str], Field(
        description="A .csv / .xlsx file (relative paths are read from the data folder; ./data on the host in "
                    "Docker), or a folder to list. Omit to list the data folder.")] = None,
    sheet: Sheet = None,
    header_row: HeaderRow = None,
    sample_rows: Annotated[int, Field(description="Data rows to show", ge=0, le=50)] = 5,
) -> str:
    """Look at a file before importing it: sheet names, the detected header
    row and columns, row count, the first rows, and which import template
    its columns match best (with required columns it lacks). Given a folder
    (or nothing), lists the importable files in it."""
    path = resolve_path(file_path, allow_dir=True)
    if os.path.isdir(path):
        entries = []
        for name in sorted(os.listdir(path))[:500]:
            full = os.path.join(path, name)
            if os.path.isdir(full):
                entries.append({"name": name + "/", "type": "folder"})
            elif name.lower().endswith(FILE_EXTENSIONS):
                stat = os.stat(full)
                entries.append({"name": name, "bytes": stat.st_size,
                                "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="minutes")})
        shown, omitted = cap(entries, 200)
        result = {"folder": path, "data_folder": _data_dir(), "entries": shown}
        if omitted:
            result["entries_omitted"] = omitted
        return out(result)

    rows, info = await anyio.to_thread.run_sync(read_table_file, file_path or "", sheet, header_row)
    result = {**{k: v for k, v in info.items() if v is not None}, "sample_rows": rows[:sample_rows]}
    templates = (await call("GET", "/import/templates")).get("templates") or []
    result["likely_imports"] = _match_templates(info["columns"], templates)
    return out(result)


def _key(name: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


#: A template counts as a match only when the file has one of these columns: what makes it that kind of
#: file (a customer + product file is not a price list without a price or discount column).
_DEFINING_COLUMNS = {
    "prices": ("unit_price", "discount_pct"),
    "deductions": ("deduction_type", "rate", "amount_per_unit"),
}


def _month_columns(columns: List[str]) -> List[str]:
    """Headers that are months (2027-01, Jan-27, an Excel date, or a bare month name): a wide layout."""
    return [c for c in columns if normalize_month(c) or month_from_name(c)]


def _match_templates(columns: List[str], templates: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Rank templates by how many of the file's headers they recognise (a hint;
    the import decides). Two or more month headers (one column per month)
    stand for the sales template's period + quantity, as forecast_import_sales
    reads that wide layout.
    """
    months = _month_columns(columns)
    wide = len(months) >= 2
    headers = {_key(c): c for c in columns if c not in months}
    ranked = []
    for t in templates:
        kind = t.get("kind")
        matched: Dict[str, str] = {}
        missing = []
        for c in t.get("columns") or []:
            names = [c.get("name"), *(c.get("aliases") or [])]
            hit = next((headers[_key(n)] for n in names if _key(n) in headers), None)
            if hit:
                matched[hit] = c.get("name")
            elif kind == "sales" and wide and c.get("name") in ("period", "quantity"):
                continue
            elif c.get("required"):
                missing.append(c.get("name"))
        defining = _DEFINING_COLUMNS.get(kind)
        if defining and not set(defining) & set(matched.values()):
            continue
        if not matched:
            continue
        score = len(matched) - 2 * len(missing)
        entry: Dict[str, Any] = {"kind": kind, "tool": KIND_TOOLS.get(kind), "columns": matched,
                                 "missing_required": missing}
        if kind == "sales" and wide:
            entry["month_columns"] = f"{len(months)} ({months[0]} .. {months[-1]}): one column per month, read as " \
                                     f"period + quantity"
            score += 2 + len(months)
        entry["_score"] = score
        ranked.append(entry)
    ranked.sort(key=lambda r: -r["_score"])
    for r in ranked:
        r.pop("_score")
    return ranked[:2]


@register("forecast_import_records", UPSERT)
async def import_records(
    table: Annotated[Literal["customers", "units", "products", "machines", "labor_rates"], Field(
        description="customers (customer_id, customer_name, customer_type!, region); units = products (unit_id, "
                    "unit_name, unit_type, base_price, bom_id, bom_version, router_id); machines (machine_id, "
                    "machine_name, machine_rate $/h, available_hours_per_month or available_minutes_per_month); "
                    "labor_rates (rate_id, rate_name, rate_amount $/h)")],
    rows: Rows = None,
    file_path: FilePath = None,
    sheet: Sheet = None,
    header_row: HeaderRow = None,
    mode: Annotated[Literal["upsert", "skip"], Field(
        description="upsert: create new ids and update existing ones, changing only the columns the row fills "
                    "(a blank cell never erases a stored value); skip: leave existing rows alone")] = "upsert",
    dry_run: DryRun = False,
    detail: Detail = False,
) -> str:
    """Bulk create / update master data from a file or rows: customers
    (customer_type drives customer-type prices and deduction rules),
    products (base_price is the fallback price; bom_id / router_id link
    costing), machines and labor rates. The tool for loading the user's
    files: it reads file_path, accepts the header aliases of
    forecast_list_import_templates and reports per-row errors without
    stopping the rest. Rows match by id, or by name when the id is blank;
    new rows without an id get one in the existing style (CUST-001,
    PROD-001, M0001, RATE-001). Load in this order: customers, products,
    machines, labor rates, then routings and BOMs. Returns counts (created,
    updated, skipped), the ids by action and, in updated_records, which
    fields each update changed."""
    data, source = await load_rows_async(rows, file_path, sheet, header_row)
    body = await call("POST", "/import/records", json={"table": table, "rows": data, "mode": mode, "dry_run": dry_run})
    result = _summarize(body, source, detail)
    key = {"customers": "customer_id", "machines": "machine_id", "labor_rates": "rate_id"}.get(body.get("table"), "unit_id")
    result.update(_records_summary(body.get("records") or [], key, detail))
    return out(result)


@register("forecast_import_routings", DESTRUCTIVE)
async def import_routings(
    rows: Annotated[Optional[List[Dict[str, Any]]], Field(
        description="One row per operation: router_id, router_name?, sequence? (10, 20, 30 assigned when blank), "
                    "machine_id (id or name), machine_minutes, labor_minutes (or machine_hours / labor_hours), "
                    "labor_type_id (a labor rate id or name), operation_description?, unit_id? (links the product). "
                    "Give rows OR file_path.")] = None,
    file_path: FilePath = None,
    sheet: Sheet = None,
    header_row: HeaderRow = None,
    mode: Annotated[Literal["replace", "append"], Field(
        description="replace (default): each routing in the input gets exactly these operations (its other "
                    "operations are deleted); append: upsert by router_id + sequence")] = "replace",
    auto_create: Annotated[bool, Field(
        description="Create unknown machines / labor types as $0/h placeholders (reported as warnings) and put "
                    "steps with no machine on a 'MANUAL' bench; default false = row error")] = False,
    link_products: Annotated[bool, Field(description="Set router_id on products named in a unit_id column")] = True,
    dry_run: DryRun = False,
    detail: Detail = False,
) -> str:
    """Import routings (labor and machine minutes per unit by operation).
    Import machines and labor rates first. Reports per routing the labor /
    machine minutes and cost per unit and any cost warnings. A routing with
    any bad row is skipped as a whole."""
    data, source = await load_rows_async(rows, file_path, sheet, header_row)
    body = await call("POST", "/import/routings", json={
        "rows": data, "mode": mode, "auto_create": auto_create, "link_products": link_products, "dry_run": dry_run,
    })
    result = _summarize(body, source, detail)
    routings = []
    for r in body.get("routings") or []:
        item = {k: r.get(k) for k in ("router_id", "action", "operations", "labor_minutes_per_unit",
                                      "machine_minutes_per_unit", "labor_cost_per_unit", "machine_cost_per_unit")}
        if r.get("cost_warnings"):
            item["cost_warnings"] = r["cost_warnings"]
        item["used_by_products"] = r.get("used_by_products") or []
        routings.append(item)
    put_capped(result, "routings", routings, 1000 if detail else 50)
    for key in ("created_machines", "created_labor_types"):
        if body.get(key):
            result[key] = body[key]
    return out(result)


@register("forecast_import_boms", DESTRUCTIVE)
async def import_boms(
    rows: Annotated[Optional[List[Dict[str, Any]]], Field(
        description="One row per component: bom_id (parent / assembly), bom_name?, version? ('1' -> '1.0', default "
                    "1.0), bom_line? (numbered in file order when blank), part_number?, material_description, qty "
                    "(per finished unit), unit?, unit_price (cost per component unit) or material_cost (extended; "
                    "the other is derived), unit_id? (links that product to this BOM + version). "
                    "Give rows OR file_path.")] = None,
    file_path: FilePath = None,
    sheet: Sheet = None,
    header_row: HeaderRow = None,
    mode: Annotated[Literal["replace", "append"], Field(
        description="replace (default): each BOM + version in the input gets exactly these lines (its other lines "
                    "are deleted); append: upsert by bom_id + version + bom_line")] = "replace",
    link_products: Annotated[bool, Field(
        description="Point products named in a unit_id column at the BOM + version")] = True,
    dry_run: DryRun = False,
    detail: Detail = False,
) -> str:
    """Import bills of materials. material_cost per line = qty x unit_price
    (computed by the server). Reports per BOM the line count, material cost
    per finished unit and the products linked / using it. Import products
    first so unit_id links resolve. A BOM with any bad row is skipped as a
    whole (never half-replaced). BOMs are company-wide master data."""
    data, source = await load_rows_async(rows, file_path, sheet, header_row)
    body = await call("POST", "/import/boms", json={
        "rows": data, "mode": mode, "link_products": link_products, "dry_run": dry_run,
    })
    result = _summarize(body, source, detail)
    boms = [{k: b.get(k) for k in ("bom_id", "version", "action", "lines", "material_cost_per_unit", "used_by_products")}
            for b in body.get("boms") or []]
    if result.get("summary"):
        result["summary"].pop("material_cost_by_bom", None)  # repeated in boms
    put_capped(result, "boms", boms, 1000 if detail else 50)
    return out(result)


RuleMode = Annotated[Literal["upsert", "append", "replace"], Field(
    description="upsert (default): update the rule with the same id or the same scope (scenario, customer, "
                "customer type, product(, product type), start, end), else create; append: always create; "
                "replace: first DELETE every rule of the scenarios in the input (company-wide when blank)")]


@register("forecast_import_prices", DESTRUCTIVE)
async def import_prices(
    rows: Annotated[Optional[List[Dict[str, Any]]], Field(
        description="One price rule per row: unit_id (blank = every product, discounts only) or unit_type (product "
                    "type), customer_type and/or "
                    "customer_id (blank = every customer), unit_price (fixed price) OR discount_pct (off the "
                    "product's base_price: 40%, 0.40 or 40), start_period / end_period (YYYY-MM, blank = open), "
                    "forecast_id? (blank = company-wide), notes?. Customers / products may be given by name. "
                    "Give rows OR file_path.")] = None,
    file_path: FilePath = None,
    sheet: Sheet = None,
    header_row: HeaderRow = None,
    forecast_id: Annotated[Optional[str], Field(
        description="Scenario for rows without one; omit for company-wide prices (every scenario)")] = None,
    mode: RuleMode = "upsert",
    dry_run: DryRun = False,
    detail: Detail = False,
) -> str:
    """Import the price list (what each customer type / customer pays per
    product or product type). Most specific rule wins: customer + product >
    customer + product type > customer (all products) > type + product >
    type + product type > type (all products) > product (all customers) >
    product type > the product's base_price. Sales rows without a price are
    priced from this list on import; forecast_apply_price_list re-prices
    existing rows."""
    data, source = await load_rows_async(rows, file_path, sheet, header_row)
    payload: Dict[str, Any] = {"rows": data, "mode": mode, "dry_run": dry_run}
    if forecast_id:
        payload["forecast_id"] = forecast_id
    body = await call("POST", "/import/prices", json=payload)
    result = _summarize(body, source, detail)
    result.update(_records_summary(body.get("records") or [], "price_id", detail))
    return out(result)


@register("forecast_import_deductions", DESTRUCTIVE)
async def import_deductions(
    rows: Annotated[Optional[List[Dict[str, Any]]], Field(
        description="One rule per row: deduction_type (terms_allowance, chargebacks, other_allowance, excise_tax, "
                    "commission; also 'terms', 'FET', 'excise', 'commission'), name?, method (percent_of_gross "
                    "default; embedded_tax default for excise; percent_of_net; per_unit), rate ('2%', 0.02 or 2 "
                    "all mean 2%; 0 = exemption), amount_per_unit (per_unit), scope: customer_id / customer_type / "
                    "unit_id / unit_type (blank = any), start_period / end_period (YYYY-MM), forecast_id?. "
                    "Give rows OR file_path.")] = None,
    file_path: FilePath = None,
    sheet: Sheet = None,
    header_row: HeaderRow = None,
    forecast_id: Annotated[Optional[str], Field(
        description="Scenario for rows without one; omit for company-wide rules (every scenario)")] = None,
    mode: RuleMode = "upsert",
    dry_run: DryRun = False,
    detail: Detail = False,
) -> str:
    """Import sales deduction rules: terms allowance, chargebacks, other
    allowances and excise tax come off gross sales (net sales = gross -
    deductions); commissions are a selling expense. FET embedded in the
    price at 11%: deduction_type excise_tax, method embedded_tax, rate 11%
    (tax = price x 0.11 / 1.11). One terms / chargebacks / excise /
    commission rule applies to each sale (the most specific; on equal
    specificity the scenario rule); other_allowance rows are programs keyed
    by name (Co-op, Freight allowance): every matching program applies and
    they add up, so give each allowance its own name. Types stack. Check
    with forecast_deduction_preview."""
    data, source = await load_rows_async(rows, file_path, sheet, header_row)
    payload: Dict[str, Any] = {"rows": data, "mode": mode, "dry_run": dry_run}
    if forecast_id:
        payload["forecast_id"] = forecast_id
    body = await call("POST", "/import/deductions", json=payload)
    result = _summarize(body, source, detail)
    result.update(_records_summary(body.get("records") or [], "rule_id", detail))
    return out(result)
