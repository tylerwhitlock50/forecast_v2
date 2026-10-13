"""
Actual results: one pipeline for every way actuals get in, an audit trail of
batches, and month-end close.

Every write path (the CSV / XLSX upload, JSON entries from the Budget vs
Actual screen, the MCP tools) runs the same steps:

1. read      file -> rows (header row detected; blank rows skipped; footer
             rows such as 'Total' with no month are skipped and reported)
2. map       column names through aliases (Month -> period, Account ->
             line_item, Actual -> amount, Ref -> reference_id, Memo -> notes).
             When several columns map to one field, the line item is the
             column whose values are statement lines ('Account Name' beats a
             GL-number 'Account'); the report says which column was used.
3. validate  every row: period is a month (normalize_month: '2026-04',
             'Apr 2026', '4/30/2026' ...), line_item is a statement line
             (key, label or alias, ignoring case: 'Net Sales' -> revenue),
             amount is a number in US format ('$1,234.50', '(500)' = -500,
             '500-' = -500); a decimal comma or misplaced commas ('12.500,00',
             '1,23') are refused, never guessed. A BLANK amount rejects the
             row ('no amount'); it is never stored as 0. '0' (or an accounting
             dash) is a confirmed zero.
4. duplicates  two rows with the same period + line_item + reference_id are
             an error naming both rows, unless ``combine_duplicates`` adds
             them up.
5. compare   with what is stored: each row is inserted / updated / unchanged
             / rejected. Amounts are stored per period + line + reference and
             ADDED UP, so ``line_totals`` gives each month + line of the
             request its stored total before, the request's amount, the total
             after and the rows under other references; a warning names them
             whenever they exist (a summary on top of detail would double
             count). ``replace_lines`` removes those rows in the same
             transaction (undo puts them back), so the line's month total
             becomes the request's: the hand-entry form uses it.
6. controls  ``control_total`` (the file must add up to it within 0.005) and
             closed months (a row that would change a closed month refuses
             the whole request). ``stored_after_total`` /
             ``control_matches_after`` compare the control total with the
             stored months after the import.
7. write     in ONE transaction. Atomic by default: any rejected row means
             nothing is written. ``allow_partial`` writes the valid rows and
             reports the rest; ``dry_run`` writes nothing and reports what
             would happen.

Every write records an ``actuals_batches`` row (batch_id 'ACT-xxxxxxxx')
and stamps ``actuals.batch_id`` on the rows it inserted or changed (a row it
found unchanged keeps the batch that stored it). Rows a batch inserted get
the actual_id '<batch_id>-<n>'. What the batch asked for is kept in
``actuals_batch_entries`` (each row written or found unchanged as sent, the
other references a replace_lines request named, a deletion's filter) and the
rows it overwrote, removed with replace_lines, found unchanged or deleted in
``actuals_batch_changes``. Undoing a batch leaves every month + line it
touched exactly as replaying, in their original order, the batches still in
place (not undone; deletions included) on top of the rows that predate them
would: rows it added go, rows it changed go back to the latest earlier batch
still in place (an undone batch's value never comes back), rows it replaced
come back unless a batch still in place wrote, replaced or deleted them since,
a later correction is kept and a row a later batch found unchanged stays with
that batch. Undoing a set of batches in any order leaves the same rows.

HTTP outcome (the routes and the MCP tools rely on it):

* a write that is refused writes NOTHING and answers 422 (409 when the
  reason is a closed month) with the full report in ``data`` and the
  explanation in ``detail``;
* a dry run always answers 200 with the same report: ``outcome`` says what a
  real import would do ('imported', 'partial', 'unchanged' or 'rejected').
"""
from __future__ import annotations

import csv
import difflib
import io
import math
import re
import uuid
from datetime import date, datetime, time as dt_time, timezone
from decimal import Decimal
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from services.financial_statements import SALES_DEDUCTION_DETAIL_LINES, VARIANCE_LINE_ITEMS, period_range
from utils.period_normalizer import normalize_month, normalize_period, period_kind

#: The sales deductions total line; a month carries either it or the detail lines (SALES_DEDUCTION_DETAIL_LINES).
DEDUCTIONS_TOTAL_LINE = "total_sales_deductions"

#: Two amounts closer than this are the same amount (and a control total ties within it).
TOLERANCE = 0.005
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
FILE_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm")


class ActualsError(ValueError):
    """A request-level problem (unreadable file, unknown batch, ...). ``status`` is the HTTP code."""

    def __init__(self, message: str, status: int = 422, data: Optional[Dict[str, Any]] = None):
        super().__init__(message)
        self.status = status
        self.data = data


def utc_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def month_label(period: Optional[str]) -> str:
    """'2026-04' -> 'Apr 2026' (anything else unchanged)."""
    if isinstance(period, str) and re.match(r"^\d{4}-\d{2}$", period):
        year, month = period.split("-")
        if 1 <= int(month) <= 12:
            return f"{date(int(year), int(month), 1):%b} {year}"
    return str(period)


def _months_text(periods: Iterable[str]) -> str:
    return ", ".join(month_label(p) for p in periods)


def _plural(n: int, word: str, plural: Optional[str] = None) -> str:
    return f"{n} {word if n == 1 else (plural or word + 's')}"


def _money(value: Optional[float]) -> str:
    if value is None:
        return "-"
    text = f"{abs(value):,.2f}"
    return f"({text})" if value < 0 else text


def _blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and value != value:
        return True
    return isinstance(value, str) and not value.strip()


def _text(value: Any) -> Optional[str]:
    """A cell as text: 901.0 -> '901', dates -> YYYY-MM-DD, blank -> None."""
    if _blank(value):
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def json_safe(value: Any) -> Any:
    """
    A spreadsheet cell as a JSON value, so every caller (the routes, the MCP
    server posting rows as JSON) can pass read_table's rows on: dates ->
    'YYYY-MM-DD' (a time of day other than midnight is kept: 'YYYY-MM-DD
    HH:MM:SS'), times -> 'HH:MM:SS', Decimal -> float, NaN -> None; numbers and
    text are unchanged and anything else becomes text.
    """
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, datetime):
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dt_time):
        return value.isoformat()
    if isinstance(value, Decimal):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


def _end_sentence(text: str) -> str:
    """'... line' -> '... line.'; text already ending in . ? or ! is left alone (no '..' or '?.')."""
    text = (text or "").rstrip()
    return text if not text or text[-1] in ".?!" else text + "."


# ---------------------------------------------------------------------------
# Columns
# ---------------------------------------------------------------------------

COLUMN_ALIASES: Dict[str, Tuple[str, ...]] = {
    "period": ("month", "fiscal_period", "posting_period", "gl_period", "accounting_period", "period_name",
               "month_end", "period_end", "period_ending", "month_ending", "date", "yyyy_mm"),
    "line_item": ("line", "line_name", "statement_line", "line_item_name", "account", "account_name",
                  "gl_account", "gl_line", "category", "caption", "item", "financial_line"),
    "amount": ("actual", "actuals", "actual_amount", "value", "amount_usd", "usd", "net_amount", "net",
               "balance", "total", "activity", "net_change", "amt"),
    "reference_id": ("reference", "ref", "ref_id", "ref_no", "reference_no", "reference_number", "document",
                     "document_number", "document_no", "doc", "doc_no", "journal", "journal_entry", "je",
                     "je_number", "entry_id", "voucher", "reference_id"),
    "notes": ("note", "memo", "comment", "comments", "description", "remarks"),
    "source": ("source_system", "origin"),
}


def header_key(name: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


_COLUMN_LOOKUP: Dict[str, Tuple[str, int]] = {}
for _canonical, _aliases in COLUMN_ALIASES.items():
    for _rank, _alias in enumerate((_canonical, *_aliases)):
        _COLUMN_LOOKUP.setdefault(header_key(_alias), (_canonical, _rank))


def column_field(name: Any) -> Optional[str]:
    hit = _COLUMN_LOOKUP.get(header_key(name))
    return hit[0] if hit else None


def map_columns(row: Dict[str, Any], prefer: Optional[Dict[str, str]] = None) -> Tuple[Dict[str, Any], List[str]]:
    """
    Rename a raw row's keys to the canonical fields; returns (mapped, unknown column names).
    When several columns map to one field the best alias wins, unless ``prefer`` names the
    column to use for that field (see choose_columns).
    """
    out: Dict[str, Any] = {}
    ranks: Dict[str, int] = {}
    unknown: List[str] = []
    for raw_key, value in row.items():
        hit = _COLUMN_LOOKUP.get(header_key(raw_key))
        if not hit:
            if not _blank(raw_key):
                unknown.append(str(raw_key))
            continue
        canonical, rank = hit
        if prefer and prefer.get(canonical) == str(raw_key):
            rank = -1
        if canonical in out and rank >= ranks[canonical]:
            continue
        if canonical in out and _blank(value):
            continue
        out[canonical] = value
        ranks[canonical] = rank
    return out, unknown


FIELD_NAMES = {"period": "period", "line_item": "line item", "amount": "amount", "reference_id": "reference",
               "notes": "notes", "source": "source"}


def _quoted(names: Sequence[str]) -> str:
    quoted = [f"'{n}'" for n in names]
    return quoted[0] if len(quoted) == 1 else ", ".join(quoted[:-1]) + f" and {quoted[-1]}"


def choose_columns(rows: List[Tuple[int, Dict[str, Any]]]) -> Tuple[Dict[str, str], List[str]]:
    """
    When several columns of a file map to one field (a trial balance with 'Account' = the GL
    number and 'Account Name'), pick one and say which. The line item is chosen by its values:
    the column whose values are statement lines on the most rows (ties: the better alias).
    Other fields keep the better alias. Returns ({field: column to use}, warnings).
    """
    candidates: Dict[str, List[Tuple[int, str]]] = {}
    seen = set()
    for _row, raw in rows:
        for key in raw:
            if key in seen:
                continue
            seen.add(key)
            hit = _COLUMN_LOOKUP.get(header_key(key))
            if hit:
                candidates.setdefault(hit[0], []).append((hit[1], str(key)))
    prefer: Dict[str, str] = {}
    warnings: List[str] = []
    resolved_cache: Dict[str, bool] = {}

    def is_line(value: Any) -> bool:
        key = str(value)
        if key not in resolved_cache:
            resolved_cache[key] = resolve_line_item(value)[0] is not None
        return resolved_cache[key]

    for field, columns in candidates.items():
        if len(columns) < 2:
            continue
        columns.sort()
        names = [name for _rank, name in columns]
        label = FIELD_NAMES.get(field, field)
        if field != "line_item":
            warnings.append(f"Columns {_quoted(names)} could each be the {label}: used '{names[0]}'.")
            continue
        scored = []
        for rank, name in columns:
            values = [raw.get(name) for _row, raw in rows if not _blank(raw.get(name))]
            resolved = sum(1 for v in values if is_line(v))
            scored.append((resolved, -rank, name, values))
        best = max(scored, key=lambda s: (s[0], s[1]))
        if best[0] == 0:
            warnings.append(f"Columns {_quoted(names)} could each be the line item, but none holds statement "
                            f"line names, so '{names[0]}' was used. Map the accounts to statement lines (e.g. "
                            f"'Net Sales', 'Payroll'), or remove or rename the columns that are not the line item.")
            continue
        prefer[field] = best[2]
        others = []
        for resolved, _r, name, values in scored:
            if name == best[2]:
                continue
            sample = next((_text(v) for v in values if not is_line(v)), None)
            others.append(f"'{name}'" + (f" (e.g. '{sample}')" if sample else ""))
        rows_with = len([1 for _row, raw in rows if not _blank(raw.get(best[2]))])
        warnings.append(f"Used '{best[2]}' for the line item ({best[0]} of {rows_with} rows are statement lines); "
                        f"not used: {', '.join(others)}.")
    return prefer, warnings


# ---------------------------------------------------------------------------
# Line items
# ---------------------------------------------------------------------------

#: Common names for each statement line (matched ignoring case, spaces and punctuation; '&' = 'and').
LINE_ALIASES: Dict[str, Tuple[str, ...]] = {
    "gross_sales": ("gross revenue", "total gross sales", "sales gross", "gross sales revenue"),
    # The gross-to-net detail lines (Sales Deductions is their sum: a file carries one or the other per month).
    "terms_allowance": ("terms", "terms allowances", "terms discount", "terms discounts", "cash discount",
                        "cash discounts", "early payment discount", "early payment discounts",
                        "prompt payment discount", "prompt payment discounts", "payment terms discount"),
    "chargebacks": ("chargeback", "charge backs", "charge back", "chargeback allowance", "chargebacks allowances",
                    "customer chargebacks"),
    "other_allowances": ("other allowances", "other allowance", "other discounts", "other discounts and allowances",
                         "other sales allowances", "rebates", "volume rebates", "customer rebates",
                         "co op advertising", "coop allowance"),
    "excise_tax": ("excise", "excise taxes", "fet", "federal excise tax", "federal excise taxes", "federal excise",
                   "fet tax", "excise tax fet", "fet excise tax", "firearms excise tax"),
    "total_sales_deductions": ("deductions", "sales deductions", "total deductions", "total sales deductions",
                               "discounts and allowances", "sales discounts", "sales discounts and allowances",
                               "sales returns and allowances", "gross to net", "gross to net deductions"),
    "revenue": ("net sales", "net revenue", "net revenues", "total net sales", "total revenue", "sales net",
                "revenue net", "net sales revenue", "revenues"),
    "cogs_materials": ("materials", "material", "material cost", "materials cost", "direct materials",
                       "direct material", "cogs materials", "raw materials"),
    "cogs_labor": ("direct labor", "direct labour", "direct labor cost", "cogs labor", "cogs direct labor"),
    "cogs_manufacturing": ("machine", "machine cost", "machine costs", "machine time", "machine and mfg",
                           "machine mfg", "cogs machine", "cogs manufacturing", "manufacturing cost"),
    "factory_overhead": ("manufacturing overhead", "plant overhead", "mfg overhead", "cogs overhead",
                         "cogs factory overhead"),
    "other_cogs": ("other cost of sales", "other cost of goods sold", "cogs other"),
    "total_cogs": ("cogs", "cost of goods sold", "cost of sales", "total cost of sales",
                   "total cost of goods sold"),
    "gross_profit": ("gross margin",),
    "commissions": ("commission", "sales commission", "commission expense", "commissions expense"),
    "payroll": ("salaries", "wages", "salaries and wages", "wages and salaries", "payroll expense",
                "payroll and benefits", "personnel", "personnel costs", "staff costs"),
    "admin_expense": ("admin", "admin expense", "administrative", "administrative expense",
                      "administrative expenses", "general and administrative", "g and a",
                      "g and a expense", "g and a expenses"),
    "operating_expense": ("opex", "operating expense", "other operating expenses", "operating costs",
                          "other opex"),
    "depreciation": ("depreciation expense", "depreciation and amortization", "d and a"),
    "operating_income": ("ebit", "operating profit", "income from operations", "operating income loss"),
    "interest_expense": ("interest", "interest paid"),
    "income_tax": ("income taxes", "income tax expense", "provision for income taxes", "tax expense"),
    "net_income": ("net profit", "net earnings", "net income loss", "net profit loss"),
    "capex": ("capital expenditures", "capital expenditure", "capital spending", "purchases of fixed assets",
              "purchase of equipment", "fixed asset purchases", "capital purchases"),
    "closing_cash": ("cash", "ending cash", "cash balance", "closing cash balance", "ending cash balance",
                     "cash at end of period", "end of period cash"),
    # Lines added by the labor classification (only used when the engine lists them).
    "indirect_labor": ("indirect labor", "indirect labour", "factory indirect labor"),
    "labor_variance": ("labor variance", "labour variance", "direct labor variance"),
}

#: Names that could mean more than one line: the message says which to use.
AMBIGUOUS_LINES: Dict[str, str] = {
    "sales": "'{text}' could be gross or net sales: use 'Gross Sales' or 'Net Sales' (after deductions)",
    "labor": "'{text}' could be direct labor or payroll: use 'Direct Labor' (COGS) or 'Payroll' (operating expense)",
    "labour": "'{text}' could be direct labor or payroll: use 'Direct Labor' (COGS) or 'Payroll' (operating expense)",
    "tax": "'{text}' is ambiguous: use 'Income Tax' or 'Excise tax (FET)' (a sales deduction)",
    "taxes": "'{text}' is ambiguous: use 'Income Tax' or 'Excise tax (FET)' (a sales deduction)",
    "overhead": "'{text}' is ambiguous: use 'Factory Overhead' (COGS) or 'Admin Expenses'",
    "profit": "'{text}' is ambiguous: use 'Gross Profit', 'Operating Income' or 'Net Income'",
    "expenses": "'{text}' is ambiguous: use 'Operating Expenses', 'Admin Expenses' or 'Payroll'",
}


def line_key(text: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text).lower().replace("&", " and "))


def _line_lookup() -> Dict[str, str]:
    """normalised name -> line_item key: keys, then labels, then aliases (first wins)."""
    lookup: Dict[str, str] = {}
    for key in VARIANCE_LINE_ITEMS:
        lookup.setdefault(line_key(key), key)
    for key, meta in VARIANCE_LINE_ITEMS.items():
        lookup.setdefault(line_key(meta.get("label", key)), key)
    for key, aliases in LINE_ALIASES.items():
        if key not in VARIANCE_LINE_ITEMS:
            continue
        for alias in aliases:
            lookup.setdefault(line_key(alias), key)
    return lookup


def line_label(key: Optional[str]) -> Optional[str]:
    meta = VARIANCE_LINE_ITEMS.get(key or "")
    return meta.get("label", key) if meta else key


def resolve_line_item(value: Any) -> Tuple[Optional[str], Optional[str]]:
    """(line_item key, None) or (None, why not)."""
    text = _text(value)
    if text is None:
        return None, "no line item"
    norm = line_key(text)
    lookup = _line_lookup()
    if norm in lookup:
        return lookup[norm], None
    if norm in AMBIGUOUS_LINES:
        return None, AMBIGUOUS_LINES[norm].format(text=text)
    labels = [meta.get("label", key) for key, meta in VARIANCE_LINE_ITEMS.items()]
    by_norm = {line_key(label): label for label in labels}
    close = difflib.get_close_matches(norm, list(by_norm), n=2, cutoff=0.75)
    hint = f". Did you mean {' or '.join(repr(by_norm[c]) for c in close)}?" if close else ""
    return None, f"line item '{text}' is not a statement line{hint}"


# ---------------------------------------------------------------------------
# Values
# ---------------------------------------------------------------------------

_ACCOUNTING_ZEROS = {"-", "--", "–", "—", "−"}
_TOTAL_RE = re.compile(r"^\s*(grand\s+|sub\s*-?\s*)?totals?\b", re.IGNORECASE)
#: US number format: commas only between groups of three digits, '.' for decimals ('1,234.50', '1234.5', '1.5E+03').
_NUMBER_RE = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)?(?:\.\d*)?(?:[eE][+-]?\d+)?$")
#: European format: '.' between thousands and / or a decimal comma ('12.500,00', '1.234,50', '1,23', '1.234.567').
_DECIMAL_COMMA_RE = re.compile(r"^[+-]?(?:\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+,\d+)$")
#: Spaces between thousands ('1 000'): plain, no-break, figure, thin and narrow no-break spaces.
_SPACES = (" ", "\u00a0", "\u2007", "\u2009", "\u202f")


def parse_amount(value: Any) -> Tuple[Optional[float], Optional[str]]:
    """
    (amount, None) or (None, why not). Blank is 'no amount', never 0. Text is read in US
    format ('$1,234.50', '(500)' = -500, '500-' = -500, '1 000'); a decimal comma or
    misplaced commas ('12.500,00', '1,23', '1,2,3') are refused, never guessed: stripping
    the commas would turn 12.500,00 into 12.50.
    """
    if _blank(value):
        return None, "no amount (leave the row out, or enter 0 for a confirmed zero)"
    if isinstance(value, bool):
        return None, f"amount must be a number, got {value}"
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            return None, f"amount {value} is not a number"
        return float(value), None
    raw = str(value).strip()
    text = raw.replace("$", "").replace("\u2212", "-")
    for space in _SPACES:
        text = text.replace(space, "")
    if text in _ACCOUNTING_ZEROS:
        return 0.0, None
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
    elif text.startswith("-(") and text.endswith(")"):
        negative, text = True, text[2:-1]
    elif len(text) > 1 and text.endswith("-") and re.match(r"^\d[\d.,]*-$", text):
        negative, text = True, text[:-1]   # SAP-style trailing minus: 500-
    if not _NUMBER_RE.match(text):
        if _DECIMAL_COMMA_RE.match(text):
            return None, (f"amount '{raw}' looks like a European number (decimal comma, '.' between thousands): "
                          f"save the file with '.' for decimals and ',' or nothing between thousands "
                          f"(e.g. 12500.00 or 12,500.00), then import again")
        if "," in text:
            return None, f"amount '{raw}' is not a number: commas may only separate thousands (e.g. 1,234.50)"
        return None, f"amount '{raw}' is not a number"
    try:
        number = float(text.replace(",", ""))
    except ValueError:
        return None, f"amount '{raw}' is not a number"
    if not math.isfinite(number):
        return None, f"amount '{raw}' is not a number"
    return (-number if negative else number), None


def parse_period(value: Any) -> Tuple[Optional[str], Optional[str]]:
    if _blank(value):
        return None, "no period (month)"
    month = normalize_month(value)
    if month:
        return month, None
    shown = _text(value)
    kind = period_kind(normalize_period(value))
    if kind == "quarter":
        return None, f"period '{shown}' is a quarter: actuals are entered by month (YYYY-MM, e.g. 2026-04)"
    if kind == "week":
        return None, f"period '{shown}' is a week: actuals are entered by month (YYYY-MM, e.g. 2026-04)"
    return None, f"period '{shown}' is not a month: use YYYY-MM, e.g. 2026-04"


def parse_control_total(value: Any) -> Optional[float]:
    if _blank(value):
        return None
    amount, error = parse_amount(value)
    if error:
        raise ActualsError(f"control_total: {error}")
    return amount


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------

def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _read_csv(raw: bytes) -> List[List[Any]]:
    text = _decode(raw)
    sample = text[:4096]
    first = sample.splitlines()[0] if sample.splitlines() else ""
    delimiter = ","
    if "," not in first:
        for candidate in ("\t", ";", "|"):
            if candidate in first:
                delimiter = candidate
                break
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(raw: bytes, sheet: Optional[str]) -> Tuple[List[List[Any]], str]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - openpyxl is a dependency
        raise ActualsError("Excel files cannot be read on this server: save the file as CSV")
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ActualsError(f"The file could not be read as an Excel workbook ({exc}). Save it as .xlsx or .csv.")
    try:
        names = workbook.sheetnames
        if sheet:
            match = next((n for n in names if n.strip().lower() == str(sheet).strip().lower()), None)
            if match is None:
                raise ActualsError(f"Sheet '{sheet}' is not in the workbook (sheets: {', '.join(names)})")
            candidates = [match]
        else:
            active = workbook.active.title if workbook.active is not None else names[0]
            candidates = [active] + [n for n in names if n != active]
        chosen_grid: Optional[List[List[Any]]] = None
        chosen_name = candidates[0]
        for name in candidates:
            grid = [list(row) for row in workbook[name].iter_rows(values_only=True)]
            if len(grid) > MAX_ROWS + 50:
                raise ActualsError(f"Sheet '{name}' has more than {MAX_ROWS:,} rows; split the file")
            if chosen_grid is None:
                chosen_grid, chosen_name = grid, name
            if _detect_header(grid, strict=True) is not None:
                return grid, name
        return chosen_grid or [], chosen_name
    finally:
        workbook.close()


def _detect_header(grid: Sequence[Sequence[Any]], strict: bool = False, scan: int = 30) -> Optional[int]:
    """0-based header row: the first row (of ``scan``) with 2+ known column names; else the first filled row."""
    first = None
    for index, row in enumerate(list(grid)[:scan]):
        cells = [c for c in (row or []) if not _blank(c)]
        if not cells:
            continue
        if first is None:
            first = index
        known = {column_field(c) for c in cells if isinstance(c, str) and column_field(c)}
        if len(known) >= 2:
            return index
    return None if strict else first


def read_table(filename: Optional[str], raw: bytes, sheet: Optional[str] = None,
               header_row: Optional[int] = None) -> Tuple[List[Tuple[int, Dict[str, Any]]], Dict[str, Any]]:
    """
    Parse an uploaded CSV / XLSX into [(spreadsheet row number, {header: value})].
    The header row is ``header_row`` (1-based) when given, else detected (a report title above it is
    skipped); blank rows are dropped. Values are JSON-safe (json_safe: Excel dates -> 'YYYY-MM-DD'),
    so the rows can be posted on as JSON entries.
    """
    name = (filename or "").strip()
    lower = name.lower()
    if len(raw) > MAX_FILE_BYTES:
        raise ActualsError(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB; split it", status=413)
    if not raw.strip():
        raise ActualsError("The file is empty")
    if lower.endswith(".xls"):
        raise ActualsError("Old .xls workbooks cannot be read: save the file as .xlsx or .csv")
    is_xlsx = lower.endswith((".xlsx", ".xlsm")) or raw[:4] == b"PK\x03\x04"
    sheet_name = None
    if is_xlsx:
        grid, sheet_name = _read_xlsx(raw, sheet)
    else:
        if sheet:
            raise ActualsError("sheet applies to Excel files only")
        grid = _read_csv(raw)
    if header_row is not None:
        if header_row < 1 or header_row > len(grid):
            raise ActualsError(f"header_row {header_row} is outside the file ({len(grid)} rows)")
        header_index = header_row - 1
    else:
        header_index = _detect_header(grid)
    if header_index is None:
        raise ActualsError("The file has no rows")
    headers: List[Optional[str]] = []
    for i, cell in enumerate(grid[header_index] or []):
        text = _text(cell)
        headers.append(text if text else None)
    rows: List[Tuple[int, Dict[str, Any]]] = []
    for offset, raw_row in enumerate(grid[header_index + 1:], start=header_index + 2):
        record: Dict[str, Any] = {}
        for i, value in enumerate(raw_row or []):
            value = json_safe(value)
            if i >= len(headers):
                if not _blank(value):
                    record[f"column_{i + 1}"] = value
                continue
            key = headers[i] if headers[i] is not None else f"column_{i + 1}"
            if key in record and _blank(value):
                continue
            record[key] = value
        if all(_blank(v) for v in record.values()):
            continue
        rows.append((offset, record))
    if len(rows) > MAX_ROWS:
        raise ActualsError(f"The file has more than {MAX_ROWS:,} rows; split it")
    info = {"filename": name or None, "sheet": sheet_name, "header_row": header_index + 1,
            "columns": [h for h in headers if h]}
    return rows, info


def rows_from_entries(entries: Any, first_row_number: int = 1,
                      row_numbers: Optional[Sequence[Any]] = None) -> List[Tuple[int, Dict[str, Any]]]:
    """
    JSON entries -> [(row number, entry)]. Row numbers start at ``first_row_number``, or are given per
    entry by ``row_numbers`` (the spreadsheet rows of a file read elsewhere, e.g. by the MCP server).
    """
    if not isinstance(entries, list) or not entries:
        raise ActualsError("entries must be a non-empty list of {period, line_item, amount, reference_id?, notes?}")
    if row_numbers is not None:
        if not isinstance(row_numbers, list) or len(row_numbers) != len(entries):
            raise ActualsError("row_numbers must list one row number per entry")
        try:
            numbers = [int(n) for n in row_numbers]
        except (TypeError, ValueError):
            raise ActualsError("row_numbers must be whole numbers")
    else:
        numbers = [first_row_number + i for i in range(len(entries))]
    out: List[Tuple[int, Dict[str, Any]]] = []
    for index, entry in enumerate(entries):
        if not isinstance(entry, dict):
            raise ActualsError(f"entries[{index}] must be an object of column: value")
        out.append((numbers[index], entry))
    return out


# ---------------------------------------------------------------------------
# Closed months
# ---------------------------------------------------------------------------

def closed_periods(cursor, periods: Optional[Iterable[str]] = None) -> List[str]:
    cursor.execute("SELECT period FROM actuals_periods WHERE status = 'closed' ORDER BY period")
    closed = [r[0] for r in cursor.fetchall()]
    if periods is None:
        return closed
    wanted = set(periods)
    return [p for p in closed if p in wanted]


def closed_in_window(cursor, start_period: str, end_period: str) -> List[str]:
    cursor.execute("SELECT period FROM actuals_periods WHERE status = 'closed' AND period >= ? AND period <= ? "
                   "ORDER BY period", (start_period, end_period))
    return [r[0] for r in cursor.fetchall()]


def closed_message(periods: Sequence[str], action: str = "change") -> str:
    many = len(periods) != 1
    return (f"{_months_text(periods)} {'are' if many else 'is'} closed: actuals for a closed month cannot be "
            f"{'changed' if action == 'change' else action}. Reopen {'them' if many else 'it'} (with a reason) first.")


# ---------------------------------------------------------------------------
# The import pipeline
# ---------------------------------------------------------------------------

def _new_batch_id(cursor) -> str:
    while True:
        batch_id = f"ACT-{uuid.uuid4().hex[:8]}"
        cursor.execute("SELECT 1 FROM actuals_batches WHERE batch_id = ?", (batch_id,))
        if not cursor.fetchone():
            return batch_id


def _note_or_total_row(row: Dict[str, Any], mapped: Dict[str, Any]) -> Optional[str]:
    """A footer row ('Total', 'Grand total' with no month) or a lone note: skipped and reported, not imported."""
    cells = [v for v in row.values() if not _blank(v)]
    if not cells:
        return None
    if normalize_month(mapped.get("period")) is None:
        first_text = _text(cells[0]) if isinstance(cells[0], str) else None
        for text in (first_text, _text(mapped.get("line_item"))):
            if text and _TOTAL_RE.match(text) and resolve_line_item(text)[0] is None:
                return text
    if len(cells) == 1 and isinstance(cells[0], str):
        text = cells[0].strip()
        if text.lower().startswith(("note:", "notes:", "*")) or len(text.split()) >= 5:
            return text
    return None


def _totals(parsed: List[Dict[str, Any]]) -> Dict[str, Any]:
    """File totals by line item, by month and overall, over every row with a readable amount."""
    by_line: Dict[str, Dict[str, Any]] = {}
    by_period: Dict[str, Dict[str, Any]] = {}
    grand = 0.0
    counted = 0
    order = {key: i for i, key in enumerate(VARIANCE_LINE_ITEMS)}
    for item in parsed:
        amount = item.get("amount")
        if amount is None:
            continue
        counted += 1
        grand += amount
        key = item.get("line_item")
        line_name = key or f"?{item.get('line_text') or '(blank)'}"
        entry = by_line.setdefault(line_name, {
            "line_item": key, "label": line_label(key) if key else (item.get("line_text") or "(blank)"),
            "recognised": bool(key), "amount": 0.0, "rows": 0,
        })
        entry["amount"] += amount
        entry["rows"] += 1
        period = item.get("period")
        period_name = period or f"?{item.get('period_text') or '(blank)'}"
        pentry = by_period.setdefault(period_name, {
            "period": period, "label": month_label(period) if period else (item.get("period_text") or "(blank)"),
            "valid": bool(period), "amount": 0.0, "rows": 0,
        })
        pentry["amount"] += amount
        pentry["rows"] += 1
    lines = sorted(by_line.values(), key=lambda e: (not e["recognised"], order.get(e["line_item"], 999), e["label"]))
    periods = sorted(by_period.values(), key=lambda e: (not e["valid"], e["period"] or "", e["label"]))
    for entry in (*lines, *periods):
        entry["amount"] = round(entry["amount"], 2)
    return {"by_line_item": lines, "by_period": periods, "grand_total": round(grand, 2), "rows": counted}


def _tie_out_warnings(planned: List[Dict[str, Any]]) -> List[str]:
    """Gross - deductions = net sales and net sales - COGS = gross profit, where a month has all three."""
    by_period: Dict[str, Dict[str, float]] = {}
    for item in planned:
        if item.get("reference_id"):
            continue
        by_period.setdefault(item["period"], {})[item["line_item"]] = item["amount"]
    out: List[str] = []
    for lines in by_period.values():
        # Deductions given as the detail lines (terms, chargebacks, other allowances, excise) tie out as their sum.
        details = [k for k in SALES_DEDUCTION_DETAIL_LINES if k in lines]
        if details and DEDUCTIONS_TOTAL_LINE not in lines:
            lines[DEDUCTIONS_TOTAL_LINE] = sum(lines[k] for k in details)
    checks = (
        ("gross_sales", "total_sales_deductions", "revenue", "Gross Sales - Sales Deductions"),
        ("revenue", "total_cogs", "gross_profit", "Net Sales - Total COGS"),
    )
    for period in sorted(by_period):
        lines = by_period[period]
        for a, b, c, label in checks:
            if a in lines and b in lines and c in lines and c in VARIANCE_LINE_ITEMS:
                expected = lines[a] - lines[b]
                if abs(expected - lines[c]) > 1.0:
                    out.append(f"{month_label(period)}: {label} = {_money(expected)} but {line_label(c)} is "
                               f"{_money(lines[c])} (difference {_money(lines[c] - expected)})")
    return out


def _sign_warnings(planned: List[Dict[str, Any]]) -> List[str]:
    """Expense lines entered as negative numbers are usually a sign-convention slip."""
    flagged = [item for item in planned
               if item["amount"] < 0 and VARIANCE_LINE_ITEMS.get(item["line_item"], {}).get("direction") == "expense"
               and item["line_item"] not in ("labor_variance",)]
    if not flagged:
        return []
    shown = "; ".join(f"row {item['row']} {line_label(item['line_item'])} {month_label(item['period'])} "
                      f"{_money(item['amount'])}" for item in flagged[:5])
    more = f" and {len(flagged) - 5} more" if len(flagged) > 5 else ""
    return [f"{_plural(len(flagged), 'expense amount')} {'is' if len(flagged) == 1 else 'are'} negative ({shown}{more}). "
            f"Costs are entered as positive numbers; a negative cost lowers the expense."]


#: Records what a write overwrote ('updated') or removed ('removed') so undo can put it back exactly.
_RECORD_CHANGE = """INSERT OR IGNORE INTO actuals_batch_changes (batch_id, actual_id, previous_amount, previous_notes,
        previous_source, previous_batch_id, previous_updated_at, change_type, previous_period, previous_line_item,
        previous_reference_id, previous_created_at)
    SELECT ?, actual_id, amount, notes, source, batch_id, updated_at, ?, period, line_item, reference_id, created_at
    FROM actuals WHERE actual_id = ?"""

_LOG_ENTRY = """INSERT INTO actuals_batch_entries (batch_id, seq, action, period, line_item, reference_id, amount,
        notes, source, actual_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"""


def _log_request(cursor, batch_id: str, planned: List[Dict[str, Any]], parsed: List[Dict[str, Any]],
                 replace_lines: bool) -> None:
    """
    Keep what the batch asked for, so undo can replay it: every row it wrote or found unchanged (amount,
    notes and source as sent, the row id) and, with replace_lines, the other references the request
    named for a month + line it wrote (a rejected row's reference is not replaced away).
    """
    entries = [(item["action"], item["period"], item["line_item"], item["reference_id"], item["amount"],
                item["notes"], item["source"], item["actual_id"]) for item in planned]
    if replace_lines:
        pairs = {(i["period"], i["line_item"]) for i in planned}
        written = {(i["period"], i["line_item"], i["reference_id"]) for i in planned}
        named = {(p["period"], p["line_item"], p["reference_id"]) for p in parsed
                 if p["period"] and p["line_item"] and (p["period"], p["line_item"]) in pairs}
        entries.extend(("keep", *key, None, None, None, None) for key in sorted(named - written))
    cursor.executemany(_LOG_ENTRY, [(batch_id, seq, *entry) for seq, entry in enumerate(entries)])


def _refs_text(rows: Sequence[Dict[str, Any]], limit: int = 3) -> str:
    shown = "; ".join(f"{r['reference_id'] or '(no reference)'} {_money(r['amount'])}" for r in rows[:limit])
    return shown + (f" and {len(rows) - limit} more" if len(rows) > limit else "")


def _line_totals(planned: List[Dict[str, Any]], parsed: List[Dict[str, Any]],
                 stored_by_pair: Dict[Tuple[str, str], List[Dict[str, Any]]],
                 replace_lines: bool) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Per month + line in the request: the stored total before, the amount in the request, the
    total after, and the rows stored under other references (references the request does not
    mention for that month and line). Amounts are stored per (period, line, reference) and
    added up, so those rows count in the month's line total unless replace_lines removes them.
    Returns (line totals, stored rows to remove).
    """
    refs_in_request: Dict[Tuple[str, str], set] = {}
    for item in parsed:
        if item["period"] and item["line_item"]:
            refs_in_request.setdefault((item["period"], item["line_item"]), set()).add(item["reference_id"])
    by_pair: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
    for item in planned:
        by_pair.setdefault((item["period"], item["line_item"]), []).append(item)
    order = {key: i for i, key in enumerate(VARIANCE_LINE_ITEMS)}
    totals: List[Dict[str, Any]] = []
    to_remove: List[Dict[str, Any]] = []
    for period, line_item in sorted(by_pair, key=lambda k: (k[0], order.get(k[1], 999), k[1])):
        items = by_pair[(period, line_item)]
        planned_refs = {i["reference_id"] for i in items}
        rows_now = stored_by_pair.get((period, line_item), [])
        others = [r for r in rows_now if (r["reference_id"] or "") not in refs_in_request[(period, line_item)]]
        removed_ids = {r["actual_id"] for r in others} if replace_lines else set()
        before = sum(float(r["amount"] or 0) for r in rows_now)
        request_amount = sum(i["amount"] for i in items)
        kept = sum(float(r["amount"] or 0) for r in rows_now
                   if (r["reference_id"] or "") not in planned_refs and r["actual_id"] not in removed_ids)
        if replace_lines:
            to_remove.extend(others)
        totals.append({
            "period": period, "label": month_label(period), "line_item": line_item,
            "line_label": line_label(line_item),
            "before": round(before, 2), "file_amount": round(request_amount, 2), "after": round(kept + request_amount, 2),
            "other_references": [{"reference_id": r["reference_id"] or "", "amount": round(float(r["amount"] or 0), 2),
                                  "batch_id": r.get("batch_id"), "actual_id": r["actual_id"]} for r in others],
            "other_references_amount": round(sum(float(r["amount"] or 0) for r in others), 2),
            "replaced": bool(removed_ids),
        })
    return totals, to_remove


def _other_reference_warnings(line_totals: List[Dict[str, Any]], replace_lines: bool) -> List[str]:
    """Say when a month's line holds amounts under references the import does not include (they add up)."""
    flagged = [t for t in line_totals if t["other_references"]]
    out: List[str] = []
    for t in flagged[:5]:
        others = t["other_references"]
        what = f"{t['label']} {t['line_label']}"
        if replace_lines:
            out.append(f"{what}: replacing {_plural(len(others), 'amount')} stored under other references "
                       f"({_refs_text(others)}); the month's line total becomes {_money(t['after'])} "
                       f"(was {_money(t['before'])}). Undoing this import puts them back.")
            continue
        batches = sorted({o["batch_id"] for o in others if o.get("batch_id")})
        undo_hint = (f", or undo {'import' if len(batches) == 1 else 'imports'} {', '.join(batches)}"
                     if batches else "")
        out.append(f"{what} already holds {_plural(len(others), 'amount')} under other references "
                   f"({_refs_text(others)}) that this import does not include: they are kept and added, so the "
                   f"month's line total becomes {_money(t['after'])}, not {_money(t['file_amount'])}. If this import "
                   f"replaces them, import it with 'Replace these months' lines' (replace_lines){undo_hint}.")
    if len(flagged) > 5:
        more = len(flagged) - 5
        out.append(f"{more} more month{'s' if more != 1 else ''} and line{'s' if more != 1 else ''} "
                   f"{'hold' if more != 1 else 'holds'} amounts under other references (see line_totals).")
    return out


def run_import(conn, rows: List[Tuple[int, Dict[str, Any]]], *, dry_run: bool = False, allow_partial: bool = False,
               combine_duplicates: bool = False, control_total: Optional[float] = None,
               source: Optional[str] = None, filename: Optional[str] = None, note: Optional[str] = None,
               default_notes: Optional[str] = None, file_info: Optional[Dict[str, Any]] = None,
               replace_lines: bool = False) -> Tuple[int, Dict[str, Any]]:
    """
    Validate, reconcile and (unless dry_run or refused) write actuals in one transaction.
    Returns (HTTP status, report). See the module docstring for the rules.

    ``replace_lines``: for every month + line in the request, rows stored under other
    references are removed in the same transaction (recorded, so undo puts them back), so
    the month's line total becomes the request's. Without it they are kept and added, and the
    report warns. ``line_totals`` gives every month + line's stored total before and after.
    """
    cursor = conn.cursor()
    warnings: List[str] = []
    skipped: List[Dict[str, Any]] = []
    unknown_columns: List[str] = []
    parsed: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []

    # 1-3: map and validate every row (several candidate columns for one field: choose_columns)
    prefer, column_warnings = choose_columns(rows)
    warnings.extend(column_warnings)
    for row_no, raw in rows:
        mapped, unknown = map_columns(raw, prefer)
        for column in unknown:
            if column not in unknown_columns and not column.startswith("column_"):
                unknown_columns.append(column)
        footer = _note_or_total_row(raw, mapped)
        if footer is not None:
            skipped.append({"row": row_no, "text": footer[:80]})
            continue
        problems: List[str] = []
        period, period_error = parse_period(mapped.get("period"))
        line_item, line_error = resolve_line_item(mapped.get("line_item"))
        amount, amount_error = parse_amount(mapped.get("amount"))
        for problem in (period_error, line_error, amount_error):
            if problem:
                problems.append(problem)
        reference_id = _text(mapped.get("reference_id")) or ""
        if len(reference_id) > 120:
            problems.append("reference is longer than 120 characters")
        notes = _text(mapped.get("notes")) or default_notes
        item = {
            "row": row_no, "period": period, "line_item": line_item, "reference_id": reference_id,
            "amount": amount, "notes": notes, "source": _text(mapped.get("source")) or source,
            "period_text": _text(mapped.get("period")), "line_text": _text(mapped.get("line_item")),
            "problems": problems,
        }
        parsed.append(item)

    if not rows:
        raise ActualsError("No rows to import")
    if not parsed:
        raise ActualsError("No data rows to import: every row is blank, a note or a total")
    if unknown_columns:
        warnings.append(f"Ignored columns (not recognised): {', '.join(unknown_columns)}")
    if skipped:
        shown = "; ".join(f"row {s['row']} ('{s['text']}')" for s in skipped[:5])
        what = "looks like a total or a note" if len(skipped) == 1 else "look like totals or notes"
        warnings.append(f"Skipped {_plural(len(skipped), 'row')} that {what}, not data: {shown}"
                        + (f" and {len(skipped) - 5} more" if len(skipped) > 5 else ""))
    present = {key for _row, raw in rows for key in map_columns(raw, prefer)[0]}
    missing_columns = [c for c in ("period", "line_item", "amount") if c not in present]
    if missing_columns:
        found = (file_info or {}).get("columns") or sorted({str(k) for _row, raw in rows for k in raw})
        raise ActualsError(
            f"No {' / '.join(missing_columns)} column. The file needs one row per month and line with the "
            f"columns period, line_item and amount (optional: reference_id, notes). Found: {', '.join(found) or 'none'}.")

    totals = _totals(parsed)
    request_errors: List[str] = []
    control_difference = None
    control_matches = None
    if control_total is not None:
        control_difference = round(totals["grand_total"] - control_total, 2)
        control_matches = abs(totals["grand_total"] - control_total) <= TOLERANCE
        if not control_matches:
            request_errors.append(
                f"The file adds up to {_money(totals['grand_total'])} but the control total is "
                f"{_money(control_total)} (difference {_money(control_difference)}).")

    # 4: duplicates within the request
    # Only rows that are otherwise valid can clash: a row already rejected (bad amount, ...) is left out
    # anyway, so it only gets a pointer to the valid row with the same key.
    groups: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
    for item in parsed:
        if item["period"] and item["line_item"] and not item["problems"]:
            groups.setdefault((item["period"], item["line_item"], item["reference_id"]), []).append(item)
    for item in parsed:
        if item["problems"] and item["period"] and item["line_item"]:
            same = groups.get((item["period"], item["line_item"], item["reference_id"])) or []
            if same:
                rows_text = _list([g["row"] for g in same])
                verb = "has" if len(same) == 1 else "have"
                item["problems"].append(f"{'row' if len(same) == 1 else 'rows'} {rows_text} {verb} the same month, "
                                        f"line and reference")
    duplicates: List[Dict[str, Any]] = []
    planned: List[Dict[str, Any]] = []
    for (period, line_item, reference_id), group in groups.items():
        first = group[0]
        if len(group) == 1:
            if not first["problems"]:
                planned.append({**first, "rows": [first["row"]]})
            continue
        numbers = [g["row"] for g in group]
        ref_text = f" with reference '{reference_id}'" if reference_id else " with no reference"
        what = f"{month_label(period)} {line_label(line_item)}{ref_text}"
        amounts = [g["amount"] for g in group]
        amounts_text = ", ".join(_money(a) for a in amounts)
        duplicates.append({"rows": numbers, "period": period, "line_item": line_item, "reference_id": reference_id,
                           "amounts": amounts, "combined": bool(combine_duplicates)})
        if combine_duplicates:
            notes = "; ".join(dict.fromkeys(g["notes"] for g in group if g["notes"])) or None
            total = sum(amounts)
            planned.append({**first, "amount": total, "notes": notes, "rows": numbers})
            warnings.append(f"Added up {_plural(len(group), 'row')} for {what} (rows {_list(numbers)}): "
                            f"{' + '.join(_money(a) for a in amounts)} = {_money(total)}")
        else:
            message = (f"rows {_list(numbers)} are {'both' if len(numbers) == 2 else 'all'} {what} "
                       f"(amounts {amounts_text}). Give them different "
                       f"references, remove the extra rows, or turn on 'add up duplicates'")
            for g in group:
                g["problems"].append(f"duplicate: {message}")

    for item in parsed:
        if item["problems"]:
            errors.append({"row": item["row"], "error": "; ".join(item["problems"])})

    # 5: compare with what is stored
    if not dry_run:
        conn.execute("BEGIN IMMEDIATE")
    try:
        periods = sorted({p["period"] for p in planned})
        stored: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
        stored_by_pair: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
        if periods:
            marks = ",".join("?" * len(periods))
            cursor.execute(f"SELECT actual_id, period, line_item, reference_id, amount, notes, source, batch_id "
                           f"FROM actuals WHERE period IN ({marks}) ORDER BY period, line_item, reference_id",
                           periods)
            columns = [d[0] for d in cursor.description]
            for record in cursor.fetchall():
                row = dict(zip(columns, record))
                stored[(row["period"], row["line_item"], row["reference_id"] or "")] = row
                stored_by_pair.setdefault((row["period"], row["line_item"]), []).append(row)
        closed = set(closed_periods(cursor, periods))
        closed_hits: Dict[str, int] = {}
        for item in planned:
            current = stored.get((item["period"], item["line_item"], item["reference_id"]))
            item["previous_amount"] = None if current is None else current["amount"]
            item["actual_id"] = None if current is None else current["actual_id"]
            if current is None:
                item["action"] = "insert"
                item["changes"] = ["amount"]
            else:
                changes = []
                if abs(float(current["amount"] or 0) - item["amount"]) > TOLERANCE:
                    changes.append("amount")
                if item["notes"] is not None and (item["notes"] or "") != (current["notes"] or ""):
                    changes.append("notes")
                item["action"] = "update" if changes else "unchanged"
                item["changes"] = changes
            if item["action"] != "unchanged" and item["period"] in closed:
                closed_hits[item["period"]] = closed_hits.get(item["period"], 0) + 1
                item["closed"] = True

        # Other references for the same month + line: they add up, or replace_lines removes them.
        line_totals, to_remove = _line_totals(planned, parsed, stored_by_pair, replace_lines)
        for row in to_remove:
            if row["period"] in closed:
                closed_hits[row["period"]] = closed_hits.get(row["period"], 0) + 1

        to_write = [i for i in planned if i["action"] in ("insert", "update")]
        counts = {
            "inserted": sum(1 for i in planned if i["action"] == "insert"),
            "updated": sum(1 for i in planned if i["action"] == "update"),
            "unchanged": sum(1 for i in planned if i["action"] == "unchanged"),
            "rejected": len(errors),
            "replaced": len(to_remove),
        }
        warnings.extend(_other_reference_warnings(line_totals, replace_lines))
        stored_before_total = round(sum(t["before"] for t in line_totals), 2)
        stored_after_total = round(sum(t["after"] for t in line_totals), 2)
        control_difference_after = None
        control_matches_after = None
        if control_total is not None:
            control_difference_after = round(stored_after_total - control_total, 2)
            control_matches_after = abs(stored_after_total - control_total) <= TOLERANCE
            if control_matches and not control_matches_after:
                why = []
                if any(t["other_references"] for t in line_totals) and not replace_lines:
                    why.append("amounts stored under other references are counted too")
                if errors:
                    why.append("rows with problems are left out")
                warnings.append(f"The file ties to the control total {_money(control_total)}, but after this import "
                                f"its months and lines hold {_money(stored_after_total)} in total (difference "
                                f"{_money(control_difference_after)})" + (f": {' and '.join(why)}" if why else "")
                                + ". See line_totals.")
        warnings.extend(_sign_warnings(planned))
        warnings.extend(_tie_out_warnings(planned))

        refused_reason = None
        status_code = 200
        if closed_hits:
            hits = "; ".join(month_label(p) + " (" + _plural(n, "row") + ")" for p, n in sorted(closed_hits.items()))
            refused_reason = (f"This would change closed months: {hits}. Reopen them (with a reason) or remove "
                              f"those rows. Rows for a closed month whose amounts already match are fine.")
            status_code = 409
        elif request_errors:
            refused_reason = " ".join(request_errors)
            status_code = 422
        elif errors and not allow_partial:
            first_errors = "; ".join(f"row {e['row']}: {e['error']}" for e in errors[:3])
            more = f" (and {_plural(len(errors) - 3, 'more problem')})" if len(errors) > 3 else ""
            verb = "has a problem" if len(errors) == 1 else "have problems"
            refused_reason = (f"{_plural(len(errors), 'row')} of {len(parsed)} {verb}: "
                              f"{_end_sentence(first_errors + more)} "
                              f"Fix the file and import again, or import only the valid rows.")
            status_code = 422
        elif not planned:
            refused_reason = (f"None of the {_plural(len(parsed), 'row')} can be imported. First problem: "
                              f"row {errors[0]['row']}: {_end_sentence(errors[0]['error'])}")
            status_code = 422

        if refused_reason:
            outcome = "rejected"
        elif not to_write and not to_remove:
            outcome = "unchanged"
        elif errors:
            outcome = "partial"
        else:
            outcome = "imported"

        batch_id = None
        written = False
        if outcome in ("imported", "partial") and not dry_run:
            batch_id = _new_batch_id(cursor)
            now = utc_now()
            notes_for_batch = [note.strip()] if note and note.strip() else []
            combined = [d for d in duplicates if d["combined"]]
            if combined:
                notes_for_batch.append(f"Added up duplicate rows for {_plural(len(combined), 'line')}")
            cursor.execute(
                """INSERT INTO actuals_batches (batch_id, created_at, source, filename, rows_in_file, inserted, updated,
                       unchanged, rejected, file_total, control_total, status, note, replaced, replace_lines,
                       entries_logged)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1)""",
                (batch_id, now, source, filename, len(parsed) + len(skipped), counts["inserted"], counts["updated"],
                 counts["unchanged"], counts["rejected"], totals["grand_total"], control_total,
                 "partial" if outcome == "partial" else "imported", " | ".join(notes_for_batch) or None,
                 counts["replaced"], 1 if replace_lines else 0))
            # replace_lines: rows under other references go, recorded so undo puts them back exactly.
            for row in to_remove:
                cursor.execute(_RECORD_CHANGE, (batch_id, "removed", row["actual_id"]))
                cursor.execute("DELETE FROM actuals WHERE actual_id = ?", (row["actual_id"],))
            # A row this import found already stored is its write too (undoing the import that stored it
            # first leaves it with this one): recorded, the row itself is not touched.
            for item in planned:
                if item["action"] == "unchanged":
                    cursor.execute(_RECORD_CHANGE, (batch_id, "confirmed", item["actual_id"]))
            sequence = 0
            for item in to_write:
                if item["action"] == "insert":
                    sequence += 1
                    actual_id = f"{batch_id}-{sequence:04d}"
                    item["actual_id"] = actual_id
                    cursor.execute(
                        """INSERT INTO actuals (actual_id, period, line_item, amount, reference_id, source, notes,
                               created_at, updated_at, batch_id)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (actual_id, item["period"], item["line_item"], item["amount"], item["reference_id"],
                         item["source"], item["notes"], now, now, batch_id))
                else:
                    # Keep what this import overwrites so undo can put it back.
                    cursor.execute(_RECORD_CHANGE, (batch_id, "updated", item["actual_id"]))
                    cursor.execute(
                        """UPDATE actuals SET amount = ?, notes = COALESCE(?, notes), source = COALESCE(?, source),
                               updated_at = ?, batch_id = ?
                           WHERE actual_id = ?""",
                        (item["amount"], item["notes"], item["source"], now, batch_id, item["actual_id"]))
            _log_request(cursor, batch_id, planned, parsed, replace_lines)
            conn.commit()
            written = True
        elif not dry_run:
            conn.rollback()
    except Exception:
        if not dry_run:
            conn.rollback()
        raise

    action_names = {"insert": "inserted", "update": "updated", "unchanged": "unchanged"}
    row_report: List[Dict[str, Any]] = []
    for item in planned:
        entry = {
            "row": item["row"], "period": item["period"], "line_item": item["line_item"],
            "label": line_label(item["line_item"]), "reference_id": item["reference_id"],
            "amount": round(item["amount"], 2), "previous_amount": item.get("previous_amount"),
            "action": action_names[item["action"]],
        }
        if len(item["rows"]) > 1:
            entry["rows"] = item["rows"]
        if item["action"] == "update":
            entry["changes"] = item["changes"]
        if item.get("closed"):
            entry["error"] = f"{month_label(item['period'])} is closed"
        row_report.append(entry)
    for item in parsed:
        if item["problems"]:
            row_report.append({
                "row": item["row"], "period": item["period"], "line_item": item["line_item"],
                "label": line_label(item["line_item"]) if item["line_item"] else item["line_text"],
                "reference_id": item["reference_id"], "amount": item["amount"], "previous_amount": None,
                "action": "rejected", "error": "; ".join(item["problems"]),
            })
    row_report.sort(key=lambda r: r["row"])

    closed_list = [{"period": p, "label": month_label(p), "rows": n} for p, n in sorted(closed_hits.items())]
    report: Dict[str, Any] = {
        "dry_run": dry_run,
        "written": written,
        "outcome": outcome,
        "batch_id": batch_id,
        "source": source,
        "filename": filename,
        "rows_in_file": len(parsed) + len(skipped),
        **counts,
        "imported": (counts["inserted"] + counts["updated"]) if written else 0,
        "errors": errors,
        "request_errors": request_errors,
        "duplicates": duplicates,
        "closed_periods": closed_list,
        "skipped_rows": skipped,
        "totals": totals,
        "control_total": control_total,
        "control_difference": control_difference,
        "control_matches": control_matches,
        "line_totals": line_totals,
        "stored_before_total": stored_before_total,
        "stored_after_total": stored_after_total,
        "control_difference_after": control_difference_after,
        "control_matches_after": control_matches_after,
        "warnings": warnings,
        "rows": row_report,
        "options": {"allow_partial": allow_partial, "combine_duplicates": combine_duplicates,
                    "replace_lines": replace_lines},
    }
    if file_info:
        report["file"] = {k: v for k, v in file_info.items() if v is not None}
    report["message"] = _import_message(report, refused_reason)
    if dry_run:
        return 200, report
    return (status_code if outcome == "rejected" else 200), report


def _list(numbers: Sequence[int]) -> str:
    numbers = list(numbers)
    if len(numbers) <= 2:
        return " and ".join(str(n) for n in numbers)
    return ", ".join(str(n) for n in numbers[:-1]) + f" and {numbers[-1]}"


def _import_message(report: Dict[str, Any], refused_reason: Optional[str]) -> str:
    counts = f"{report['inserted']} new, {report['updated']} changed, {report['unchanged']} unchanged"
    if report.get("replaced"):
        counts += f", {report['replaced']} replaced"
    if report["dry_run"]:
        if report["outcome"] == "rejected":
            return f"Check only, nothing written. The import would be refused: {refused_reason}"
        if report["outcome"] == "unchanged":
            return f"Check only, nothing written. Nothing would change: all {_plural(report['unchanged'], 'amount')} already match."
        rejected = f"; {_plural(report['rejected'], 'row')} would be left out" if report["rejected"] else ""
        return f"Check only, nothing written. Importing would give {counts}{rejected}."
    if report["outcome"] == "rejected":
        return f"Nothing imported. {refused_reason}"
    if report["outcome"] == "unchanged":
        return f"Nothing changed: all {_plural(report['unchanged'], 'amount')} already match what is stored."
    if report["outcome"] == "partial":
        return (f"Imported the valid rows: {counts}; {_plural(report['rejected'], 'row')} left out "
                f"(batch {report['batch_id']}).")
    return f"Imported actuals: {counts} (batch {report['batch_id']})."


# ---------------------------------------------------------------------------
# Deletes
# ---------------------------------------------------------------------------

def delete_actuals(conn, *, period: Optional[str] = None, line_item: Optional[str] = None,
                   actual_id: Optional[str] = None, source: Optional[str] = None) -> Dict[str, Any]:
    """
    Delete actuals by id, period and / or line item; refused (409) when a matching row is in a closed month.
    The deletion is a batch (status 'deleted', never undone): the rows it deleted are kept in
    actuals_batch_changes ('deleted') and its filter in actuals_batch_entries, so undoing an earlier
    import never brings a deleted row back.
    """
    clauses, params = [], []
    line_key_value = None
    if actual_id:
        clauses.append("actual_id = ?")
        params.append(actual_id)
    if period:
        clauses.append("period = ?")
        params.append(period)
    if line_item:
        line_key_value, error = resolve_line_item(line_item)
        if error:
            raise ActualsError(error[0].upper() + error[1:])
        clauses.append("line_item = ?")
        params.append(line_key_value)
    if not clauses:
        raise ActualsError("Provide actual_id, period and/or line_item")
    where = " AND ".join(clauses)
    conn.execute("BEGIN IMMEDIATE")
    try:
        cursor = conn.cursor()
        cursor.execute(f"SELECT period, COUNT(*), COALESCE(SUM(amount), 0) FROM actuals WHERE {where} GROUP BY period",
                       params)
        found = cursor.fetchall()
        closed = closed_periods(cursor, [r[0] for r in found])
        if closed:
            raise ActualsError(closed_message(closed, "deleted"), status=409, data={"closed_periods": closed})
        batch_id = _new_batch_id(cursor) if found else None
        if batch_id:
            cursor.execute(f"SELECT period, line_item, COALESCE(reference_id, '') FROM actuals WHERE {where}", params)
            keys = sorted(set(cursor.fetchall()))
            cursor.execute(
                f"""INSERT OR IGNORE INTO actuals_batch_changes (batch_id, actual_id, previous_amount, previous_notes,
                        previous_source, previous_batch_id, previous_updated_at, change_type, previous_period,
                        previous_line_item, previous_reference_id, previous_created_at)
                    SELECT ?, actual_id, amount, notes, source, batch_id, updated_at, 'deleted', period, line_item,
                           COALESCE(reference_id, ''), created_at
                    FROM actuals WHERE {where}""", (batch_id, *params))
            # The filter: a row id is its month, line and reference; otherwise every reference of the
            # month and / or line (NULL = any).
            filters = keys if actual_id else [(period, line_key_value, None)]
            cursor.executemany(_LOG_ENTRY, [(batch_id, seq, "delete", p, li, ref, None, None, None, None)
                                            for seq, (p, li, ref) in enumerate(filters)])
        cursor.execute(f"DELETE FROM actuals WHERE {where}", params)
        deleted = cursor.rowcount
        if deleted:
            what = []
            if period:
                what.append(month_label(period))
            if line_key_value:
                what.append(line_label(line_key_value) or line_item)
            if actual_id:
                what.append(f"row {actual_id}")
            cursor.execute(
                """INSERT INTO actuals_batches (batch_id, created_at, source, filename, rows_in_file, inserted, updated,
                       unchanged, rejected, file_total, control_total, status, note, entries_logged)
                   VALUES (?, ?, ?, NULL, 0, 0, 0, 0, 0, NULL, NULL, 'deleted', ?, 1)""",
                (batch_id, utc_now(), source or "delete",
                 f"Deleted {_plural(deleted, 'row')} ({', '.join(what)}; total {_money(sum(r[2] for r in found))})"))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {"deleted": deleted, "periods": sorted(r[0] for r in found), "batch_id": batch_id}


# ---------------------------------------------------------------------------
# Batches
# ---------------------------------------------------------------------------

BATCH_COLUMNS = ("batch_id", "created_at", "source", "filename", "rows_in_file", "inserted", "updated", "unchanged",
                 "rejected", "file_total", "control_total", "status", "note", "replaced")
CHANGE_COLUMNS = ("batch_id", "actual_id", "previous_amount", "previous_notes", "previous_source", "previous_batch_id",
                  "previous_updated_at", "change_type", "previous_period", "previous_line_item",
                  "previous_reference_id", "previous_created_at")
_ROW_COLUMNS = ("actual_id", "period", "line_item", "reference_id", "amount", "notes", "source", "batch_id",
                "created_at", "updated_at")
_ENTRY_COLUMNS = ("batch_id", "seq", "action", "period", "line_item", "reference_id", "amount", "notes", "source",
                  "actual_id")
#: Entry actions that write a row (or find it already stored, which counts as the batch's write too).
_WRITE_ACTIONS = ("insert", "update", "unchanged")


def _batch_rows(cursor, batch_id: str) -> List[Dict[str, Any]]:
    cursor.execute("SELECT actual_id, period, line_item, reference_id, amount FROM actuals WHERE batch_id = ? "
                   "ORDER BY period, line_item, reference_id", (batch_id,))
    columns = [d[0] for d in cursor.description]
    return [dict(zip(columns, r)) for r in cursor.fetchall()]


def _change_records(cursor, where: str, params: Sequence[Any]) -> List[Dict[str, Any]]:
    cursor.execute(f"SELECT {', '.join('c.' + c for c in CHANGE_COLUMNS)} FROM actuals_batch_changes c "
                   f"WHERE {where}", tuple(params))
    out = []
    for record in cursor.fetchall():
        change = dict(zip(CHANGE_COLUMNS, record))
        change["change_type"] = change["change_type"] or "updated"
        out.append(change)
    return out


def _batch_info(cursor) -> Dict[str, Dict[str, Any]]:
    """Every batch in the order it ran (rowid): position, status, time, replace_lines, entries logged or not."""
    cursor.execute("SELECT batch_id, status, created_at, COALESCE(replace_lines, 0), COALESCE(entries_logged, 0) "
                   "FROM actuals_batches ORDER BY rowid")
    return {r[0]: {"pos": i, "status": r[1], "created_at": r[2], "replace_lines": bool(r[3]), "logged": bool(r[4])}
            for i, r in enumerate(cursor.fetchall())}


def _resolve_earlier(cursor, actual_id: str, change: Dict[str, Any],
                     statuses: Dict[str, str]) -> Tuple[Optional[Dict[str, Any]], List[str]]:
    """
    The state a change record holds, walked back past undone imports (their own change records): an
    undone import's value never comes back. None: an undone import added the row, or there is no earlier
    value on record. Returns (state, undone imports skipped).
    """
    skipped: List[str] = []
    state = change
    while True:
        owner = state.get("previous_batch_id")
        if not owner or owner in skipped or statuses.get(owner) != "undone":
            return state, skipped
        skipped.append(owner)
        if str(actual_id).startswith(f"{owner}-"):
            return None, skipped
        earlier = _change_records(cursor, "c.batch_id = ? AND c.actual_id = ?", (owner, actual_id))
        if not earlier:
            return None, skipped
        state = earlier[0]


def _record_state(change: Dict[str, Any]) -> Dict[str, Any]:
    """The row a change record kept (its state before the batch changed, removed, found or deleted it)."""
    return {"actual_id": change["actual_id"], "amount": change["previous_amount"], "notes": change["previous_notes"],
            "source": change["previous_source"], "batch_id": change["previous_batch_id"],
            "created_at": change["previous_created_at"], "updated_at": change["previous_updated_at"]}


def _row_state(row: Dict[str, Any]) -> Dict[str, Any]:
    return {k: row.get(k) for k in ("actual_id", "amount", "notes", "source", "batch_id", "created_at", "updated_at")}


def _same_amount(a: Any, b: Any, tolerance: float = 1e-9) -> bool:
    return abs(float(a or 0) - float(b or 0)) <= tolerance


def _same_content(a: Dict[str, Any], b: Dict[str, Any]) -> bool:
    return _same_amount(a["amount"], b["amount"]) and (a["notes"] or "") == (b["notes"] or "")


def _same_state(a: Optional[Dict[str, Any]], b: Optional[Dict[str, Any]]) -> bool:
    """Same row as far as anyone can tell: presence, amount, notes, source and the batch that owns it."""
    if a is None or b is None:
        return a is None and b is None
    return (_same_content(a, b) and (a["source"] or "") == (b["source"] or "")
            and (a["batch_id"] or "") == (b["batch_id"] or ""))


def _undo_pairs(cursor, batch_id: str, rows: List[Dict[str, Any]],
                records: List[Dict[str, Any]]) -> List[Tuple[str, str]]:
    """Every month + line a batch acted on: what it asked for, the rows it owns, changed, removed or added."""
    prefix = f"{batch_id}-"
    pairs = {(r["period"], r["line_item"]) for r in rows}
    pairs |= {(c["previous_period"], c["previous_line_item"]) for c in records if c.get("previous_period")}
    cursor.execute("SELECT DISTINCT period, line_item FROM actuals_batch_entries WHERE batch_id = ? "
                   "AND period IS NOT NULL AND line_item IS NOT NULL", (batch_id,))
    pairs |= set(cursor.fetchall())
    cursor.execute("SELECT DISTINCT period, line_item FROM actuals WHERE substr(actual_id, 1, ?) = ?",
                   (len(prefix), prefix))
    pairs |= set(cursor.fetchall())
    cursor.execute("SELECT DISTINCT previous_period, previous_line_item FROM actuals_batch_changes "
                   "WHERE substr(actual_id, 1, ?) = ? AND previous_period IS NOT NULL", (len(prefix), prefix))
    pairs |= set(cursor.fetchall())
    return sorted(p for p in pairs if p[0] and p[1])


def _load_pairs(cursor, pairs: Sequence[Tuple[str, str]]) -> Dict[Tuple[str, str], Dict[str, Any]]:
    """For each month + line: the rows stored now, every change record and every logged entry (deletions included)."""
    models: Dict[Tuple[str, str], Dict[str, Any]] = {pair: {"rows": {}, "records": [], "entries": []}
                                                     for pair in pairs}
    by_period: Dict[str, List[Tuple[str, str]]] = {}
    for pair in models:
        by_period.setdefault(pair[0], []).append(pair)
    periods = sorted(by_period)
    for start in range(0, len(periods), 500):
        chunk = periods[start:start + 500]
        marks = ",".join("?" * len(chunk))
        cursor.execute(f"SELECT {', '.join(_ROW_COLUMNS)} FROM actuals WHERE period IN ({marks})", chunk)
        for record in cursor.fetchall():
            row = dict(zip(_ROW_COLUMNS, record))
            model = models.get((row["period"], row["line_item"]))
            if model is not None:
                model["rows"][row["reference_id"] or ""] = _row_state(row)
        for change in _change_records(cursor, f"c.previous_period IN ({marks})", chunk):
            model = models.get((change["previous_period"], change["previous_line_item"]))
            if model is not None:
                model["records"].append(change)
        cursor.execute(f"SELECT {', '.join(_ENTRY_COLUMNS)} FROM actuals_batch_entries WHERE period IN ({marks})",
                       chunk)
        for record in cursor.fetchall():
            entry = dict(zip(_ENTRY_COLUMNS, record))
            if entry["line_item"] is None:          # a deletion of a whole month
                for pair in by_period.get(entry["period"], []):
                    models[pair]["entries"].append(entry)
            elif (entry["period"], entry["line_item"]) in models:
                models[(entry["period"], entry["line_item"])]["entries"].append(entry)
    cursor.execute(f"SELECT {', '.join(_ENTRY_COLUMNS)} FROM actuals_batch_entries "
                   f"WHERE action = 'delete' AND period IS NULL")
    for record in cursor.fetchall():                # a deletion of a line in every month
        entry = dict(zip(_ENTRY_COLUMNS, record))
        for pair, model in models.items():
            if entry["line_item"] is None or entry["line_item"] == pair[1]:
                model["entries"].append(entry)
    return models


def _pair_program(model: Dict[str, Any], info: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    """
    One month + line as a program: ``ops`` = [(batch, op)] in the order the batches ran, and what
    ``_pair_base`` needs to tell each touched reference's state before the first batch touched it. An op
    holds the batch's writes there (reference, amount / notes / source as sent, row id, times, action),
    the references a replace_lines batch kept (None: it replaced nothing), references it removed one by
    one (a deletion; a batch from before the entries were logged) and ``clear`` (a deletion of the line).
    References no batch touched are not part of it and are never changed by undo.
    """
    def pos(b: Optional[str]) -> int:
        return info[b]["pos"] if b in info else -1

    ops: Dict[str, Dict[str, Any]] = {}
    inserted_at: Dict[str, int] = {}
    keeps: Dict[str, set] = {}

    def op_for(b: str) -> Dict[str, Any]:
        return ops.setdefault(b, {"writes": [], "keep": None, "remove": [], "clear": False})

    for entry in sorted(model["entries"], key=lambda e: (pos(e["batch_id"]), e["seq"])):
        b = entry["batch_id"]
        if b not in info or not info[b]["logged"]:
            continue
        action = entry["action"]
        if action in _WRITE_ACTIONS:
            ref = entry["reference_id"] or ""
            when = info[b]["created_at"]
            op_for(b)["writes"].append((ref, {"amount": entry["amount"], "notes": entry["notes"],
                                              "source": entry["source"]}, entry["actual_id"], when, when, action))
            if action == "insert":
                inserted_at.setdefault(ref, info[b]["pos"])
        elif action == "keep":
            keeps.setdefault(b, set()).add(entry["reference_id"] or "")
        elif action == "delete":
            if entry["reference_id"] is None:
                op_for(b)["clear"] = True
            else:
                op_for(b)["remove"].append(entry["reference_id"] or "")
    for b, op in ops.items():
        if info[b]["replace_lines"] and op["writes"]:
            op["keep"] = {w[0] for w in op["writes"]} | keeps.get(b, set())

    rows, records = model["rows"], model["records"]
    legacy = {b for b, i in info.items() if not i["logged"]}
    if legacy:
        _legacy_ops(rows, records, legacy, info, op_for, inserted_at)

    first_record: Dict[str, Tuple[int, Dict[str, Any]]] = {}
    for change in records:
        ref = change["previous_reference_id"] or ""
        at = pos(change["batch_id"])
        if ref not in first_record or at < first_record[ref][0]:
            first_record[ref] = (at, change)
    tracked = set(first_record) | set(inserted_at) | {w[0] for op in ops.values() for w in op["writes"]}
    return {"ops": sorted(ops.items(), key=lambda kv: pos(kv[0])), "tracked": tracked,
            "first_record": first_record, "inserted_at": inserted_at}


def _pair_base(cursor, program: Dict[str, Any], statuses: Dict[str, str]) -> Dict[str, Optional[Dict[str, Any]]]:
    """
    Each touched reference's state before the first batch touched it (None = no row): what the earliest
    change record there kept, or no row when a batch inserted it first. A state an undone batch wrote
    (one from before change tracking) walks back past it.
    """
    base: Dict[str, Optional[Dict[str, Any]]] = {}
    for ref in program["tracked"]:
        found = program["first_record"].get(ref)
        inserted = program["inserted_at"].get(ref)
        if found is None or (inserted is not None and inserted < found[0]):
            base[ref] = None
            continue
        state, _skipped = _resolve_earlier(cursor, found[1]["actual_id"], found[1], statuses)
        base[ref] = None if state is None else _record_state(state)
    return base


def _legacy_ops(rows: Dict[str, Dict[str, Any]], records: List[Dict[str, Any]], legacy: set,
                info: Dict[str, Dict[str, Any]], op_for, inserted_at: Dict[str, int]) -> None:
    """
    Batches imported before their entries were logged: their writes come from the rows they still own and
    from the change records later batches kept of them (a row id '<batch>-n' = a row the batch added), and
    the rows they replaced away from their own 'removed' records.
    """
    def owner_of_id(actual_id: Any) -> Optional[str]:
        text = str(actual_id or "")
        head = text.rsplit("-", 1)[0] if "-" in text else ""
        return head if head in legacy else None

    def written_by(b: str, ref: str) -> Optional[Dict[str, Any]]:
        row = rows.get(ref)
        if row is not None and row["batch_id"] == b:
            return row
        for change in records:
            if (change["previous_reference_id"] or "") == ref and change["previous_batch_id"] == b:
                return _record_state(change)
        return None

    writes: Dict[str, Dict[str, Tuple]] = {}
    ids: Dict[str, str] = {}
    for ref, row in rows.items():
        ids.setdefault(row["actual_id"], ref)
    for change in records:
        ids.setdefault(change["actual_id"], change["previous_reference_id"] or "")
    for actual_id, ref in ids.items():
        b = owner_of_id(actual_id)
        if b is None:
            continue
        inserted_at[ref] = min(inserted_at.get(ref, info[b]["pos"]), info[b]["pos"])
        state = written_by(b, ref)
        if state is not None:
            writes.setdefault(b, {})[ref] = (state, actual_id, "insert")
    own_updates = {(c["batch_id"], c["previous_reference_id"] or "") for c in records
                   if c["batch_id"] in legacy and c["change_type"] == "updated"}
    for b, ref in own_updates:
        state = written_by(b, ref)
        if state is not None:
            writes.setdefault(b, {}).setdefault(ref, (state, state["actual_id"], "update"))
    for ref, row in rows.items():
        b = row["batch_id"]
        if b in legacy and owner_of_id(row["actual_id"]) != b and (b, ref) not in own_updates:
            # Changed before change tracking: no earlier amount on record, so it counts as added.
            inserted_at[ref] = min(inserted_at.get(ref, info[b]["pos"]), info[b]["pos"])
            writes.setdefault(b, {}).setdefault(ref, (row, row["actual_id"], "insert"))
    for b, by_ref in writes.items():
        op = op_for(b)
        for ref in sorted(by_ref):
            state, actual_id, action = by_ref[ref]
            op["writes"].append((ref, {"amount": state["amount"], "notes": state["notes"], "source": state["source"]},
                                 actual_id, state["created_at"], state["updated_at"], action))
    for change in records:
        if change["batch_id"] in legacy and change["change_type"] == "removed":
            op_for(change["batch_id"])["remove"].append(change["previous_reference_id"] or "")


def _replay(base: Dict[str, Optional[Dict[str, Any]]], ordered: List[Tuple[str, Dict[str, Any]]],
            info: Dict[str, Dict[str, Any]], skip: Optional[str]) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, str]]:
    """
    Run a month + line's batches that are not undone (and not ``skip``) in order on top of ``base``, with
    the import's own rules: a new reference is inserted; a stored one is updated when the amount or the
    notes differ (notes / source left blank keep the stored ones) and otherwise left with the batch that
    stored it; replace_lines removes the references the batch did not name. Returns (state by reference,
    the batch that last removed each reference that is gone).
    """
    state = {ref: dict(s) for ref, s in base.items() if s is not None}
    removed_by: Dict[str, str] = {}
    for b, op in ordered:
        if b == skip or info.get(b, {}).get("status") == "undone":
            continue
        for ref, value, actual_id, created_at, updated_at, _action in op["writes"]:
            current = state.get(ref)
            if current is None:
                state[ref] = {"actual_id": actual_id, "amount": value["amount"], "notes": value["notes"],
                              "source": value["source"], "batch_id": b, "created_at": created_at,
                              "updated_at": updated_at}
                removed_by.pop(ref, None)
            elif (not _same_amount(current["amount"], value["amount"], TOLERANCE)
                  or (value["notes"] is not None and (value["notes"] or "") != (current["notes"] or ""))):
                current.update(amount=value["amount"], batch_id=b, updated_at=updated_at)
                if value["notes"] is not None:
                    current["notes"] = value["notes"]
                if value["source"] is not None:
                    current["source"] = value["source"]
        if op["clear"]:
            gone = list(state)
        else:
            gone = [ref for ref in state if op["keep"] is not None and ref not in op["keep"]]
            gone += [ref for ref in op["remove"] if ref in state and ref not in gone]
        for ref in gone:
            state.pop(ref, None)
            removed_by[ref] = b
    return state, removed_by


def list_batches(conn, limit: int = 50, offset: int = 0) -> Dict[str, Any]:
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM actuals_batches")
    total = cursor.fetchone()[0]
    cursor.execute(f"SELECT {', '.join(BATCH_COLUMNS)} FROM actuals_batches ORDER BY created_at DESC, rowid DESC "
                   f"LIMIT ? OFFSET ?", (limit, offset))
    batches = [dict(zip(BATCH_COLUMNS, r)) for r in cursor.fetchall()]
    closed = set(closed_periods(cursor))
    for batch in batches:
        batch["replaced"] = batch.get("replaced") or 0
        cursor.execute("SELECT period, COUNT(*), SUM(CASE WHEN actual_id LIKE ? THEN 1 ELSE 0 END) "
                       "FROM actuals WHERE batch_id = ? GROUP BY period ORDER BY period",
                       (f"{batch['batch_id']}-%", batch["batch_id"]))
        per_period = cursor.fetchall()
        batch["rows_current"] = sum(r[1] for r in per_period)
        batch["rows_added_current"] = sum(r[2] or 0 for r in per_period)
        replaced_periods = []
        if batch["replaced"] and batch["status"] not in ("undone", "deleted"):
            cursor.execute("SELECT DISTINCT previous_period FROM actuals_batch_changes WHERE batch_id = ? "
                           "AND change_type = 'removed'", (batch["batch_id"],))
            replaced_periods = [r[0] for r in cursor.fetchall() if r[0]]
        # The months the batch changed: the rows it owns and the months it replaced rows in (undo puts them back).
        batch["periods"] = sorted({r[0] for r in per_period} | set(replaced_periods))
        batch["closed_periods"] = [p for p in batch["periods"] if p in closed]
        if batch["status"] in ("undone", "deleted"):
            batch["can_undo"] = False
            batch["undo_blocked_reason"] = ("Already undone" if batch["status"] == "undone"
                                            else "A deletion is recorded for the audit trail; it cannot be undone")
        elif batch["closed_periods"]:
            batch["can_undo"] = False
            batch["undo_blocked_reason"] = closed_message(batch["closed_periods"], "removed")
        else:
            batch["can_undo"] = True
            batch["undo_blocked_reason"] = None
    return {"batches": batches, "total_count": total, "limit": limit, "offset": offset}


def _rows_word(n: int, singular: str, plural_: str) -> str:
    return singular if n == 1 else plural_


def _imports_text(ids: Sequence[str]) -> str:
    return f"{'import' if len(ids) == 1 else 'imports'} {', '.join(ids)}"


def undo_batch(conn, batch_id: str, dry_run: bool = False) -> Dict[str, Any]:
    """
    Undo an import, as if it had never run: every month + line it touched ends up exactly as replaying,
    in their original order, the batches still in place (not undone; deletions included) on top of the
    rows that predate them would leave it. So, for the common cases:

    * rows it added are deleted;
    * rows it changed go back to the value of the latest earlier import still in place (an import
      already undone is skipped, so its value never comes back); with none, they are removed;
    * rows it removed with replace_lines come back, unless a batch still in place wrote, replaced or
      deleted them since;
    * a row it stored that a later import still in place found unchanged stays, owned by that import
      (the later import would have stored it); a later correction is kept.

    Undoing a set of imports in any order therefore leaves the same rows. Each batch's request is kept in
    actuals_batch_entries (what it wrote, found unchanged, kept and deleted) and the rows it overwrote,
    removed, found or deleted in actuals_batch_changes; batches from before the request was logged are
    replayed from those change records. Only the months + lines the batch touched are recomputed, and
    only rows whose outcome depends on the batch change. Refused for closed months.
    """
    if not dry_run:
        conn.execute("BEGIN IMMEDIATE")
    try:
        cursor = conn.cursor()
        cursor.execute(f"SELECT {', '.join(BATCH_COLUMNS)} FROM actuals_batches WHERE batch_id = ?", (batch_id,))
        found = cursor.fetchone()
        if not found:
            raise ActualsError(f"Import batch '{batch_id}' not found", status=404)
        batch = dict(zip(BATCH_COLUMNS, found))
        batch["replaced"] = batch.get("replaced") or 0
        if batch["status"] == "undone":
            raise ActualsError(f"Batch {batch_id} was already undone", status=409)
        if batch["status"] == "deleted":
            raise ActualsError(f"Batch {batch_id} records a deletion; a deletion cannot be undone", status=409)
        info = _batch_info(cursor)
        statuses = {b: i["status"] for b, i in info.items()}
        prefix = f"{batch_id}-"
        rows = _batch_rows(cursor, batch_id)
        added = [r for r in rows if str(r["actual_id"]).startswith(prefix)]
        changed = [r for r in rows if not str(r["actual_id"]).startswith(prefix)]
        records = _change_records(cursor, "c.batch_id = ?", (batch_id,))
        updates = {c["actual_id"]: c for c in records if c["change_type"] == "updated"}
        removals = [c for c in records if c["change_type"] == "removed"]
        untracked = [] if info[batch_id]["logged"] else [r for r in changed if r["actual_id"] not in updates]

        # Replay every month + line the batch touched with and without it; change only what differs.
        pairs = _undo_pairs(cursor, batch_id, rows, records)
        models = _load_pairs(cursor, pairs)
        undone_with = {**statuses}
        undone_without = {**statuses, batch_id: "undone"}
        here = info[batch_id]["pos"]
        changes: List[Dict[str, Any]] = []
        targets: Dict[Tuple[str, str, str], Optional[Dict[str, Any]]] = {}
        past_undone: List[Tuple[List[str], bool]] = []
        blocked = 0
        removed_again: List[str] = []
        dropped: List[str] = []
        changed_by_pair: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
        for r in changed:
            changed_by_pair.setdefault((r["period"], r["line_item"]), []).append(r)
        removals_by_pair: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
        for change in removals:
            removals_by_pair.setdefault((change["previous_period"], change["previous_line_item"]), []).append(change)
        for period, line_item in pairs:
            model = models[(period, line_item)]
            program = _pair_program(model, info)
            ordered = program["ops"]
            with_state, _ = _replay(_pair_base(cursor, program, undone_with), ordered, info, None)
            without_state, removed_by = _replay(_pair_base(cursor, program, undone_without), ordered, info, batch_id)
            pair_changes: Dict[str, Dict[str, Any]] = {}
            for ref in sorted(set(with_state) | set(without_state)):
                target = without_state.get(ref)
                current = model["rows"].get(ref)
                if _same_state(with_state.get(ref), target) or _same_state(current, target):
                    continue
                kind = ("put_back" if current is None else "remove" if target is None
                        else "restore" if not _same_content(current, target) else "reassign")
                pair_changes[ref] = {"kind": kind, "period": period, "line_item": line_item, "reference_id": ref,
                                     "current": current, "target": target}
            changes.extend(pair_changes.values())
            # What the notes say: rows that go back past undone imports, replaced rows not put back.
            for r in changed_by_pair.get((period, line_item), []):
                ref = r["reference_id"] or ""
                target = without_state.get(ref)
                targets[(period, line_item, ref)] = target
                after = info[target["batch_id"]]["pos"] if target and target.get("batch_id") in info else -1
                skipped = sorted({b for b, op in ordered if statuses.get(b) == "undone" and b != batch_id
                                  and after < info[b]["pos"] < here
                                  and any(w[0] == ref and w[5] in ("insert", "update") for w in op["writes"])},
                                 key=lambda b: info[b]["pos"])
                if skipped:
                    past_undone.append((skipped, target is None))
            for change in removals_by_pair.get((period, line_item), []):
                ref = change["previous_reference_id"] or ""
                if ref in pair_changes:          # put back, or described by the change it gets
                    continue
                if ref in without_state:
                    blocked += 1
                elif removed_by.get(ref) and removed_by[ref] != batch_id:
                    removed_again.append(removed_by[ref])
                else:
                    dropped.append(change["previous_batch_id"] or "")

        removes = [c for c in changes if c["kind"] == "remove"]
        restores = [c for c in changes if c["kind"] == "restore"]
        put_back = [c for c in changes if c["kind"] == "put_back"]
        reassigns = [c for c in changes if c["kind"] == "reassign"]

        # Rows that left this import since (changed by a later entry and kept, or deleted).
        cursor.execute("""SELECT COUNT(*) FROM actuals WHERE COALESCE(batch_id, '') <> ?
                            AND (substr(actual_id, 1, ?) = ? OR actual_id IN (
                                 SELECT actual_id FROM actuals_batch_changes
                                 WHERE batch_id = ? AND COALESCE(change_type, 'updated') = 'updated'))""",
                       (batch_id, len(prefix), prefix, batch_id))
        changed_since = cursor.fetchone()[0]
        gone = max((batch["inserted"] or 0) + (batch["updated"] or 0) - len(rows), 0)
        changed_since = min(changed_since, gone)
        deleted_since = gone - changed_since
        legacy = len(updates) + len(untracked) < (batch["updated"] or 0)

        periods = sorted({r["period"] for r in rows} | {c["period"] for c in changes})
        # A closed month blocks the undo when a row there would change amount, notes or presence, or when
        # the batch owns a row there; handing a row to another import with the same amount does not.
        blocking = {r["period"] for r in rows} | {c["period"] for c in changes if c["kind"] != "reassign"}
        closed = closed_periods(cursor, sorted(blocking))

        def describe(c: Dict[str, Any], state: Dict[str, Any]) -> Dict[str, Any]:
            return {"period": c["period"], "line_item": c["line_item"], "label": line_label(c["line_item"]),
                    "reference_id": c["reference_id"], "amount": state["amount"], "batch_id": state["batch_id"]}

        preview = {
            "batch_id": batch_id,
            "batch": batch,
            "rows_to_remove": len(removes),
            "rows_to_restore": len(restores),
            "rows_to_put_back": len(put_back),
            "rows_to_reassign": len(reassigns),
            "added_rows": len(added),
            "changed_rows": len(changed),
            "rows_changed_since": changed_since,
            "rows_deleted_since": deleted_since,
            "periods": periods,
            "closed_periods": closed,
            "amount_removed": round(sum(float(c["current"]["amount"] or 0) for c in removes), 2),
            "changed": [{"period": r["period"], "line_item": r["line_item"], "label": line_label(r["line_item"]),
                         "reference_id": r["reference_id"], "amount": r["amount"],
                         "restores_to": (targets.get((r["period"], r["line_item"], r["reference_id"] or ""))
                                         or {}).get("amount"),
                         "action": ("restore" if targets.get((r["period"], r["line_item"], r["reference_id"] or ""))
                                    is not None else "remove")}
                        for r in changed[:50]],
            "put_back_rows": [describe(c, c["target"]) for c in put_back[:50]],
            "reassigned_rows": [{**describe(c, c["target"]), "from_batch_id": c["current"]["batch_id"]}
                                for c in reassigns[:50]],
            "dry_run": dry_run,
        }
        notes = []
        if restores:
            notes.append(f"{_plural(len(restores), 'row')} existed before this import and "
                         f"{_rows_word(len(restores), 'goes', 'go')} back to the earlier amount.")
        back_to = [s for s, gone_too in past_undone if not gone_too]
        if back_to:
            ids = sorted({b for skipped in back_to for b in skipped})
            notes.append(f"{_plural(len(back_to), 'row')} {_rows_word(len(back_to), 'goes', 'go')} back past "
                         f"{_imports_text(ids)} (already undone) to the amount before "
                         f"{_rows_word(len(ids), 'it', 'them')}.")
        gone_with = [s for s, gone_too in past_undone if gone_too]
        if gone_with:
            ids = sorted({b for skipped in gone_with for b in skipped})
            notes.append(f"{_plural(len(gone_with), 'row')} {_rows_word(len(gone_with), 'is', 'are')} removed: no "
                         f"import that is still in place wrote {_rows_word(len(gone_with), 'it', 'them')} before "
                         f"(already undone: {', '.join(ids)}).")
        if untracked:
            notes.append(f"{_plural(len(untracked), 'row')} changed by this import "
                         f"{_rows_word(len(untracked), 'has', 'have')} no earlier amount on record (imported before "
                         f"change tracking) and {_rows_word(len(untracked), 'is', 'are')} removed.")
        if put_back:
            notes.append(f"{_plural(len(put_back), 'row')} replaced under other references "
                         f"{_rows_word(len(put_back), 'is', 'are')} put back.")
        if reassigns:
            owners = sorted({c["target"]["batch_id"] for c in reassigns if c["target"]["batch_id"]})
            whose = (f"{_imports_text(owners)}, still in place, {_rows_word(len(owners), 'has', 'have')} the same "
                     f"amount and {_rows_word(len(owners), 'keeps', 'keep')} {_rows_word(len(reassigns), 'it', 'them')}"
                     if owners else f"{_rows_word(len(reassigns), 'it was', 'they were')} stored before any import")
            notes.append(f"{_plural(len(reassigns), 'row')} {_rows_word(len(reassigns), 'keeps its', 'keep their')} "
                         f"amount: {whose}.")
        if blocked:
            notes.append(f"{_plural(blocked, 'row')} this import replaced {_rows_word(blocked, 'is', 'are')} not put "
                         f"back: the same month, line and reference has been written again since.")
        if removed_again:
            ids = sorted(set(removed_again), key=lambda b: info.get(b, {}).get("pos", -1))
            notes.append(f"{_plural(len(removed_again), 'row')} this import replaced "
                         f"{_rows_word(len(removed_again), 'is', 'are')} not put back: {_imports_text(ids)}, still in "
                         f"place, {_rows_word(len(ids), 'removes', 'remove')} "
                         f"{_rows_word(len(removed_again), 'it', 'them')} as well.")
        if dropped:
            ids = sorted({b for b in dropped if b})
            notes.append(f"{_plural(len(dropped), 'row')} this import replaced {_rows_word(len(dropped), 'is', 'are')} "
                         f"not put back: the import that wrote {_rows_word(len(dropped), 'it', 'them')}"
                         + (f" ({', '.join(ids)})" if ids else "") + " was undone.")
        if legacy and gone:
            notes.append(f"{_plural(gone, 'row')} from this import {_rows_word(gone, 'has', 'have')} been changed or "
                         f"deleted since by later entries; undo leaves {_rows_word(gone, 'it', 'them')} as "
                         f"{_rows_word(gone, 'it is', 'they are')}.")
        else:
            if changed_since:
                notes.append(f"{_plural(changed_since, 'row')} from this import "
                             f"{_rows_word(changed_since, 'has', 'have')} been changed since by later entries and "
                             f"{_rows_word(changed_since, 'is', 'are')} kept.")
            if deleted_since:
                notes.append(f"{_plural(deleted_since, 'row')} from this import "
                             f"{_rows_word(deleted_since, 'has', 'have')} been deleted since (or replaced by a later "
                             f"import); nothing to undo for {_rows_word(deleted_since, 'it', 'them')}.")
        preview["notes"] = notes
        if closed:
            if dry_run:
                preview["can_undo"] = False
                preview["message"] = closed_message(closed, "removed")
                return preview
            raise ActualsError(closed_message(closed, "removed"), status=409, data=preview)
        preview["can_undo"] = True
        extra_text = ((f" and put back {_plural(len(put_back), 'replaced row')}" if put_back else "")
                      + (f"; {_plural(len(reassigns), 'row')} {_rows_word(len(reassigns), 'stays', 'stay')} with "
                         f"the import still in place that has the same amount" if reassigns else ""))
        if dry_run:
            preview["message"] = (f"Undo would remove {_plural(len(removes), 'row')}, restore "
                                  f"{_plural(len(restores), 'row')} to the earlier amounts{extra_text}"
                                  + (f" in {_months_text(periods)}" if periods else "") + ".")
            return preview
        for c in removes:
            cursor.execute("DELETE FROM actuals WHERE actual_id = ?", (c["current"]["actual_id"],))
        for c in restores + reassigns:
            t = c["target"]
            cursor.execute(
                """UPDATE actuals SET amount = ?, notes = ?, source = ?, batch_id = ?,
                       created_at = COALESCE(?, created_at), updated_at = COALESCE(?, updated_at)
                   WHERE actual_id = ?""",
                (t["amount"], t["notes"], t["source"], t["batch_id"], t["created_at"], t["updated_at"],
                 c["current"]["actual_id"]))
        for c in put_back:
            t = c["target"]
            actual_id = t["actual_id"]
            cursor.execute("SELECT 1 FROM actuals WHERE actual_id = ?", (actual_id,))
            if not actual_id or cursor.fetchone():
                actual_id = f"PUT-{uuid.uuid4().hex[:12]}"
            cursor.execute(
                """INSERT INTO actuals (actual_id, period, line_item, amount, reference_id, source, notes,
                       created_at, updated_at, batch_id)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (actual_id, c["period"], c["line_item"], t["amount"], c["reference_id"], t["source"], t["notes"],
                 t["created_at"] or t["updated_at"] or utc_now(), t["updated_at"] or t["created_at"] or utc_now(),
                 t["batch_id"]))
        removed, restored = len(removes), len(restores)
        stamp = utc_now()
        undo_note = (f"Undone {stamp[:10]} {stamp[11:16]} UTC: removed {_plural(removed, 'row')}, "
                     f"restored {_plural(restored, 'row')}"
                     + (f", put back {_plural(len(put_back), 'replaced row')}" if put_back else "")
                     + (f", kept {_plural(len(reassigns), 'row')} with the import still in place" if reassigns else ""))
        cursor.execute("UPDATE actuals_batches SET status = 'undone', note = CASE WHEN note IS NULL OR note = '' "
                       "THEN ? ELSE note || ' | ' || ? END WHERE batch_id = ?", (undo_note, undo_note, batch_id))
        conn.commit()
    except Exception:
        if not dry_run:
            conn.rollback()
        raise
    preview["removed"] = removed
    preview["restored"] = restored
    preview["put_back"] = len(put_back)
    preview["reassigned"] = len(reassigns)
    preview["message"] = (f"Undid import {batch_id}: removed {_plural(removed, 'row')}, restored "
                          f"{_plural(restored, 'row')} to the earlier amounts{extra_text}"
                          + (f" in {_months_text(periods)}" if periods else "") + ".")
    return preview


# ---------------------------------------------------------------------------
# Month-end close
# ---------------------------------------------------------------------------

PERIOD_COLUMNS = ("period", "status", "closed_at", "closed_by", "note", "updated_at")
#: The close / reopen trail kept in actuals_periods.note (newest first) is cut to this many characters.
NOTE_TRAIL_MAX = 2000


def _note_trail(newest: Optional[str], trail: Optional[str]) -> Optional[str]:
    """'Reopened ... by bob: reason | Aug close v1': the newest entry first, earlier ones kept after it."""
    text = " | ".join(part.strip() for part in (newest, trail) if part and part.strip())
    if len(text) > NOTE_TRAIL_MAX:
        text = text[:NOTE_TRAIL_MAX - 3].rstrip() + "..."
    return text or None


def list_periods(conn, start_period: Optional[str] = None, end_period: Optional[str] = None) -> Dict[str, Any]:
    """Every month with actuals or a close record (or every month of the window), with its close status."""
    cursor = conn.cursor()
    cursor.execute(f"SELECT {', '.join(PERIOD_COLUMNS)} FROM actuals_periods")
    status = {r[0]: dict(zip(PERIOD_COLUMNS, r)) for r in cursor.fetchall()}
    cursor.execute("SELECT period, COUNT(*), MAX(updated_at) FROM actuals GROUP BY period")
    stats = {r[0]: {"rows": r[1], "last_updated": r[2]} for r in cursor.fetchall()}
    cursor.execute("SELECT DISTINCT period, line_item FROM actuals")
    lines_by_period: Dict[str, set] = {}
    for period, line in cursor.fetchall():
        lines_by_period.setdefault(period, set()).add(line)
    if start_period and end_period:
        months = period_range(start_period, end_period)
    else:
        months = sorted(set(status) | set(stats))
    out = []
    for p in months:
        s = status.get(p) or {}
        st = stats.get(p) or {"rows": 0, "last_updated": None}
        # The sales deductions are either the total or the detail lines (as budget vs actual counts them):
        # a month given the total alone needs no detail lines, and detail lines alone give the total.
        lines = lines_by_period.get(p, set()) & set(VARIANCE_LINE_ITEMS)
        details = lines & set(SALES_DEDUCTION_DETAIL_LINES)
        required = set(VARIANCE_LINE_ITEMS)
        if DEDUCTIONS_TOTAL_LINE in lines and not details:
            required -= set(SALES_DEDUCTION_DETAIL_LINES)
        if details and DEDUCTIONS_TOTAL_LINE not in lines:
            lines = lines | {DEDUCTIONS_TOTAL_LINE}
        st = {**st, "lines_entered": len(lines & required)}
        lines_total = len(required)
        out.append({
            "period": p,
            "label": month_label(p),
            "status": s.get("status") or "open",
            "closed": (s.get("status") == "closed"),
            "closed_at": s.get("closed_at") if s.get("status") == "closed" else None,
            "closed_by": s.get("closed_by") if s.get("status") == "closed" else None,
            "note": s.get("note"),
            "updated_at": s.get("updated_at"),
            "rows": st["rows"],
            "lines_entered": st["lines_entered"],
            "lines_total": lines_total,
            "last_updated": st["last_updated"],
        })
    return {"periods": out, "closed_periods": [p["period"] for p in out if p["closed"]]}


def close_period(conn, period: str, closed_by: Optional[str] = None, note: Optional[str] = None) -> Dict[str, Any]:
    cursor = conn.cursor()
    cursor.execute(f"SELECT {', '.join(PERIOD_COLUMNS)} FROM actuals_periods WHERE period = ?", (period,))
    found = cursor.fetchone()
    current = dict(zip(PERIOD_COLUMNS, found)) if found else None
    if current and current["status"] == "closed":
        who = f" by {current['closed_by']}" if current.get("closed_by") else ""
        return {"period": period, "status": "closed", "changed": False, "closed_at": current["closed_at"],
                "closed_by": current["closed_by"], "note": current["note"],
                "message": f"{month_label(period)} was already closed{who}."}
    now = utc_now()
    closed_by = (closed_by or "").strip() or None
    note = (note or "").strip() or None
    # A month closed again keeps its trail (the reopen reason and earlier close notes), newest first.
    stored_note = _note_trail(note, current["note"] if current else None)
    cursor.execute(
        """INSERT INTO actuals_periods (period, status, closed_at, closed_by, note, updated_at)
           VALUES (?, 'closed', ?, ?, ?, ?)
           ON CONFLICT(period) DO UPDATE SET status = 'closed', closed_at = excluded.closed_at,
               closed_by = excluded.closed_by, note = excluded.note, updated_at = excluded.updated_at""",
        (period, now, closed_by, stored_note, now))
    conn.commit()
    cursor.execute("SELECT COUNT(*) FROM actuals WHERE period = ?", (period,))
    rows = cursor.fetchone()[0]
    return {"period": period, "status": "closed", "changed": True, "closed_at": now, "closed_by": closed_by,
            "note": stored_note, "rows": rows,
            "message": f"Closed {month_label(period)}" + (f" ({_plural(rows, 'actual')} locked)" if rows else
                                                          " (no actuals entered yet)") + "."}


def reopen_period(conn, period: str, reason: Optional[str], reopened_by: Optional[str] = None) -> Dict[str, Any]:
    reason = (reason or "").strip()
    if not reason:
        raise ActualsError("A reason is required to reopen a closed month")
    cursor = conn.cursor()
    cursor.execute(f"SELECT {', '.join(PERIOD_COLUMNS)} FROM actuals_periods WHERE period = ?", (period,))
    found = cursor.fetchone()
    current = dict(zip(PERIOD_COLUMNS, found)) if found else None
    if not current or current["status"] != "closed":
        return {"period": period, "status": "open", "changed": False,
                "message": f"{month_label(period)} is not closed; nothing to reopen."}
    now = utc_now()
    who = (reopened_by or "").strip()
    entry = f"Reopened {now[:10]}{' by ' + who if who else ''}: {reason}"
    if current.get("closed_by") or current.get("closed_at"):
        entry += f" (was closed {str(current.get('closed_at') or '')[:10]}{' by ' + current['closed_by'] if current.get('closed_by') else ''})"
    note = _note_trail(entry, current.get("note"))
    cursor.execute("UPDATE actuals_periods SET status = 'open', note = ?, updated_at = ? WHERE period = ?",
                   (note, now, period))
    conn.commit()
    return {"period": period, "status": "open", "changed": True, "reason": reason, "reopened_by": who or None,
            "note": note, "message": f"Reopened {month_label(period)}: {reason}"}
