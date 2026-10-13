"""
Excel workbook export of the forecast.

One .xlsx that carries, and explains, the revenue -> COGS forecast for a
scenario and a reporting window:

    Read Me            scenario, window, how every number is computed
                       (deduction formulas, rule precedence, COGS, labor),
                       headline totals and every warning
    Gross Margin       gross sales -> deductions -> net sales -> COGS ->
                       gross profit, by month
    Income Statement   the full monthly P&L
    By Product / By Customer / By Customer Type
                       window totals plus monthly units and net sales
    Labor Forecast     direct labor hours, cost and FTEs by labor type
    Machine Hours      hours required vs capacity, utilization, cost
    Unit Costs         per-unit BOM / routing cost of each product sold
    Sales Detail       every forecast sales line, deductions applied
    Deduction Rules    discount / allowance / excise / commission rules
    Price List         customer and customer-type prices
    Cash Flow          monthly cash flow (totals carry balances)
    Balance Sheet      month-end balances from the opening balance sheet

Nothing is computed here: every number comes from
services.financial_statements (statements) and services.operations_forecast
(COGS / labor forecasts and sales detail), so the workbook always agrees
with the API and the screens. This module only lays the numbers out with
human labels and Excel formatting. Cells hold numbers, never formatted
strings, so the sheets can be summed, pivoted and charted.
"""
from __future__ import annotations

import io
import math
import os
import re
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services.financial_statements import AVG_DAYS_PER_MONTH, build_statements, load_assumptions, period_range
from services.operations_forecast import (
    DEFAULT_HOURS_PER_FTE,
    MACHINE_TOTALS_NOTE,
    build_cogs_forecast,
    build_labor_forecast,
    detail_totals,
    sales_detail_report,
)
from services.pricing import (
    DEDUCTION_EXAMPLE,
    DEDUCTION_METHODS,
    DEDUCTION_TYPES,
    PRICE_METHODS,
    deduction_specificity,
    load_deduction_rules,
    load_price_rules,
    price_source,
    price_specificity,
    rule_price,
)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SHEET_NAMES: List[str] = [
    "Read Me",
    "Gross Margin",
    "Income Statement",
    "By Product",
    "By Customer",
    "By Customer Type",
    "Labor Forecast",
    "Machine Hours",
    "Unit Costs",
    "Sales Detail",
    "Deduction Rules",
    "Price List",
    "Cash Flow",
    "Balance Sheet",
]

SHEET_DESCRIPTIONS: Dict[str, str] = {
    "Gross Margin": "Monthly bridge from gross sales through each deduction to net sales, product COGS, the "
                    "period costs (labor variance, indirect labor, overhead, other COGS) and gross profit. Ties to "
                    "the Income Statement.",
    "Income Statement": "Full monthly profit and loss, the same numbers as the Reporting screens.",
    "By Product": "Units, gross sales, deductions, net sales, product COGS and contribution margin per product, "
                  "with units and net sales by month.",
    "By Customer": "The same breakdown per customer.",
    "By Customer Type": "The same breakdown per customer type (channel).",
    "Labor Forecast": "Direct labor hours, cost and FTEs by labor type by month (routing labor minutes x units), "
                      "the routing standard against the pay of employees classed Direct (paid hours, utilization, "
                      "variance), plus labor hours by product.",
    "Machine Hours": "Machine hours required against available capacity, utilization and machine cost by month. "
                     "The machine cost Total row ties to the Income Statement's Machine / manufacturing line.",
    "Unit Costs": "Per-unit material, labor and machine cost of each product sold, from its BOM and routing, "
                  "with data-quality warnings.",
    "Sales Detail": "Every forecast sales line with its deductions, net sales, product COGS and the deduction rules "
                    "applied, each amount rounded to the cent. Use the filter buttons on the header row.",
    "Deduction Rules": "The discount, allowance, excise tax and commission rules visible to the selected scenario(s).",
    "Price List": "Customer, customer-type and product prices used to price sales rows imported without a price.",
    "Cash Flow": "Monthly cash flow. The Total column carries the window's opening and closing cash (balances, "
                 "not sums of monthly balances).",
    "Balance Sheet": "Month-end balance sheet starting from the opening balances, including excise tax payable.",
}

# ---------------------------------------------------------------------------
# Statement line labels (shared with the CSV export in reporting_routes)
# ---------------------------------------------------------------------------

#: The COGS lines that are period costs: in Total COGS and gross profit, not in the per-product figures.
PERIOD_COSTS = "the direct labor variance, indirect labor, factory overhead and other COGS"
#: Gross profit as the income statement and the COGS forecast compute it.
GROSS_PROFIT_FORMULA = ("Gross profit = net sales - total COGS, where total COGS = product COGS + direct labor "
                        "variance (payroll vs standard) + indirect labor (factory) + factory overhead + other COGS.")

INCOME_STATEMENT_LABELS: Dict[str, str] = {
    "gross_sales": "Gross sales",
    "terms_allowance": "Terms allowance",
    "chargebacks": "Chargebacks allowance",
    "other_allowances": "Other discounts & allowances",
    "excise_tax": "Excise tax (FET)",
    "total_sales_deductions": "Total sales deductions",
    "revenue": "Net sales",
    "units_sold": "Units sold",
    "cogs_materials": "Materials",
    "cogs_labor": "Direct labor (standard)",
    "labor_variance": "Direct labor variance (payroll vs standard)",
    "cogs_manufacturing": "Machine / manufacturing",
    "indirect_labor": "Indirect labor (factory)",
    "factory_overhead": "Factory overhead",
    "other_cogs": "Other COGS",
    "total_cogs": "Total COGS",
    "gross_profit": "Gross profit",
    "gross_margin_pct": "Gross margin %",
    "commissions": "Sales commissions",
    "payroll": "Payroll",
    "admin_expense": "Admin expenses",
    "operating_expense": "Operating expenses",
    "depreciation": "Depreciation",
    "total_operating_expenses": "Total operating expenses",
    "operating_income": "Operating income",
    "operating_margin_pct": "Operating margin %",
    "interest_expense": "Interest expense",
    "pre_tax_income": "Pre-tax income",
    "income_tax": "Income tax",
    "net_income": "Net income",
    "net_margin_pct": "Net margin %",
}

CASH_FLOW_LABELS: Dict[str, str] = {
    "net_income": "Net income",
    "depreciation": "Depreciation",
    "change_in_accounts_receivable": "(Increase) / decrease in accounts receivable",
    "change_in_inventory": "(Increase) / decrease in inventory",
    "change_in_accounts_payable": "Increase / (decrease) in accounts payable",
    "change_in_excise_tax_payable": "Increase / (decrease) in excise tax payable",
    "operating_cash_flow": "Cash from operations",
    "capex": "Capital expenditures",
    "investing_cash_flow": "Cash from investing",
    "loan_proceeds": "Loan proceeds",
    "principal_payments": "Loan principal payments",
    "financing_cash_flow": "Cash from financing",
    "net_cash_flow": "Net change in cash",
    "opening_cash": "Opening cash",
    "closing_cash": "Closing cash",
    "interest_paid": "Memo: interest paid (in net income)",
    "taxes_paid": "Memo: income tax (in net income)",
}

BALANCE_SHEET_LABELS: Dict[str, str] = {
    "cash": "Cash",
    "accounts_receivable": "Accounts receivable",
    "inventory": "Inventory",
    "total_current_assets": "Total current assets",
    "fixed_assets_gross": "Fixed assets at cost",
    "accumulated_depreciation": "Accumulated depreciation",
    "fixed_assets_net": "Net fixed assets",
    "total_assets": "Total assets",
    "accounts_payable": "Accounts payable",
    "excise_tax_payable": "Excise tax payable",
    "other_liabilities": "Other liabilities",
    "total_current_liabilities": "Total current liabilities",
    "debt": "Debt",
    "total_liabilities": "Total liabilities",
    "opening_equity": "Opening equity",
    "retained_earnings": "Retained earnings since opening",
    "total_equity": "Total equity",
    "total_liabilities_and_equity": "Total liabilities & equity",
    "balance_check": "Balance check (assets - liabilities - equity)",
    "working_capital": "Working capital",
}

STATEMENT_LABELS: Dict[str, Dict[str, str]] = {
    "income_statement": INCOME_STATEMENT_LABELS,
    "cash_flow": CASH_FLOW_LABELS,
    "balance_sheet": BALANCE_SHEET_LABELS,
}

PERCENT_LINES = frozenset({"gross_margin_pct", "operating_margin_pct", "net_margin_pct",
                           "contribution_margin_pct", "deduction_pct_of_gross"})

_BOLD_LINES: Dict[str, frozenset] = {
    "income_statement": frozenset({"gross_sales", "total_sales_deductions", "revenue", "total_cogs", "gross_profit",
                                   "total_operating_expenses", "operating_income", "pre_tax_income", "net_income"}),
    "cash_flow": frozenset({"operating_cash_flow", "investing_cash_flow", "financing_cash_flow", "net_cash_flow",
                            "closing_cash"}),
    "balance_sheet": frozenset({"total_current_assets", "total_assets", "total_current_liabilities",
                                "total_liabilities", "total_equity", "total_liabilities_and_equity"}),
}
_INDENT_LINES: Dict[str, frozenset] = {
    "income_statement": frozenset({"terms_allowance", "chargebacks", "other_allowances", "excise_tax",
                                   "cogs_materials", "cogs_labor", "labor_variance", "cogs_manufacturing",
                                   "indirect_labor", "factory_overhead", "other_cogs", "commissions", "payroll", "admin_expense", "operating_expense",
                                   "depreciation", "gross_margin_pct", "operating_margin_pct", "net_margin_pct"}),
    "cash_flow": frozenset({"depreciation", "change_in_accounts_receivable", "change_in_inventory",
                            "change_in_accounts_payable", "change_in_excise_tax_payable", "capex", "loan_proceeds",
                            "principal_payments"}),
    "balance_sheet": frozenset({"cash", "accounts_receivable", "inventory", "fixed_assets_gross",
                                "accumulated_depreciation", "accounts_payable", "excise_tax_payable",
                                "other_liabilities", "opening_equity", "retained_earnings"}),
}


def humanize(key: str) -> str:
    """'change_in_inventory' -> 'Change in inventory' (fallback label for new lines)."""
    text = str(key).replace("_", " ").strip()
    return text[:1].upper() + text[1:]


def line_label(section: str, key: str) -> str:
    """Human label for a statement line key; unknown keys are humanized."""
    return STATEMENT_LABELS.get(section, {}).get(key) or humanize(key)


def statement_lines(statements: Dict[str, Any], section: str) -> List[Tuple[str, str]]:
    """
    (key, label) for every line of a statement section, in the engine's
    order. Lines come from the engine's rows, so a line added to
    build_statements shows up in the exports automatically.
    """
    rows = statements.get(section, {}).get("by_period") or []
    if not rows:
        return []
    return [(k, line_label(section, k)) for k in rows[0].keys()
            if k != "period" and _is_number(rows[0][k])]


# ---------------------------------------------------------------------------
# Sales detail columns and totals (shared with GET /reporting/sales-detail)
# ---------------------------------------------------------------------------

#: (key, header, format kind) for the Sales Detail sheet, in column order.
SALES_DETAIL_COLUMNS: List[Tuple[str, str, Optional[str]]] = [
    ("period", "Period", None),
    ("forecast_id", "Scenario", None),
    ("sale_id", "Sale ID", None),
    ("customer_id", "Customer ID", None),
    ("customer_name", "Customer", None),
    ("customer_type", "Customer type", None),
    ("unit_id", "Product ID", None),
    ("unit_name", "Product", None),
    ("unit_type", "Product type", None),
    ("quantity", "Quantity", "units"),
    ("unit_price", "Unit price", "price"),
    ("gross_sales", "Gross sales", "money"),
    ("terms_allowance", "Terms allowance", "money"),
    ("chargebacks", "Chargebacks allowance", "money"),
    ("other_allowances", "Other discounts & allowances", "money"),
    ("excise_tax", "Excise tax (FET)", "money"),
    ("total_sales_deductions", "Total sales deductions", "money"),
    ("net_sales", "Net sales", "money"),
    ("commissions", "Sales commissions", "money"),
    ("materials", "Materials", "money"),
    ("labor", "Direct labor", "money"),
    ("machine", "Machine", "money"),
    ("product_cogs", "Product COGS", "money"),
    ("contribution_margin", "Contribution margin", "money"),
    ("labor_hours", "Labor hours", "hours"),
    ("machine_hours", "Machine hours", "hours"),
    ("rules_applied", "Deduction rules applied", None),
]

def sales_detail_totals(rows: Iterable[Dict[str, Any]]) -> Dict[str, float]:
    """
    Column totals over sales detail rows (every numeric column except the
    unit price). Totals are sums of the rows as exported (each money value a
    cent amount), so they equal a SUM of the column in Excel. The
    statements add up the unrounded lines instead; the difference is the
    ``rounding`` that operations_forecast.sales_detail_report returns.
    """
    return detail_totals(rows)


def sales_detail_columns(rows: Sequence[Dict[str, Any]]) -> List[Tuple[str, str, Optional[str]]]:
    """SALES_DETAIL_COLUMNS plus any extra keys the rows carry (humanized)."""
    known = {c[0] for c in SALES_DETAIL_COLUMNS}
    extra: List[Tuple[str, str, Optional[str]]] = []
    for row in rows[:1]:
        for key, value in row.items():
            if key not in known:
                extra.append((key, humanize(key), "money" if _is_number(value) else None))
    return SALES_DETAIL_COLUMNS + extra


# ---------------------------------------------------------------------------
# Labels for rule / cost metadata
# ---------------------------------------------------------------------------

COST_WARNING_TEXT: Dict[str, str] = {
    "missing_bom": "No BOM (or the BOM version has no lines): material cost is $0",
    "missing_router": "No routing (or the routing has no operations): labor and machine cost are $0",
    "router_missing_machine": "A routing step uses a machine that does not exist: its machine cost is $0",
    "router_labor_without_type": "A routing step has labor minutes but no labor type: its labor cost is $0",
    "router_missing_labor_rate": "A routing step uses a labor type with no rate: its labor cost is $0",
}

PRICE_SOURCE_LABELS: Dict[str, str] = {
    "customer_product": "Customer + product",
    "customer_all": "Customer, all products",
    "type_product": "Customer type + product",
    "type_all": "Customer type, all products",
    "product": "Product, all customers",
    "all": "All products and customers",
}

TREATMENT_LABELS: Dict[str, str] = {
    "contra_revenue": "Sales deduction (gross to net sales)",
    "selling_expense": "Selling expense (below gross profit)",
}

# ---------------------------------------------------------------------------
# Formatting
# ---------------------------------------------------------------------------

MONEY = "#,##0.00"
PRICE = "#,##0.00##"   # cents, plus up to two more decimals when the price has them (23.994)
MONEY_WHOLE = "#,##0"
PERCENT = "0.0%"
RATE_PERCENT = "0.00%"
UNITS_WHOLE = "#,##0"
UNITS_FRACTION = "#,##0.##"
HOURS = "#,##0.00"
FTE = "0.00"
MINUTES = "#,##0.00"

_FORMATS = {
    "money": MONEY,
    "price": PRICE,
    "money0": MONEY_WHOLE,
    "pct": PERCENT,
    "rate": RATE_PERCENT,
    "hours": HOURS,
    "fte": FTE,
    "minutes": MINUTES,
    "int": UNITS_WHOLE,
}

_NAVY = "1F3864"
_FONT = "Calibri"  # the workbook's default font; named so bold / colored cells match it
_HEADER_FONT = Font(name=_FONT, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor=_NAVY)
_TOTAL_FILL = PatternFill("solid", fgColor="E7ECF5")
_OVER_FILL = PatternFill("solid", fgColor="F8CBAD")
_TITLE_FONT = Font(name=_FONT, bold=True, size=16, color=_NAVY)
_SECTION_FONT = Font(name=_FONT, bold=True, size=12, color=_NAVY)
_BOLD = Font(name=_FONT, bold=True)
_NOTE_FONT = Font(name=_FONT, italic=True, color="595959")
_TOP_BORDER = Border(top=Side(style="thin", color="7F7F7F"))
_WRAP = Alignment(wrap_text=True, vertical="top")
_MAX_WIDTH = 50
_MIN_WIDTH = 8
_SIZE_SAMPLE_ROWS = 2000


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _clean(value: Any) -> Any:
    """Cell-safe value: numbers stay numbers, NaN/inf become blank, others text."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        # NaN / inf cannot be stored; "+ 0.0" turns a rounded -0.0 into 0.0.
        return value + 0.0 if math.isfinite(value) else None
    if _is_number(value):
        return value
    if isinstance(value, (list, tuple, set)):
        value = "; ".join(str(v) for v in value)
    # Control characters (from imported names / notes) are not allowed in xlsx.
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _set(ws, row: int, col: int, value: Any):
    """Write a value; text is always stored as text, never as a formula."""
    cell = ws.cell(row=row, column=col, value=_clean(value))
    if isinstance(cell.value, str) and cell.value.startswith("="):
        # openpyxl treats "=..." as a formula; an imported name or note must
        # not become one (formula injection, or a workbook Excel must repair).
        cell.data_type = "s"
    return cell


def _pct(value: Any) -> Optional[float]:
    """Engine percentages are 0-100; Excel percentages are fractions."""
    return round(float(value) / 100.0, 6) if _is_number(value) else None


class _Writer:
    """Small helper around a worksheet: styled cells, sections and sizing."""

    def __init__(self, ws):
        self.ws = ws
        self.row = 1
        self._unsized: set = set()

    def cell(self, row: int, col: int, value: Any, kind: Optional[str] = None, *, bold: bool = False,
             indent: int = 0, fill: Optional[PatternFill] = None, border: Optional[Border] = None,
             align_left: bool = False):
        cell = _set(self.ws, row, col, value)
        if _is_number(cell.value):
            if kind == "units":
                cell.number_format = UNITS_WHOLE if float(cell.value).is_integer() else UNITS_FRACTION
            elif kind in _FORMATS:
                cell.number_format = _FORMATS[kind]
            elif kind is None:
                cell.number_format = MONEY
        if bold:
            cell.font = _BOLD
        if indent or align_left:
            cell.alignment = Alignment(indent=indent, horizontal="left")
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
        return cell

    def header(self, headers: Sequence[str], row: Optional[int] = None, label_cols: int = 1) -> int:
        """Styled header row; the first `label_cols` (text) columns are left-aligned, the rest centered."""
        row = self.row if row is None else row
        for i, text in enumerate(headers, start=1):
            cell = _set(self.ws, row, i, text)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="left" if i <= label_cols else "center", vertical="center")
        self.row = row + 1
        return row

    def section(self, text: str, font: Font = _SECTION_FONT) -> None:
        cell = _set(self.ws, self.row, 1, text)
        cell.font = font
        self._unsized.add(cell.coordinate)
        self.row += 1

    def note(self, text: str, col: int = 1) -> None:
        cell = _set(self.ws, self.row, col, text)
        cell.font = _NOTE_FONT
        self._unsized.add(cell.coordinate)
        self.row += 1

    def blank(self, n: int = 1) -> None:
        self.row += n

    def style_row(self, row: int, ncols: int, *, bold: bool = True, fill: Optional[PatternFill] = _TOTAL_FILL,
                  border: Optional[Border] = _TOP_BORDER) -> None:
        for col in range(1, ncols + 1):
            cell = self.ws.cell(row=row, column=col)
            if bold:
                cell.font = _BOLD
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border

    def finish(self, freeze: Optional[str] = None, widths: Optional[Dict[str, float]] = None) -> None:
        if freeze:
            self.ws.freeze_panes = freeze
            # Repeat the header row on every printed page.
            self.ws.print_title_rows = "1:1"
        self._autosize()
        for col, width in (widths or {}).items():
            self.ws.column_dimensions[col].width = width
        # Print landscape; narrow sheets fit one page wide (wide ones would
        # shrink to unreadable, so they print across pages instead).
        self.ws.page_setup.orientation = "landscape"
        if self.ws.max_column <= 16:
            self.ws.sheet_properties.pageSetUpPr.fitToPage = True
            self.ws.page_setup.fitToWidth = 1
            self.ws.page_setup.fitToHeight = 0

    def _autosize(self) -> None:
        widths: Dict[str, int] = {}
        max_row = self.ws.max_row
        rows = list(self.ws.iter_rows(min_row=1, max_row=min(max_row, _SIZE_SAMPLE_ROWS)))
        if max_row > _SIZE_SAMPLE_ROWS:
            rows += list(self.ws.iter_rows(min_row=max_row, max_row=max_row))
        for row in rows:
            for cell in row:
                if cell.value is None or cell.coordinate in self._unsized:
                    continue
                length = _display_length(cell.value, cell.number_format)
                if cell.alignment is not None and cell.alignment.indent:
                    length += int(cell.alignment.indent) * 2
                letter = cell.column_letter
                widths[letter] = max(widths.get(letter, 0), length)
        for letter, length in widths.items():
            self.ws.column_dimensions[letter].width = min(max(length + 2, _MIN_WIDTH), _MAX_WIDTH)


def _display_length(value: Any, number_format: str) -> int:
    if _is_number(value):
        if "%" in number_format:
            return len(f"{value * 100:,.2f}%")
        if number_format in (MONEY_WHOLE, UNITS_WHOLE):
            return len(f"{value:,.0f}")
        if number_format == PRICE:
            text = f"{value:,.4f}"
            return len(text[:-2] + text[-2:].rstrip("0"))
        return len(f"{value:,.2f}")
    return max((len(part) for part in str(value).split("\n")), default=0)


# ---------------------------------------------------------------------------
# Data collection
# ---------------------------------------------------------------------------

def _lookup(cursor, sql: str) -> Dict[Any, Any]:
    cursor.execute(sql)
    return {row[0]: row[1:] for row in cursor.fetchall()}


def collect_workbook_data(cursor, forecast_ids: Sequence[str], start_period: str, end_period: str,
                          hours_per_fte: Optional[float] = None) -> Dict[str, Any]:
    """Everything the workbook shows, built by the engine and forecast services."""
    ids = [str(f).strip() for f in forecast_ids if f and str(f).strip()]
    if not ids:
        raise ValueError("At least one forecast_id is required")
    periods = period_range(start_period, end_period)
    hours_per_fte = float(hours_per_fte) if hours_per_fte is not None else DEFAULT_HOURS_PER_FTE
    if not math.isfinite(hours_per_fte) or hours_per_fte <= 0:
        raise ValueError("hours_per_fte must be greater than zero")

    statements = build_statements(cursor, ids, start_period, end_period)
    cogs = build_cogs_forecast(cursor, ids, start_period, end_period)
    labor = build_labor_forecast(cursor, ids, start_period, end_period, hours_per_fte=hours_per_fte)
    detail = sales_detail_report(cursor, ids, start_period, end_period)
    # The statements' assumptions block is rounded to 2 decimals for display
    # (a 25.5% tax rate reads 0.26); the Read Me states the stored values.
    assumptions, _found = load_assumptions(cursor, ids[0])

    scenario_names = {fid: vals[0] for fid, vals in _lookup(cursor, "SELECT forecast_id, name FROM forecast").items()}
    customers = {cid: vals[0] for cid, vals in
                 _lookup(cursor, "SELECT customer_id, customer_name FROM customers").items()}
    units = _lookup(cursor, "SELECT unit_id, unit_name, unit_type, base_price FROM units")

    return {
        "forecast_ids": ids,
        "periods": periods,
        "start_period": periods[0],
        "end_period": periods[-1],
        "generated_at": datetime.now(),
        "hours_per_fte": hours_per_fte,
        "statements": statements,
        "cogs": cogs,
        "labor": labor,
        "sales_detail": detail["rows"],
        "sales_detail_totals": detail["totals"],
        "sales_detail_report_totals": detail["report_totals"],
        "sales_detail_rounding": detail["rounding"],
        "assumptions": assumptions,
        "deduction_rules": load_deduction_rules(cursor, ids, include_inactive=True),
        "price_rules": load_price_rules(cursor, ids, include_inactive=True),
        "scenario_names": scenario_names,
        "customers": customers,
        "units": units,
    }


def _scenario_label(data: Dict[str, Any], forecast_id: Optional[str]) -> str:
    if not forecast_id:
        return "All scenarios (company-wide)"
    name = data["scenario_names"].get(forecast_id)
    return f"{name} ({forecast_id})" if name else f"{forecast_id} (not found)"


def _collect_warnings(data: Dict[str, Any]) -> List[Tuple[str, str, str]]:
    """(level, code, message) for every distinct warning in the reports."""
    out: List[Tuple[str, str, str]] = []
    seen = set()
    for fid in data["forecast_ids"]:
        if fid not in data["scenario_names"]:
            out.append(("warning", "scenario_not_found", f"Scenario '{fid}' does not exist; it contributes nothing."))
    sources = (data["statements"].get("warnings") or [], data["cogs"].get("warnings") or [],
               data["labor"].get("warnings") or [])
    for warnings in sources:
        for w in warnings:
            code = str(w.get("code") or "warning")
            key = (code, w.get("machine_id"))
            if key in seen:
                continue
            seen.add(key)
            out.append((str(w.get("level") or "warning"), code, str(w.get("message") or "")))
            products = w.get("products")
            if code == "cost_warnings" and isinstance(products, dict):
                for unit_id, codes in sorted(products.items(), key=lambda kv: str(kv[0])):
                    name = (data["units"].get(unit_id) or (None,))[0]
                    who = f"{unit_id} ({name})" if name else str(unit_id)
                    text = "; ".join(COST_WARNING_TEXT.get(c, c) for c in (codes or []))
                    out.append(("warning", "product_cost", f"Product {who}: {text}."))
    return out


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------

def _sheet_read_me(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    stmt = data["statements"]
    cogs = data["cogs"]
    labor = data["labor"]
    income = stmt["income_statement"]["totals"]
    summary = cogs["summary"]["totals"]
    periods = data["periods"]

    w.section("Forecast workbook", font=_TITLE_FONT)
    w.blank()

    def pair(label: str, value: Any, kind: Optional[str] = None) -> None:
        w.cell(w.row, 1, label, bold=True)
        w.cell(w.row, 2, value, kind, align_left=True)
        w.row += 1

    names = [_scenario_label(data, fid) for fid in data["forecast_ids"]]
    pair("Scenario" if len(names) == 1 else "Scenarios (combined)", "; ".join(names))
    pair("Reporting window", f"{data['start_period']} to {data['end_period']} ({len(periods)} month"
                             f"{'' if len(periods) == 1 else 's'})")
    pair("Generated", data["generated_at"].strftime("%Y-%m-%d %H:%M"))
    pair("Sales lines in window", cogs.get("line_count", len(data["sales_detail"])), "int")
    pair("Active deduction rules", cogs.get("deduction_rule_count", 0), "int")
    carry = stmt.get("carry_forward") or {}
    if carry.get("months_rolled_forward"):
        pair("Opening balances", f"Stated at the end of {carry.get('opening_period')}; rolled forward "
                                 f"{carry['months_rolled_forward']} month(s) to {carry.get('balances_as_of')}.")
    w.blank()

    w.section("Headline totals for the window")
    pair("Gross sales", income.get("gross_sales"), "money")
    pair("Total sales deductions", income.get("total_sales_deductions"), "money")
    pair("Net sales", income.get("revenue"), "money")
    pair("Product COGS (BOM + routing)", summary.get("product_cogs"), "money")
    pair("Total COGS", income.get("total_cogs"), "money")
    pair("Gross profit", income.get("gross_profit"), "money")
    pair("Gross margin % (on net sales)", _pct(income.get("gross_margin_pct")), "pct")
    pair("Sales commissions", income.get("commissions"), "money")
    pair("Net income", income.get("net_income"), "money")
    pair("Units sold", income.get("units_sold"), "units")
    pair("Direct labor hours", labor["totals"].get("hours"), "hours")
    pair("Average direct labor FTEs", labor["totals"].get("avg_fte"), "fte")
    w.blank()

    w.section("Sheets in this workbook")
    for name in SHEET_NAMES[1:]:
        pair(name, SHEET_DESCRIPTIONS.get(name, ""))
    w.blank()

    w.section("How the numbers are computed")
    hpf = data["hours_per_fte"]
    days_per_month = f"{AVG_DAYS_PER_MONTH:g}"   # the engine's working-capital divisor, 365.25 / 12 = 30.4375
    assumptions = data.get("assumptions") or stmt.get("assumptions") or {}
    explanations: List[Tuple[str, str]] = [
        ("Gross sales", "Quantity x unit price for each forecast sales line; the line's stored total is used only "
                        "when the line has no price."),
    ]
    for dtype, meta in DEDUCTION_TYPES.items():
        treatment = ("It is a sales deduction: it reduces gross sales to net sales."
                     if meta["treatment"] == "contra_revenue"
                     else "It is a selling expense below gross profit; it does not reduce net sales.")
        explanations.append((meta["label"], f"{meta['description']} {treatment}"))
    for method, formula in DEDUCTION_METHODS.items():
        explanations.append((f"Method: {method}", formula))
    explanations += [
        ("Worked example", DEDUCTION_EXAMPLE),
        ("Which rule applies", "Each sales line gets at most one rule per deduction type: the most specific active "
                               "rule whose scope matches the line and whose start / end periods cover the month. "
                               "Specificity = 8 x customer + 4 x customer type + 2 x product + 1 x product type. "
                               "Ties go to the scenario-specific rule, then the later start period. Different "
                               "deduction types stack. A rule with a 0 rate is an exemption (for example a "
                               "government customer exempt from excise tax)."),
        ("Net sales", "Gross sales - terms allowance - chargebacks allowance - other discounts & allowances - "
                      "excise tax (FET). Sales commissions are not deducted from net sales."),
        ("Product COGS", "Per-unit cost from each product's BOM and routing, times units sold. Materials = sum of "
                         "BOM line quantity x unit price; direct labor (standard) = routing labor minutes / 60 x "
                         "labor rate; machine = routing machine minutes / 60 x machine rate."),
        ("Period costs in COGS", "The direct labor variance (payroll of employees classed Direct - the routing "
                                 "standard) and indirect labor (pay of employees classed Indirect) come from "
                                 "Payroll; factory overhead and other COGS from operating expense allocations "
                                 "(categories typed factory overhead or COGS). They are in Total COGS and gross "
                                 "profit on the Gross Margin and Income Statement sheets, but not in the "
                                 "per-product, per-customer or per-customer-type figures."),
        ("Contribution margin vs gross profit", f"Contribution margin = net sales - product COGS (before "
                                                f"{PERIOD_COSTS}). {GROSS_PROFIT_FORMULA} Margin percentages are "
                                                f"on net sales."),
        ("Labor classes", "Each employee in Payroll has a labor class that decides where their pay goes. Direct "
                          "(builds product): in COGS. Direct labor (standard) stays routing labor minutes x labor "
                          "rates x units sold, and in every month with direct payroll a Direct labor variance line "
                          "adds direct payroll - standard, so total direct labor equals the pay of employees classed "
                          "Direct and nobody is counted twice. Indirect (factory support: supervisors, maintenance, "
                          "material handling): in COGS as Indirect labor (factory). Admin, or not classified: "
                          "Payroll in operating expenses. In a month with no direct payroll (nobody classed Direct, "
                          "or before a direct hire starts) direct labor is the routing standard alone. Inventory is "
                          "always valued at the standard product cost."),
        ("Labor hours and FTEs", f"Labor hours = routing labor minutes per unit x units sold / 60, by labor type. "
                                 f"FTEs = hours / {hpf:,.2f} paid hours per FTE per month (40 h x 52 weeks / 12 = "
                                 f"173.33 by default). Labor cost equals Direct labor (standard) on the Income "
                                 f"Statement. Paid hours = weekly hours x 52 / 12 of employees classed Direct; "
                                 f"utilization = standard hours / paid hours."),
        ("Machine utilization", "Machine hours required / available hours per month (the machine's available "
                                "minutes per month / 60). Above 100% means more machine time than capacity."),
        ("Payroll vs direct labor", str((labor.get("payroll_note") or {}).get("message") or "")),
        ("Working capital", f"Accounts receivable = (net sales + excise tax) x DSO / {days_per_month}; inventory = "
                            f"product COGS x DIO / {days_per_month}; accounts payable = purchases x DPO / "
                            f"{days_per_month}, where purchases = materials + factory overhead + other COGS + admin "
                            f"+ operating expenses + sales commissions; excise tax payable = excise tax x excise "
                            f"payable days / {days_per_month}. {days_per_month} is the average days in a month "
                            f"(365.25 / 12). This scenario uses DSO "
                            f"{_num_text(assumptions.get('dso_days'), 4)}, "
                            f"DIO {_num_text(assumptions.get('dio_days'), 4)}, "
                            f"DPO {_num_text(assumptions.get('dpo_days'), 4)}, excise payable days "
                            f"{_num_text(assumptions.get('excise_payable_days'), 4)} and an income tax rate of "
                            f"{_rate_text(assumptions.get('income_tax_rate'))}."),
        ("Rounding", "Sales Detail rounds every money amount of a line to the cent (unit prices keep up to 4 "
                     "decimals), so each line cross-foots and its Total row is the SUM of the lines. The statements "
                     "add up the unrounded lines, so that SUM can differ from the Income Statement by a few cents; "
                     "the Rounding row under the Sales Detail total shows the difference. On Machine Hours the "
                     "Total rows are every machine together, month by month as the statements add them (the "
                     "machine cost Total is the Income Statement's Machine / manufacturing line); each machine's "
                     "row is rounded on its own, so the machine rows can add up to a cent or two more or less "
                     "than the Total."),
    ]
    for topic, text in explanations:
        w.cell(w.row, 1, topic, bold=True)
        w.cell(w.row, 2, text).alignment = _WRAP
        w.ws.cell(row=w.row, column=1).alignment = _WRAP
        w.row += 1
    w.blank()

    w.section("Warnings")
    warnings = _collect_warnings(data)
    if not warnings:
        w.cell(w.row, 1, "None")
        w.cell(w.row, 2, "No warnings for this window.")
        w.row += 1
    for level, code, message in warnings:
        w.cell(w.row, 1, "Info" if level == "info" else "Warning", bold=True).alignment = _WRAP
        w.cell(w.row, 2, f"{message} [{code}]").alignment = _WRAP
        w.row += 1

    w.finish(widths={"A": 34, "B": 110})


def _num_text(value: Any, decimals: int = 2) -> str:
    if not _is_number(value):
        return "n/a"
    return f"{value:,.{decimals}f}".rstrip("0").rstrip(".")


def _rate_text(rate: Any) -> str:
    """A stored fraction as a percent with the precision it has: 0.255 -> '25.5%', 0.21 -> '21%'."""
    if not _is_number(rate):
        return "n/a"
    return f"{_num_text(float(rate) * 100, 4)}%"


_GROSS_MARGIN_LINES: List[Tuple[str, str, str, bool, int]] = [
    # key, label, kind, bold, indent
    ("units", "Units sold", "units", False, 0),
    ("gross_sales", "Gross sales", "money0", True, 0),
    ("terms_allowance", "Terms allowance", "money0", False, 1),
    ("chargebacks", "Chargebacks allowance", "money0", False, 1),
    ("other_allowances", "Other discounts & allowances", "money0", False, 1),
    ("excise_tax", "Excise tax (FET)", "money0", False, 1),
    ("total_sales_deductions", "Total sales deductions", "money0", True, 0),
    ("net_sales", "Net sales", "money0", True, 0),
    ("materials", "Materials", "money0", False, 1),
    ("labor", "Direct labor (standard)", "money0", False, 1),
    ("machine", "Machine", "money0", False, 1),
    ("product_cogs", "Product COGS (BOM + routing)", "money0", True, 0),
    ("contribution_margin", "Contribution margin (before period costs)", "money0", True, 0),
    ("contribution_margin_pct", "Contribution margin %", "pct", False, 1),
    ("labor_variance", "Direct labor variance (payroll vs standard)", "money0", False, 1),
    ("indirect_labor", "Indirect labor (factory)", "money0", False, 1),
    ("factory_overhead", "Factory overhead", "money0", False, 1),
    ("other_cogs", "Other COGS", "money0", False, 1),
    ("total_cogs", "Total COGS", "money0", True, 0),
    ("gross_profit", "Gross profit", "money0", True, 0),
    ("gross_margin_pct", "Gross margin %", "pct", False, 1),
    ("commissions", "Memo: sales commissions (operating expense)", "money0", False, 0),
]


def _sheet_gross_margin(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    periods = data["periods"]
    summary = data["cogs"]["summary"]
    by_period = {r["period"]: r for r in summary["by_period"]}
    totals = summary["totals"]
    w.header(["Gross margin bridge"] + periods + ["Total"])
    for key, label, kind, bold, indent in _GROSS_MARGIN_LINES:
        r = w.row
        w.cell(r, 1, label, bold=bold, indent=indent)
        values = [by_period.get(p, {}).get(key) for p in periods] + [totals.get(key)]
        for j, value in enumerate(values, start=2):
            w.cell(r, j, _pct(value) if kind == "pct" else value, kind, bold=bold)
        if key in ("net_sales", "gross_profit"):
            w.style_row(r, len(periods) + 2)
        w.row += 1
    w.blank()
    w.note(f"Contribution margin = net sales - product COGS. {GROSS_PROFIT_FORMULA} Commissions are an operating "
           f"expense below gross profit.")
    w.finish(freeze="B2")


def _sheet_statement(ws, data: Dict[str, Any], section: str) -> None:
    w = _Writer(ws)
    stmt = data["statements"]
    periods = data["periods"]
    block = stmt[section]
    rows = block["by_period"]
    is_balance = section == "balance_sheet"
    title = {"income_statement": "Income statement", "cash_flow": "Cash flow",
             "balance_sheet": "Balance sheet"}[section]
    if is_balance:
        opening = block.get("opening") or {}
        headers = [title, f"Opening ({opening.get('period', '')})"] + periods
    else:
        headers = [title] + periods + ["Total"]
    w.header(headers)
    totals = block.get("totals") or {}
    bold_lines = _BOLD_LINES.get(section, frozenset())
    indent_lines = _INDENT_LINES.get(section, frozenset())
    for key, label in statement_lines(stmt, section):
        r = w.row
        bold = key in bold_lines
        w.cell(r, 1, label, bold=bold, indent=1 if key in indent_lines else 0)
        if is_balance:
            values = [opening.get(key)] + [row.get(key) for row in rows]
        else:
            values = [row.get(key) for row in rows] + [totals.get(key)]
        if key in PERCENT_LINES:
            kind, values = "pct", [_pct(v) for v in values]
        elif key == "units_sold":
            kind = "units"
        else:
            kind = "money0"
        for j, value in enumerate(values, start=2):
            w.cell(r, j, value, kind, bold=bold)
        if key in ("revenue", "gross_profit", "net_income", "closing_cash", "total_assets",
                   "total_liabilities_and_equity"):
            w.style_row(r, len(headers))
        w.row += 1
    if section == "cash_flow":
        w.blank()
        w.note("Total column: flows are summed; opening and closing cash are the window's balances.")
    w.finish(freeze="B2")


_DIMENSION_METRICS: List[Tuple[str, str, str]] = [
    ("units", "Units", "units"),
    ("gross_sales", "Gross sales", "money"),
    ("terms_allowance", "Terms allowance", "money"),
    ("chargebacks", "Chargebacks allowance", "money"),
    ("other_allowances", "Other discounts & allowances", "money"),
    ("excise_tax", "Excise tax (FET)", "money"),
    ("total_sales_deductions", "Total sales deductions", "money"),
    ("net_sales", "Net sales", "money"),
    ("materials", "Materials", "money"),
    ("labor", "Direct labor", "money"),
    ("machine", "Machine", "money"),
    ("product_cogs", "Product COGS", "money"),
    ("contribution_margin", "Contribution margin", "money"),
    ("contribution_margin_pct", "Contribution margin %", "pct"),
    ("avg_gross_price", "Avg gross price / unit", "money"),
    ("avg_net_price", "Avg net price / unit", "money"),
    ("cost_per_unit", "Product COGS / unit", "money"),
    ("commissions", "Sales commissions (opex)", "money"),
]


def _sheet_dimension(ws, data: Dict[str, Any], groups: List[Dict[str, Any]],
                     id_columns: List[Tuple[str, str]], what: str) -> None:
    w = _Writer(ws)
    periods = data["periods"]
    summary = data["cogs"]["summary"]
    headers = ([h for _k, h in id_columns] + [h for _k, h, _f in _DIMENSION_METRICS]
               + [f"Units {p}" for p in periods] + [f"Net sales {p}" for p in periods])
    n_id = len(id_columns)
    w.header(headers, label_cols=n_id)

    def write(row: int, labels: List[Any], totals: Dict[str, Any], units_by: List[Any], net_by: List[Any],
              bold: bool = False) -> None:
        for i, value in enumerate(labels, start=1):
            w.cell(row, i, value, bold=bold)
        col = n_id + 1
        for key, _h, kind in _DIMENSION_METRICS:
            value = totals.get(key)
            w.cell(row, col, _pct(value) if kind == "pct" else value, kind, bold=bold)
            col += 1
        for value in units_by:
            w.cell(row, col, value, "units", bold=bold)
            col += 1
        for value in net_by:
            w.cell(row, col, value, "money", bold=bold)
            col += 1

    if not groups:
        w.cell(w.row, 1, f"No sales in this window, so there is nothing to show by {what}.")
        w.row += 1
    for group in groups:
        bp = group.get("by_period") or {}
        write(w.row, [group.get(k) for k, _h in id_columns], group.get("totals") or {},
              [bp.get(p, {}).get("units") for p in periods], [bp.get(p, {}).get("net_sales") for p in periods])
        w.row += 1
    if groups:
        summary_by = {r["period"]: r for r in summary["by_period"]}
        write(w.row, ["Total"] + [None] * (n_id - 1), summary["totals"],
              [summary_by.get(p, {}).get("units") for p in periods],
              [summary_by.get(p, {}).get("net_sales") for p in periods], bold=True)
        w.style_row(w.row, len(headers))
        w.row += 1
    w.blank()
    w.note(f"Contribution margin = net sales - product COGS (materials + direct labor + machine), before "
           f"{PERIOD_COSTS}, which are period costs shown on the Gross Margin sheet.")
    w.finish(freeze=f"{get_column_letter(n_id + 1)}2")


def _sheet_labor(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    labor = data["labor"]
    periods = data["periods"]
    types = labor.get("by_labor_type") or []
    totals_by = {r["period"]: r for r in labor.get("totals_by_period") or []}
    totals = labor.get("totals") or {}
    headers = ["Labor type", "Rate / hour"] + periods + ["Total"]
    ncols = len(headers)
    w.header(headers)

    blocks = [
        ("hours", "Direct labor hours", "hours", "hours"),
        ("cost", "Direct labor cost", "money", "cost"),
        ("fte", f"FTEs (hours / {data['hours_per_fte']:,.2f} per FTE per month); Total column = average", "fte",
         "avg_fte"),
    ]
    if not types:
        w.cell(w.row, 1, "No routing labor for the products sold in this window.")
        w.row += 2
        blocks = []
    for metric, title, kind, total_key in blocks:
        w.section(title)
        for entry in types:
            r = w.row
            w.cell(r, 1, entry.get("labor_type_name") or entry.get("labor_type_id"), indent=1)
            w.cell(r, 2, entry.get("hourly_rate"), "money")
            bp = entry.get("by_period") or {}
            for j, p in enumerate(periods, start=3):
                w.cell(r, j, bp.get(p, {}).get(metric), kind)
            w.cell(r, ncols, (entry.get("totals") or {}).get(total_key), kind)
            w.row += 1
        r = w.row
        w.cell(r, 1, "Total direct labor")
        for j, p in enumerate(periods, start=3):
            w.cell(r, j, totals_by.get(p, {}).get(metric), kind)
        w.cell(r, ncols, totals.get(total_key), kind)
        w.style_row(r, ncols)
        w.row += 1
        w.blank()

    if types:
        w.note(f"Peak month: {totals.get('peak_fte', 0):,.2f} FTEs. Average: {totals.get('avg_fte', 0):,.2f} FTEs.")
        w.blank()

    # Standard (routings) vs the pay of employees classed Direct, by month.
    comparison = labor.get("standard_vs_payroll") or {}
    by_month = {r["period"]: r for r in comparison.get("by_period") or []}
    comp_totals = comparison.get("totals") or {}
    w.section("Standard vs payroll (employees classed Direct)")
    for key, label, kind in (
        ("standard_hours", "Standard hours (routings)", "hours"),
        ("paid_hours", "Paid hours (Direct employees)", "hours"),
        ("utilization_pct", "Utilization (standard / paid hours)", "pct"),
        ("standard_cost", "Direct labor (standard)", "money"),
        ("direct_payroll_cost", "Direct payroll", "money"),
        ("labor_variance", "Direct labor variance (payroll vs standard)", "money"),
        ("indirect_payroll_cost", "Indirect labor (factory)", "money"),
    ):
        r = w.row
        w.cell(r, 1, label, indent=1)
        for j, p in enumerate(periods, start=3):
            value = by_month.get(p, {}).get(key)
            w.cell(r, j, _pct(value) if kind == "pct" else value, kind)
        total = comp_totals.get(key)
        w.cell(r, ncols, _pct(total) if kind == "pct" else total, kind)
        w.row += 1
    note = labor.get("payroll_note") or {}
    r = w.row
    w.cell(r, 1, "Payroll in opex (admin, not classified)", indent=1)
    w.cell(r, ncols, note.get("admin_payroll_cost"), "money")
    w.row += 1
    if note.get("message"):
        w.note(str(note["message"]))
    w.blank()

    products = labor.get("by_product") or []
    w.section("Labor by product")
    w.header(["Product", "Units", "Hours / unit", "Labor hours", "Labor cost", "Machine hours"])
    for prod in products:
        r = w.row
        name, unit_id = prod.get("unit_name"), prod.get("unit_id")
        w.cell(r, 1, f"{name} ({unit_id})" if name and unit_id else (name or unit_id), indent=1)
        w.cell(r, 2, prod.get("units"), "units")
        w.cell(r, 3, prod.get("labor_hours_per_unit"), "hours")
        w.cell(r, 4, prod.get("labor_hours"), "hours")
        w.cell(r, 5, prod.get("labor_cost"), "money")
        w.cell(r, 6, prod.get("machine_hours"), "hours")
        w.row += 1
    if not products:
        w.cell(w.row, 1, "No products sold in this window.")
        w.row += 1
    w.finish(freeze="C2")


def _sheet_machines(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    labor = data["labor"]
    periods = data["periods"]
    machines = labor.get("by_machine") or []
    headers = ["Machine", "Rate / hour", "Available hours / month"] + periods + ["Total"]
    ncols = len(headers)
    w.header(headers)
    if not machines:
        w.cell(w.row, 1, "No machine time in the routings of the products sold in this window.")
        w.row += 1
        w.finish(freeze="D2")
        return

    # Total rows: every machine together, month by month as the statements add it (the cost row is the Income
    # Statement's Machine / manufacturing line), never a sum of the rounded per-machine cells.
    totals_by_period = {r["period"]: r for r in labor.get("machine_totals_by_period") or []}
    window_totals = labor.get("machine_totals") or {}
    blocks = [
        ("hours", "Machine hours required", "hours", "hours", True),
        ("utilization_pct", "Utilization (hours required / available hours); Total column = peak month", "pct",
         "peak_utilization_pct", False),
        ("cost", "Machine cost", "money", "cost", True),
    ]
    for metric, title, kind, total_key, has_total in blocks:
        w.section(title)
        for entry in machines:
            r = w.row
            w.cell(r, 1, entry.get("machine_name") or entry.get("machine_id"), indent=1)
            w.cell(r, 2, entry.get("machine_rate"), "money")
            w.cell(r, 3, entry.get("available_hours_per_month"), "hours")
            bp = entry.get("by_period") or {}
            for j, p in enumerate(periods, start=4):
                value = bp.get(p, {}).get(metric)
                if kind == "pct":
                    value = _pct(value)
                cell = w.cell(r, j, value, kind)
                if kind == "pct" and _is_number(value) and value > 1:
                    cell.fill = _OVER_FILL
            total = (entry.get("totals") or {}).get(total_key)
            w.cell(r, ncols, _pct(total) if kind == "pct" else total, kind)
            w.row += 1
        if has_total:
            r = w.row
            w.cell(r, 1, "Total")
            for j, p in enumerate(periods, start=4):
                w.cell(r, j, totals_by_period.get(p, {}).get(metric), kind)
            w.cell(r, ncols, window_totals.get(total_key), kind)
            w.style_row(r, ncols)
            w.row += 1
        w.blank()
    w.note("Shaded cells: the machine needs more hours than it has available that month.")
    w.note(str(window_totals.get("note") or MACHINE_TOTALS_NOTE))
    w.finish(freeze="D2")


def _sheet_unit_costs(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    w.header(["Product ID", "Product", "Product type", "BOM", "BOM version", "Routing", "Material / unit",
              "Direct labor / unit", "Machine / unit", "Total cost / unit", "Labor minutes / unit",
              "Machine minutes / unit", "Data warnings"], label_cols=6)
    unit_costs = data["cogs"].get("unit_costs") or []
    for uc in unit_costs:
        r = w.row
        values = [uc.get("unit_id"), uc.get("unit_name"), uc.get("unit_type"), uc.get("bom_id"),
                  uc.get("bom_version"), uc.get("router_id")]
        for i, value in enumerate(values, start=1):
            w.cell(r, i, value)
        w.cell(r, 7, uc.get("material_per_unit"), "money")
        w.cell(r, 8, uc.get("labor_per_unit"), "money")
        w.cell(r, 9, uc.get("machine_per_unit"), "money")
        w.cell(r, 10, uc.get("total_per_unit"), "money", bold=True)
        w.cell(r, 11, uc.get("labor_minutes_per_unit"), "minutes")
        w.cell(r, 12, uc.get("machine_minutes_per_unit"), "minutes")
        w.cell(r, 13, "; ".join(COST_WARNING_TEXT.get(c, c) for c in sorted(set(uc.get("warnings") or []))))
        w.row += 1
    if not unit_costs:
        w.cell(w.row, 1, "No products sold in this window.")
        w.row += 1
    w.finish(freeze="C2")


def _sheet_sales_detail(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    rows = data["sales_detail"]
    columns = sales_detail_columns(rows)
    w.header([h for _k, h, _f in columns], label_cols=9)
    for row in rows:
        r = w.row
        for i, (key, _h, kind) in enumerate(columns, start=1):
            w.cell(r, i, row.get(key), kind)
        w.row += 1
    if rows:
        last = w.row - 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"
        totals = data.get("sales_detail_totals") or sales_detail_totals(rows)
        r = w.row
        w.cell(r, 1, "Total")
        for i, (key, _h, kind) in enumerate(columns, start=1):
            if key in totals and i > 1:
                w.cell(r, i, totals[key], kind)
        w.style_row(r, len(columns))
        w.row += 1
        rounding = data.get("sales_detail_rounding") or {}
        report_totals = data.get("sales_detail_report_totals") or {}
        if any(abs(v) >= 0.005 for v in rounding.values()):
            # Lines are rounded to the cent; the statements add up the
            # unrounded lines. Bridge the SUM of the lines to the statements.
            r = w.row
            w.cell(r, 1, "Rounding (lines rounded to the cent)")
            for i, (key, _h, kind) in enumerate(columns, start=1):
                if key in rounding and i > 1:
                    w.cell(r, i, rounding[key], kind)
            w.row += 1
            r = w.row
            w.cell(r, 1, "Total per Income Statement")
            for i, (key, _h, kind) in enumerate(columns, start=1):
                if key in rounding and key in report_totals and i > 1:
                    w.cell(r, i, report_totals[key], kind)
            w.style_row(r, len(columns))
            w.row += 1
        w.blank()
        w.note("Each line's amounts are rounded to the cent (unit price to 4 decimals), so every line cross-foots "
               "and Total is the SUM of the lines. The statements add up the unrounded lines; any difference is "
               "shown on the Rounding row.")
    else:
        w.cell(w.row, 1, "No sales lines in this window.")
        w.row += 1
    w.finish(freeze="B2")


def _sheet_deduction_rules(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    w.header(["Rule ID", "Deduction type", "Treatment", "Name", "Method", "Rate", "Amount per unit", "Scenario",
              "Customer ID", "Customer", "Customer type", "Product ID", "Product", "Product type", "Start period",
              "End period", "Active", "Specificity", "Notes"], label_cols=19)
    type_order = {t: i for i, t in enumerate(DEDUCTION_TYPES)}
    rules = sorted(data["deduction_rules"], key=lambda r: (
        type_order.get(r.get("deduction_type"), 99), -deduction_specificity(r), str(r.get("rule_id") or "")))
    for rule in rules:
        r = w.row
        meta = DEDUCTION_TYPES.get(rule.get("deduction_type") or "", {})
        unit = data["units"].get(rule.get("unit_id")) or (None,)
        values: List[Tuple[Any, Optional[str]]] = [
            (rule.get("rule_id"), None),
            (meta.get("label") or rule.get("deduction_type"), None),
            (TREATMENT_LABELS.get(meta.get("treatment", ""), meta.get("treatment")), None),
            (rule.get("name"), None),
            (rule.get("method"), None),
            (rule.get("rate"), "rate"),
            (rule.get("amount_per_unit"), "money"),
            (_scenario_label(data, rule.get("forecast_id")), None),
            (rule.get("customer_id"), None),
            (data["customers"].get(rule.get("customer_id")), None),
            (rule.get("customer_type"), None),
            (rule.get("unit_id"), None),
            (unit[0], None),
            (rule.get("unit_type"), None),
            (rule.get("start_period"), None),
            (rule.get("end_period"), None),
            ("No" if str(rule.get("is_active")) in ("0", "False") else "Yes", None),
            (deduction_specificity(rule), "int"),
            (rule.get("notes"), None),
        ]
        for i, (value, kind) in enumerate(values, start=1):
            w.cell(r, i, value, kind)
        w.row += 1
    if not rules:
        w.cell(w.row, 1, "No deduction rules are set up, so net sales equal gross sales.")
        w.row += 1
    w.blank()
    w.note("Blank scope fields mean 'any'. Rates are fractions of the base shown by the method "
           "(embedded_tax: 11% means price x 0.11 / 1.11).")
    w.finish(freeze="B2")


def _sheet_price_list(ws, data: Dict[str, Any]) -> None:
    w = _Writer(ws)
    w.header(["Price ID", "Applies to", "Product ID", "Product", "Product type", "Customer ID", "Customer", "Customer type",
              "Method", "Unit price", "Discount off base", "Product base price", "Resulting price", "Scenario",
              "Start period", "End period", "Active", "Notes"], label_cols=18)
    rules = sorted(data["price_rules"], key=lambda r: (
        str(r.get("unit_id") or "~"), -price_specificity(r), str(r.get("price_id") or "")))
    for rule in rules:
        r = w.row
        unit = data["units"].get(rule.get("unit_id")) or (None, None, None)
        base_price = unit[2] if rule.get("unit_id") else None
        resulting = rule_price(rule, base_price) if (rule.get("unit_id") or rule.get("method") != "discount_from_base") else None
        values: List[Tuple[Any, Optional[str]]] = [
            (rule.get("price_id"), None),
            (PRICE_SOURCE_LABELS.get(price_source(rule), price_source(rule)), None),
            (rule.get("unit_id"), None),
            (unit[0] if rule.get("unit_id") else (
                f"All {rule.get('unit_type')} products" if rule.get("unit_type") else "All products"), None),
            (rule.get("unit_type"), None),
            (rule.get("customer_id"), None),
            (data["customers"].get(rule.get("customer_id")), None),
            (rule.get("customer_type"), None),
            (rule.get("method") or "fixed", None),
            (rule.get("unit_price"), "price"),
            (rule.get("discount_pct"), "rate"),
            (base_price, "price"),
            (resulting, "price"),
            (_scenario_label(data, rule.get("forecast_id")), None),
            (rule.get("start_period"), None),
            (rule.get("end_period"), None),
            ("No" if str(rule.get("is_active")) in ("0", "False") else "Yes", None),
            (rule.get("notes"), None),
        ]
        for i, (value, kind) in enumerate(values, start=1):
            w.cell(r, i, value, kind)
        w.row += 1
    if not rules:
        w.cell(w.row, 1, "No price list entries; sales rows use the price they were imported with, or the "
                         "product's base price.")
        w.row += 1
    w.blank()
    w.note("The most specific matching price wins: customer + product, customer (all products), customer type + "
           "product, customer type (all products), product, then the product's base price.")
    w.note("Methods: " + "; ".join(f"{m}: {text}" for m, text in PRICE_METHODS.items()) + ".")
    w.finish(freeze="B2")


def build_workbook(cursor, forecast_ids: Sequence[str], start_period: str, end_period: str,
                   hours_per_fte: Optional[float] = None) -> Workbook:
    """
    Build the forecast workbook (see the module docstring for the sheets).

    Raises ValueError for an empty scenario list, a bad window or a
    non-positive hours_per_fte.
    """
    data = collect_workbook_data(cursor, forecast_ids, start_period, end_period, hours_per_fte)
    wb = Workbook()
    sheets = {name: (wb.active if i == 0 else wb.create_sheet()) for i, name in enumerate(SHEET_NAMES)}
    for name, ws in sheets.items():
        ws.title = name
    cogs = data["cogs"]

    _sheet_read_me(sheets["Read Me"], data)
    _sheet_gross_margin(sheets["Gross Margin"], data)
    _sheet_statement(sheets["Income Statement"], data, "income_statement")
    _sheet_dimension(sheets["By Product"], data, cogs.get("by_product") or [],
                     [("unit_id", "Product ID"), ("unit_name", "Product"), ("unit_type", "Product type")], "product")
    _sheet_dimension(sheets["By Customer"], data, cogs.get("by_customer") or [],
                     [("customer_id", "Customer ID"), ("customer_name", "Customer"),
                      ("customer_type", "Customer type")], "customer")
    _sheet_dimension(sheets["By Customer Type"], data, cogs.get("by_customer_type") or [],
                     [("customer_type", "Customer type")], "customer type")
    _sheet_labor(sheets["Labor Forecast"], data)
    _sheet_machines(sheets["Machine Hours"], data)
    _sheet_unit_costs(sheets["Unit Costs"], data)
    _sheet_sales_detail(sheets["Sales Detail"], data)
    _sheet_deduction_rules(sheets["Deduction Rules"], data)
    _sheet_price_list(sheets["Price List"], data)
    _sheet_statement(sheets["Cash Flow"], data, "cash_flow")
    _sheet_statement(sheets["Balance Sheet"], data, "balance_sheet")

    names = ", ".join(data["scenario_names"].get(f) or f for f in data["forecast_ids"])
    wb.properties.title = f"Forecast {names} {data['start_period']} to {data['end_period']}"
    wb.properties.creator = "Forecast model"
    wb.active = 0
    return wb


def workbook_bytes(cursor, forecast_ids: Sequence[str], start_period: str, end_period: str,
                   hours_per_fte: Optional[float] = None) -> bytes:
    """The workbook as .xlsx bytes (for an HTTP download)."""
    buffer = io.BytesIO()
    build_workbook(cursor, forecast_ids, start_period, end_period, hours_per_fte).save(buffer)
    return buffer.getvalue()


def write_workbook(path: str, cursor, forecast_ids: Sequence[str], start_period: str, end_period: str,
                   hours_per_fte: Optional[float] = None) -> Dict[str, Any]:
    """
    Save the workbook to `path` (parent folders are created). Returns
    {"path", "filename", "sheets", "bytes"}.
    """
    wb = build_workbook(cursor, forecast_ids, start_period, end_period, hours_per_fte)
    folder = os.path.dirname(os.path.abspath(path))
    os.makedirs(folder, exist_ok=True)
    wb.save(path)
    return {
        "path": os.path.abspath(path),
        "filename": os.path.basename(path),
        "sheets": list(wb.sheetnames),
        "bytes": os.path.getsize(path),
    }


_UNSAFE = re.compile(r"[^A-Za-z0-9._+-]+")


def scenario_names(cursor, forecast_ids: Sequence[str]) -> Dict[str, str]:
    """{forecast_id: name} of the given scenarios (unknown ids are left out)."""
    ids = [str(f).strip() for f in forecast_ids if f and str(f).strip()]
    if not ids:
        return {}
    cursor.execute(f"SELECT forecast_id, name FROM forecast WHERE forecast_id IN ({','.join('?' for _ in ids)})", ids)
    return {fid: name for fid, name in cursor.fetchall() if name is not None and str(name).strip()}


def _slug(text: Any, limit: int = 40) -> str:
    """'2027 Budget (v2)' -> '2027-budget-v2': lower-case letters and digits joined by '-'."""
    return re.sub(r"[^a-z0-9]+", "-", str(text or "").lower()).strip("-")[:limit].strip("-")


def workbook_filename(forecast_ids: Sequence[str], start_period: str, end_period: str,
                      names: Optional[Dict[str, str]] = None) -> str:
    """
    forecast_<scenario name slug>_<start>_<end>.xlsx, e.g.
    forecast_2027-budget_2027-01_2027-12.xlsx. Each scenario is named by
    the slug of its name (``names``: {forecast_id: name}), falling back to
    its id; several scenarios are joined with '+' (at most three, then
    '+N-more'). Unsafe characters are replaced.
    """
    ids = [str(f).strip() for f in forecast_ids if f and str(f).strip()] or ["forecast"]
    parts = [_slug((names or {}).get(fid)) or fid for fid in ids]
    label = "+".join(parts[:3]) + (f"+{len(parts) - 3}-more" if len(parts) > 3 else "")
    stem = f"forecast_{label}_{start_period}_{end_period}"
    return _UNSAFE.sub("_", stem).strip("._") + ".xlsx"


def export_path(data_dir: str, forecast_ids: Sequence[str], start_period: str, end_period: str,
                filename: Optional[str] = None, names: Optional[Dict[str, str]] = None) -> str:
    """
    Where to save an export: <data_dir>/exports/<filename>. A caller-supplied
    filename is reduced to a safe base name (no folders) ending in .xlsx;
    without one it is ``workbook_filename`` (``names``: {forecast_id: name}).
    """
    if filename:
        base = os.path.basename(str(filename).replace("\\", "/"))
        stem = base[:-5] if base.lower().endswith(".xlsx") else base
        stem = _UNSAFE.sub("_", stem).strip("._")
        name = (stem or "forecast") + ".xlsx"
    else:
        name = workbook_filename(forecast_ids, start_period, end_period, names)
    return os.path.join(data_dir, "exports", name)
