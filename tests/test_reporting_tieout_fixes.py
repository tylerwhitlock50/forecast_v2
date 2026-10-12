"""
Reporting tie-out findings (review round v13), one test group per finding:

* LABOR-CENT-SALES-DETAIL / SD-LABOR-CENT - the sales detail's
  statement-exact ``report_totals`` (REST, MCP and the workbook's 'Total per
  Income Statement' row) add the lines up month by month as the statements
  do, so every column equals the income statement to the cent. The 2027
  sample's direct labor sits on a half cent (449,997.575 to the float): one
  running sum over the lines gives .58, the statements' month-by-month sum .57.
* README-GP-DEFINITION - every stated gross profit formula (workbook Read Me,
  Gross Margin note, per-group footnotes) includes the direct labor variance
  and indirect labor, and applying it reproduces the workbook's gross profit.
* MCP-LABOR-NOTE-FIELDS - forecast_labor_forecast passes through every
  payroll_note field and every standard-vs-payroll monthly column of the API.
"""
import csv
import io
import json
import os
import shutil
import sys
import tempfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

_REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
_APP_DIR = os.path.join(_REPO, "app")
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from app.db.database import DatabaseManager  # noqa: E402

SAMPLE_DIR = os.path.join(_REPO, "examples", "sample_import_2027")
WINDOW = ("2027-01", "2027-12")
HEADERS = {"host": "localhost:8000", "accept": "application/json, text/event-stream",
           "content-type": "application/json"}
# Sales detail column -> income statement line.
STATEMENT_PAIRS = {
    "quantity": "units_sold",
    "gross_sales": "gross_sales",
    "terms_allowance": "terms_allowance",
    "chargebacks": "chargebacks",
    "other_allowances": "other_allowances",
    "excise_tax": "excise_tax",
    "total_sales_deductions": "total_sales_deductions",
    "net_sales": "revenue",
    "commissions": "commissions",
    "materials": "cogs_materials",
    "labor": "cogs_labor",
    "machine": "cogs_manufacturing",
}
# Sales detail columns that are not income statement lines: the COGS forecast's figures.
COGS_FORECAST_KEYS = ("product_cogs", "contribution_margin")


# ---------------------------------------------------------------------------
# Fixture: the real app on a scratch database with the 2027 sample (prices and deductions included)
# ---------------------------------------------------------------------------

def _patch_db_manager(mgr):
    originals = {}
    for name, module in list(sys.modules.items()):
        if module is None:
            continue
        if name in ("db.database", "app.db.database") or name.startswith(("api.", "app.api.", "utils.")):
            if hasattr(module, "db_manager"):
                originals[name] = module.db_manager
                module.db_manager = mgr
    return originals


def _data(res):
    assert res.status_code == 200, res.text
    return res.json()["data"]


def _csv(name):
    with open(os.path.join(SAMPLE_DIR, name), newline="", encoding="utf-8-sig") as f:
        return [dict(r) for r in csv.DictReader(f)]


def _load_sample(client):
    fid = _data(client.post("/forecast/scenario", json={"name": "2027 Budget"}))["forecast_id"]
    for table, name in (("customers", "customers.csv"), ("units", "products.csv"), ("machines", "machines.csv"),
                        ("labor_rates", "labor_rates.csv")):
        _data(client.post("/import/records", json={"table": table, "rows": _csv(name)}))
    _data(client.post("/import/routings", json={"rows": _csv("routings.csv")}))
    sheet = load_workbook(os.path.join(SAMPLE_DIR, "boms.xlsx"), data_only=True).active
    raw = list(sheet.iter_rows(values_only=True))
    _data(client.post("/import/boms", json={"rows": [dict(zip(raw[1], r)) for r in raw[2:]
                                                     if any(v is not None for v in r)]}))
    _data(client.post("/import/prices", json={"rows": _csv("price_list.csv")}))
    _data(client.post("/import/deductions", json={"rows": _csv("deduction_rules.csv")}))
    _data(client.post("/forecast/smart_import_sales", json={"forecast_id": fid, "rows": _csv("sales_forecast_2027.csv")}))
    # Classed staff, so the income statement carries a direct labor variance and indirect labor.
    base = {"department": "Production", "weekly_hours": 40, "rate_type": "hourly", "labor_type": "Production",
            "start_date": "2020-01-01", "benefits_eligible": False, "expected_raise": 0}
    for name, rate, labor_class in (("Dana Direct", 28, "direct"), ("Dale Direct", 30, "direct"),
                                    ("Ivy Indirect", 32, "indirect"), ("Abe Admin", 40, "admin")):
        _data(client.post("/payroll/employees", json={**base, "employee_name": name, "hourly_rate": rate,
                                                       "labor_class": labor_class}))
    return fid


@pytest.fixture(scope="module")
def env():
    tmp = tempfile.mkdtemp()
    mgr = DatabaseManager(database_path=os.path.join(tmp, "forecast.db"), data_dir=tmp)
    mgr.initialize()
    from app.main import app as fastapi_app
    originals = _patch_db_manager(mgr)
    try:
        with TestClient(fastapi_app) as client:
            fid = _load_sample(client)
            yield client, mgr, fid
    finally:
        for name, original in originals.items():
            sys.modules[name].db_manager = original
        mgr.close_all_connections()
        shutil.rmtree(tmp, ignore_errors=True)


def _params(fid, **extra):
    return {"forecast_ids": fid, "start_period": WINDOW[0], "end_period": WINDOW[1], **extra}


def _mcp(client, name, arguments):
    res = client.post("/mcp", headers=HEADERS, follow_redirects=False, json={
        "jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": name, "arguments": arguments}})
    assert res.status_code == 200, res.text
    result = res.json()["result"]
    text = "".join(b.get("text", "") for b in result.get("content", []) if b.get("type") == "text")
    assert not result.get("isError"), text
    return json.loads(text)


@pytest.fixture(scope="module")
def reports(env):
    client, _mgr, fid = env
    statements = _data(client.get("/reporting/statements", params=_params(fid)))
    return {
        "statements": statements,
        "income": statements["income_statement"]["totals"],
        "income_by_period": {r["period"]: r for r in statements["income_statement"]["by_period"]},
        "cogs": _data(client.get("/reporting/cogs-forecast", params=_params(fid)))["summary"]["totals"],
        "labor": _data(client.get("/reporting/labor-forecast", params=_params(fid))),
        "detail": _data(client.get("/reporting/sales-detail", params=_params(fid, limit=1))),
    }


@pytest.fixture(scope="module")
def workbook(env):
    client, _mgr, fid = env
    res = client.get("/reporting/workbook", params=_params(fid))
    assert res.status_code == 200, res.text
    return load_workbook(io.BytesIO(res.content), data_only=True)


def _label_rows(ws):
    return {row[0]: row for row in ws.iter_rows(values_only=True) if row and isinstance(row[0], str)}


def _sheet_text(ws):
    return "\n".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)


# ---------------------------------------------------------------------------
# LABOR-CENT-SALES-DETAIL / SD-LABOR-CENT
# ---------------------------------------------------------------------------

def test_the_sample_sits_on_a_half_cent_that_line_order_rounds_the_other_way(env, reports):
    """Guards the regression test's power: on this data the two addition orders disagree."""
    _client, mgr, fid = env
    from services.financial_statements import load_sales_lines, period_range, window_total
    conn = mgr.get_connection()
    try:
        lines = load_sales_lines(conn.cursor(), [fid], period_range(*WINDOW))["lines"]
    finally:
        conn.close()
    running, by_month = 0.0, {}
    for line in lines:
        running += line["labor"]
        by_month[line["period"]] = by_month.get(line["period"], 0.0) + line["labor"]
    month_by_month = window_total(by_month, period_range(*WINDOW))
    assert round(month_by_month, 2) == reports["income"]["cogs_labor"]
    assert round(running, 2) != reports["income"]["cogs_labor"]


def test_sales_detail_report_totals_equal_the_income_statement_to_the_cent(reports):
    report = reports["detail"]["report_totals"]
    income = reports["income"]
    for key, line in STATEMENT_PAIRS.items():
        assert report[key] == income[line], (key, report[key], income[line])
    for key in COGS_FORECAST_KEYS:
        assert report[key] == reports["cogs"][key], (key, report[key], reports["cogs"][key])
    # The labor forecast, COGS forecast and statement agree on the one direct labor figure.
    assert report["labor"] == reports["labor"]["totals"]["cost"] == reports["cogs"]["labor"]
    # rounding still bridges the per-line SUM to the statement figure.
    detail = reports["detail"]
    for key, value in detail["rounding"].items():
        assert value == round(report[key] - detail["totals"][key], 2), key


def test_report_totals_add_months_not_lines_on_a_minimal_half_cent_case():
    """Four one-unit lines whose labor totals 1,804.075: line order gives .08, month by month .07."""
    import sqlite3
    from services.financial_statements import build_statements
    from services.operations_forecast import sales_detail_report
    rates = {"P1": ("2027-01", 82.58), "P2": ("2027-01", 267.19), "P3": ("2027-02", 889.987),
             "P4": ("2027-02", 564.318)}
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, "half_cent.db")
    mgr = DatabaseManager(database_path=path, data_dir=tmp)
    try:
        mgr.initialize()
        conn = sqlite3.connect(path)
        c = conn.cursor()
        c.execute("INSERT INTO forecast (forecast_id, name, is_base) VALUES ('F-HC', 'Half cent', 1)")
        c.execute("INSERT INTO customers (customer_id, customer_name, customer_type) VALUES ('C1', 'Acme', 'Dealer')")
        c.execute("INSERT INTO machines (machine_id, machine_name, machine_rate) VALUES ('BENCH', 'Bench', 0)")
        for unit, (period, rate) in rates.items():
            c.execute("INSERT INTO labor_rates (rate_id, rate_name, rate_amount) VALUES (?, ?, ?)",
                      (f"L-{unit}", f"Rate {unit}", rate))
            c.execute("INSERT INTO units (unit_id, unit_name, base_price, router_id) VALUES (?, ?, 100, ?)",
                      (unit, unit, f"R-{unit}"))
            c.execute("""INSERT INTO router_operations (router_id, sequence, machine_id, machine_minutes, labor_minutes,
                         labor_type_id) VALUES (?, 10, 'BENCH', 0, 60, ?)""", (f"R-{unit}", f"L-{unit}"))
            c.execute("""INSERT INTO sales (sale_id, customer_id, unit_id, period, quantity, unit_price, total_revenue,
                         forecast_id) VALUES (?, 'C1', ?, ?, 1, 100, 100, 'F-HC')""", (f"S-{unit}", unit, period))
        c.execute("""INSERT INTO financial_assumptions (forecast_id, opening_period, opening_cash, dso_days, dio_days,
                     dpo_days, income_tax_rate) VALUES ('F-HC', '2026-12', 1000, 30, 30, 30, 0.25)""")
        conn.commit()
        cur = conn.cursor()
        running = 0.0
        for _unit, (_period, rate) in rates.items():
            running += rate
        stmt = build_statements(cur, ["F-HC"], "2027-01", "2027-02")["income_statement"]["totals"]
        report = sales_detail_report(cur, ["F-HC"], "2027-01", "2027-02")["report_totals"]
        conn.close()
        assert round(running, 2) == 1804.08                  # what one running sum over the lines gave
        assert round(stmt["cogs_labor"], 2) == 1804.07
        assert report["labor"] == 1804.07
        assert report["product_cogs"] == 1804.07
    finally:
        mgr.close_all_connections()
        shutil.rmtree(tmp, ignore_errors=True)


def test_a_month_filter_gives_that_months_statement_figures(env, reports):
    client, _mgr, fid = env
    for period in ("2027-03", "2027-12"):
        month = _data(client.get("/reporting/sales-detail", params=_params(fid, period=period, limit=1)))
        row = reports["income_by_period"][period]
        for key, line in STATEMENT_PAIRS.items():
            assert month["report_totals"][key] == row[line], (period, key)


def test_mcp_sales_tools_quote_the_statement_figures(env, reports):
    client, _mgr, fid = env
    income = reports["income"]
    listed = _mcp(client, "forecast_list_sales", {"forecast_id": fid, "limit": 0})
    detail = _mcp(client, "forecast_sales_detail", {"forecast_ids": [fid], "limit": 1})
    for answer in (listed, detail):
        assert answer["report_totals"] == reports["detail"]["report_totals"]
        assert answer["report_totals"]["labor"] == income["cogs_labor"]


def test_workbook_total_per_income_statement_row_ties_to_the_income_statement_sheet(workbook, reports):
    ws = workbook["Sales Detail"]
    header = [c.value for c in ws[1]]
    rows = _label_rows(ws)
    assert "Total per Income Statement" in rows
    statement_row = rows["Total per Income Statement"]
    income_sheet = _label_rows(workbook["Income Statement"])
    gross_margin = _label_rows(workbook["Gross Margin"])
    labels = {"gross_sales": "Gross sales", "terms_allowance": "Terms allowance",
              "chargebacks": "Chargebacks allowance", "excise_tax": "Excise tax (FET)",
              "total_sales_deductions": "Total sales deductions", "net_sales": "Net sales",
              "commissions": "Sales commissions", "materials": "Materials", "labor": "Direct labor",
              "machine": "Machine"}
    income_labels = {"gross_sales": "Gross sales", "terms_allowance": "Terms allowance",
                     "chargebacks": "Chargebacks allowance", "excise_tax": "Excise tax (FET)",
                     "total_sales_deductions": "Total sales deductions", "net_sales": "Net sales",
                     "commissions": "Sales commissions", "materials": "Materials",
                     "labor": "Direct labor (standard)", "machine": "Machine / manufacturing"}
    for key, column in labels.items():
        value = statement_row[header.index(column)]
        assert value == income_sheet[income_labels[key]][-1], (key, value)
    product_cogs = statement_row[header.index("Product COGS")]
    assert product_cogs == gross_margin["Product COGS (BOM + routing)"][-1]
    # The case the reviewers saw: .57 on the Income Statement sheet, .58 on this row.
    assert statement_row[header.index("Direct labor")] == reports["income"]["cogs_labor"]


# ---------------------------------------------------------------------------
# README-GP-DEFINITION
# ---------------------------------------------------------------------------

def test_read_me_gross_profit_formula_names_every_cogs_line_and_reproduces_gross_profit(workbook, reports):
    rows = _label_rows(workbook["Read Me"])
    formula = str(rows["Contribution margin vs gross profit"][1])
    for part in ("product COGS", "direct labor variance", "indirect labor", "factory overhead", "other COGS"):
        assert part in formula, part
    period_costs = str(rows["Period costs in COGS"][1])
    for part in ("direct labor variance", "indirect labor", "factory overhead", "other COGS"):
        assert part in period_costs, part
    # The formula, applied to the Gross Margin sheet, gives that sheet's gross profit (not contribution margin).
    margin = {label: row[-1] for label, row in _label_rows(workbook["Gross Margin"]).items()}
    variance = margin["Direct labor variance (payroll vs standard)"]
    indirect = margin["Indirect labor (factory)"]
    assert variance and indirect, "the fixture's classed staff must move gross profit"
    total_cogs = (margin["Product COGS (BOM + routing)"] + variance + indirect + margin["Factory overhead"]
                  + margin["Other COGS"])
    assert round(margin["Net sales"] - total_cogs, 2) == pytest.approx(margin["Gross profit"], abs=0.011)
    assert margin["Gross profit"] == pytest.approx(reports["income"]["gross_profit"], abs=0.011)
    assert margin["Gross profit"] != pytest.approx(margin["Contribution margin (before period costs)"], abs=1)


def test_gross_margin_note_and_group_footnotes_name_the_labor_lines(workbook):
    note = _sheet_text(workbook["Gross Margin"])
    assert "direct labor variance" in note and "indirect labor" in note
    for name in ("By Product", "By Customer", "By Customer Type"):
        text = _sheet_text(workbook[name])
        assert "before the direct labor variance, indirect labor, factory overhead and other COGS" in text, name


# ---------------------------------------------------------------------------
# MCP-LABOR-NOTE-FIELDS
# ---------------------------------------------------------------------------

def test_mcp_labor_forecast_passes_every_payroll_note_field(env, reports):
    client, _mgr, fid = env
    api = reports["labor"]
    tool = _mcp(client, "forecast_labor_forecast", {"forecast_ids": [fid], "start_period": WINDOW[0],
                                                    "end_period": WINDOW[1]})
    note = tool["payroll_note"]
    for key in ("cogs_direct_labor", "months_with_direct_pay", "standard_only_months", "standard_only_labor",
                "direct_labor_cost", "labor_variance", "indirect_payroll_cost", "admin_payroll_cost", "message"):
        assert key in api["payroll_note"], key
    for key, value in api["payroll_note"].items():
        if value is None or value == "":
            continue
        assert note[key] == value, key
    assert note["months_with_direct_pay"] == list(api["periods"])
    assert note["cogs_direct_labor"] == round(note["direct_labor_cost"] + note["labor_variance"], 2)
    svp = tool["standard_vs_payroll"]
    rows = api["standard_vs_payroll"]["by_period"]
    for metric in (k for k in rows[0] if k != "period"):
        assert svp[metric] == [r[metric] for r in rows], metric
    assert "indirect_payroll_cost" in svp


def test_a_product_customer_or_type_filter_equals_that_group_of_the_cogs_forecast(env):
    # Filtered totals are added line by line, as the COGS forecast adds a group, so they tie to the group.
    client, _mgr, fid = env
    cogs = _data(client.get("/reporting/cogs-forecast", params=_params(fid)))
    cases = [("unit_id", "by_product", "unit_id"), ("customer_id", "by_customer", "customer_id"),
             ("customer_type", "by_customer_type", "customer_type"), ("unit_type", "by_product_type", "unit_type")]
    for param, group_key, field in cases:
        for group in cogs[group_key][:3]:
            detail = _data(client.get("/reporting/sales-detail", params=_params(fid, limit=1, **{param: group[field]})))
            for key in ("gross_sales", "net_sales", "materials", "labor", "machine", "product_cogs"):
                assert detail["report_totals"][key] == group["totals"][key], (param, group[field], key)


def test_revenue_summary_totals_equal_the_income_statement(env, reports):
    client, _mgr, fid = env
    res = client.get("/source-data/revenue-summary",
                     params={"forecast_id": fid, "start_period": WINDOW[0], "end_period": WINDOW[1], "group_by": "period"})
    assert res.status_code == 200, res.text
    data = _data(res)
    income = reports["income"]
    assert data["totals"]["labor"] == income["cogs_labor"]
    assert data["totals"]["net_sales"] == income["revenue"]
    assert "direct labor variance" in data["notes"]
