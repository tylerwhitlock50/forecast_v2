"""
Reporting findings (review round v15), one test group per finding, on the 2027 sample:

* V15-WB-MACHINE-COST-TOTAL - the labor forecast carries a statement-exact machine total
  (``machine_totals`` / ``machine_totals_by_period``), and the workbook's Machine Hours Total rows
  and the MCP labor forecast / machine utilization tools quote it, so the machine cost total equals
  the income statement's Machine / manufacturing line to the cent (the per-machine rows, rounded one
  by one, add up to a cent more on this sample).
* V15-ACTUALS-GROSS-TO-NET-LINES - actuals can be recorded for terms allowance, chargebacks, other
  allowances and excise tax (FET) as their own budget vs actual lines; Sales Deductions stays their sum
  (a month carries either the detail lines or the total; both warns).
* V15-README-30.44 - the workbook Read Me states the engine's working-capital divisor (365.25 / 12).
"""
import io
import os
import re
import sys

import pytest
from openpyxl import load_workbook

_REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
_APP_DIR = os.path.join(_REPO, "app")
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from tests.test_reporting_tieout_fixes import (  # noqa: E402,F401  (fixture + helpers)
    WINDOW, _data, _label_rows, _mcp, _params, env,
)

DETAIL_LINES = ("terms_allowance", "chargebacks", "other_allowances", "excise_tax")


@pytest.fixture(scope="module")
def reports(env):
    client, _mgr, fid = env
    statements = _data(client.get("/reporting/statements", params=_params(fid)))
    return {
        "statements": statements,
        "income": statements["income_statement"]["totals"],
        "income_by_period": {r["period"]: r for r in statements["income_statement"]["by_period"]},
        "balance_by_period": {r["period"]: r for r in statements["balance_sheet"]["by_period"]},
        "labor": _data(client.get("/reporting/labor-forecast", params=_params(fid))),
    }


@pytest.fixture(scope="module")
def workbook(env):
    client, _mgr, fid = env
    res = client.get("/reporting/workbook", params=_params(fid))
    assert res.status_code == 200, res.text
    return load_workbook(io.BytesIO(res.content), data_only=True)


def _section_rows(ws, title):
    """The rows of one titled block on a sheet (from the title to the next blank row)."""
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r and r[0] == title)
    out = []
    for row in rows[start + 1:]:
        if not row or all(v is None for v in row):
            break
        out.append(row)
    return out


# ---------------------------------------------------------------------------
# V15-WB-MACHINE-COST-TOTAL
# ---------------------------------------------------------------------------

def test_the_rounded_machine_rows_add_up_to_a_cent_more_than_the_statement(reports):
    """Guards the regression tests' power: on this sample a sum of the rounded rows is off by a cent."""
    by_machine = reports["labor"]["by_machine"]
    rounded_sum = round(sum(m["totals"]["cost"] for m in by_machine), 2)
    assert rounded_sum == 1145764.59
    assert reports["income"]["cogs_manufacturing"] == 1145764.58


def test_labor_forecast_machine_totals_equal_the_income_statement(reports):
    labor = reports["labor"]
    assert labor["machine_totals"]["cost"] == reports["income"]["cogs_manufacturing"] == 1145764.58
    rows = {r["period"]: r for r in labor["machine_totals_by_period"]}
    assert list(rows) == labor["periods"]
    for period, row in reports["income_by_period"].items():
        assert rows[period]["cost"] == row["cogs_manufacturing"], period
    # Hours: every machine together, month by month (each month is the unrounded hours rounded once).
    hours = round(sum(r["hours"] for r in labor["machine_totals_by_period"]), 2)
    assert labor["machine_totals"]["hours"] == pytest.approx(hours, abs=0.011)
    assert "a cent or two" in labor["machine_totals"]["note"]


def test_workbook_machine_cost_total_row_equals_the_income_statement(workbook, reports):
    ws = workbook["Machine Hours"]
    header = [c.value for c in ws[1]]
    periods = header[3:-1]
    assert periods == reports["labor"]["periods"]
    cost_rows = _section_rows(ws, "Machine cost")
    total = next(r for r in cost_rows if r[0] == "Total")
    for j, period in enumerate(periods, start=3):
        assert total[j] == reports["income_by_period"][period]["cogs_manufacturing"], period
    assert total[-1] == reports["income"]["cogs_manufacturing"] == 1145764.58
    # The Income Statement sheet's own figure, for a reader ticking and tying the workbook.
    income_sheet = _label_rows(workbook["Income Statement"])
    assert total[-1] == income_sheet["Machine / manufacturing"][-1]
    # Hours Total row: the labor forecast's machine_totals, not a sum of the rounded machine cells.
    hours_total = next(r for r in _section_rows(ws, "Machine hours required") if r[0] == "Total")
    machine_rows = {r["period"]: r for r in reports["labor"]["machine_totals_by_period"]}
    for j, period in enumerate(periods, start=3):
        assert hours_total[j] == machine_rows[period]["hours"], period
    assert hours_total[-1] == reports["labor"]["machine_totals"]["hours"]
    # The sheet says the machine rows can differ from the Total by a cent.
    text = "\n".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)
    assert "can add up to a cent or two more or less than the total" in text


def test_mcp_labor_forecast_and_machine_utilization_quote_the_statement_total(env, reports):
    client, _mgr, fid = env
    income = reports["income"]
    labor = _mcp(client, "forecast_labor_forecast", {"forecast_ids": [fid], "start_period": WINDOW[0],
                                                     "end_period": WINDOW[1]})
    machine_totals = labor["machine_totals"]
    assert machine_totals["totals"]["cost"] == income["cogs_manufacturing"]
    assert machine_totals["by_period"]["cost"] == [reports["income_by_period"][p]["cogs_manufacturing"]
                                                   for p in labor["periods"]]
    assert machine_totals["note"]

    utilization = _mcp(client, "forecast_machine_utilization", {"forecast_id": fid})
    assert utilization["totals"]["total_cost"] == income["cogs_manufacturing"]
    assert utilization["totals"]["total_hours_required"] == reports["labor"]["machine_totals"]["hours"]
    assert utilization["totals"]["machine_count"] == len(utilization["machines"])
    assert "monthly" not in utilization["totals"]          # per-month arrays only with include_monthly
    monthly = _mcp(client, "forecast_machine_utilization", {"forecast_id": fid, "include_monthly": True})
    assert [m["cost"] for m in monthly["totals"]["monthly"]] == [
        reports["income_by_period"][p]["cogs_manufacturing"] for p in reports["labor"]["periods"]]


# ---------------------------------------------------------------------------
# V15-README-30.44
# ---------------------------------------------------------------------------

def test_read_me_working_capital_divisor_is_the_engine_constant_and_reproduces_the_balance_sheet(workbook, reports):
    from services.financial_statements import AVG_DAYS_PER_MONTH
    text = str(_label_rows(workbook["Read Me"])["Working capital"][1])
    assert "30.44" not in text.replace("30.4375", "")
    found = re.findall(r"(?:DSO|DIO|DPO|excise payable days) / (\d+(?:\.\d+)?)", text)
    assert len(found) == 4, text
    divisors = {float(d) for d in found}
    assert divisors == {AVG_DAYS_PER_MONTH} == {30.4375}
    assert "365.25 / 12" in text
    # The stated AR formula gives the balance sheet's January receivable (to rounding of the inputs).
    dso = reports["statements"]["assumptions"]["dso_days"]
    jan = reports["income_by_period"]["2027-01"]
    ar = (jan["revenue"] + jan["excise_tax"]) * dso / divisors.pop()
    assert ar == pytest.approx(reports["balance_by_period"]["2027-01"]["accounts_receivable"], abs=0.05)


# ---------------------------------------------------------------------------
# V15-ACTUALS-GROSS-TO-NET-LINES
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name, key", [
    ("Excise tax (FET)", "excise_tax"), ("FET", "excise_tax"), ("Federal excise tax", "excise_tax"),
    ("Excise tax", "excise_tax"), ("Terms allowance", "terms_allowance"), ("Cash discount", "terms_allowance"),
    ("Terms discounts", "terms_allowance"), ("Chargebacks", "chargebacks"), ("Chargebacks allowance", "chargebacks"),
    ("Chargeback", "chargebacks"), ("Other allowances", "other_allowances"),
    ("Other discounts & allowances", "other_allowances"), ("Sales Deductions", "total_sales_deductions"),
    ("Total sales deductions", "total_sales_deductions"),
])
def test_gross_to_net_line_names_resolve(name, key):
    from services.actuals import resolve_line_item
    assert resolve_line_item(name) == (key, None)


def test_the_bare_word_tax_points_at_income_tax_or_excise_tax():
    from services.actuals import resolve_line_item
    key, why = resolve_line_item("Taxes")
    assert key is None and "Income Tax" in why and "Excise tax (FET)" in why


def test_line_catalogue_lists_the_detail_lines_with_income_statement_labels(env):
    client, _mgr, _fid = env
    items = {i["line_item"]: i for i in _data(client.get("/reporting/actuals/line-items"))["line_items"]}
    labels = {"terms_allowance": "Terms allowance", "chargebacks": "Chargebacks allowance",
              "other_allowances": "Other discounts & allowances", "excise_tax": "Excise tax (FET)"}
    for key, label in labels.items():
        assert items[key]["label"] == label
        assert items[key]["direction"] == "expense" and items[key]["statement"] == "income"
    assert "fet" in items["excise_tax"]["aliases"]
    order = list(items)
    assert order.index("gross_sales") < order.index("terms_allowance") < order.index("total_sales_deductions")


def test_the_reviewers_file_imports_instead_of_being_refused(env):
    client, _mgr, _fid = env
    csv_text = ("Month,Account,Actual\n2027-01,Excise tax (FET),36100\n2027-01,Terms allowance,5900\n"
                "2027-01,Net Sales,381829.36\n")
    res = client.post("/reporting/actuals/import", params={"dry_run": "true"},
                      files={"file": ("actuals.csv", csv_text, "text/csv")})
    assert res.status_code == 200, res.text
    report = res.json()["data"]
    assert report["outcome"] == "imported" and report["errors"] == []
    assert [r["line_item"] for r in report["rows"]] == ["excise_tax", "terms_allowance", "revenue"]


def _import(client, rows):
    res = client.post("/reporting/actuals", json={"entries": [
        {"period": p, "line_item": line, "amount": amount} for p, line, amount in rows]})
    assert res.status_code == 200, res.text
    assert res.json()["data"]["outcome"] == "imported", res.json()["data"]


def _variance(client, fid, period):
    data = _data(client.get("/reporting/variance", params={"forecast_ids": fid, "start_period": period,
                                                            "end_period": period}))
    return data, {line["line_item"]: line for line in data["line_items"]}


def test_detail_lines_are_compared_with_the_statement_and_sales_deductions_is_their_sum(env, reports):
    client, _mgr, fid = env
    period = "2027-02"
    amounts = {"terms_allowance": 5000.0, "chargebacks": 1200.0, "other_allowances": 800.0, "excise_tax": 36000.0}
    _import(client, [(period, "Gross Sales", 420000), *((period, key, v) for key, v in amounts.items())])
    data, lines = _variance(client, fid, period)
    budget = reports["income_by_period"][period]
    for key, amount in amounts.items():
        row = lines[key]["by_period"][0]
        assert row["budget"] == budget[key], key        # the income statement's own line
        assert row["actual"] == amount
        assert row["variance"] == round(amount - budget[key], 2)
        assert lines[key]["coverage"] == "complete"
    total = lines["total_sales_deductions"]
    assert total["by_period"][0]["actual"] == sum(amounts.values())
    assert total["by_period"][0]["budget"] == budget["total_sales_deductions"]
    assert total["derived_periods"] == [period] and total["coverage"] == "complete"
    codes = {w["code"]: w for w in data["warnings"]}
    assert codes["sales_deductions_from_detail"]["level"] == "info"
    assert "sales_deductions_total_and_detail" not in codes
    coverage = data["period_coverage"][0]
    assert not set(DETAIL_LINES) & set(coverage["lines_missing"]) and coverage["lines_in_total"] == []


def test_a_month_with_only_the_total_does_not_report_the_detail_lines_missing(env):
    client, _mgr, fid = env
    period = "2027-03"
    _import(client, [(period, "Gross Sales", 400000), (period, "Sales Deductions", 41000)])
    data, lines = _variance(client, fid, period)
    assert lines["total_sales_deductions"]["by_period"][0]["actual"] == 41000
    assert lines["total_sales_deductions"]["derived_periods"] == []
    for key in DETAIL_LINES:
        assert lines[key]["by_period"][0]["actual"] is None
        assert lines[key]["in_total_periods"] == [period]
        assert lines[key]["totals"]["missing_periods"] == []
        assert key not in data["lines_without_actuals"]
    assert not any(line["line_item"] in DETAIL_LINES for line in data["incomplete_lines"])
    coverage = data["period_coverage"][0]
    assert coverage["lines_in_total"] == list(DETAIL_LINES)
    assert not set(DETAIL_LINES) & set(coverage["lines_missing"])
    assert coverage["lines_total"] == len(lines) - len(DETAIL_LINES)
    codes = {w["code"] for w in data["warnings"]}
    assert "sales_deductions_total_and_detail" not in codes and "sales_deductions_from_detail" not in codes


def test_both_the_total_and_detail_lines_in_one_month_warn(env):
    client, _mgr, fid = env
    period = "2027-04"
    _import(client, [(period, "Sales Deductions", 45000), (period, "FET", 36500), (period, "Terms allowance", 6000)])
    data, lines = _variance(client, fid, period)
    # The total is compared as entered; the detail lines on their own lines.
    assert lines["total_sales_deductions"]["by_period"][0]["actual"] == 45000
    assert lines["excise_tax"]["by_period"][0]["actual"] == 36500
    warning = next(w for w in data["warnings"] if w["code"] == "sales_deductions_total_and_detail")
    assert warning["periods"] == [{"period": period, "total": 45000.0, "detail_lines": ["terms_allowance", "excise_tax"],
                                   "detail_sum": 42500.0, "difference": 2500.0}]
    assert "either the detail lines or the total, never both" in warning["message"]
    assert "2027-04" in warning["message"] and "42,500.00" in warning["message"]
