"""
Labor classification (W4): direct / indirect / admin payroll on the statements.

Reviewer's evidence: direct labor in COGS (routing standard) and every
employee's pay in operating expenses were both subtracted, so production
workers entered in Payroll were counted twice.

Policy under test (automatic, no toggle):

* admin + unclassified payroll -> operating expense ``payroll`` (as before);
* indirect payroll             -> COGS ``indirect_labor``;
* direct payroll               -> COGS: ``cogs_labor`` stays the routing
  standard and ``labor_variance`` = direct payroll - standard in every month
  with direct payroll, so total direct labor = direct payroll;
* no employee classed direct   -> unchanged (standard only);
* inventory (DIO) stays on product COGS at standard.

Seed (zero burden, monthly pay so the arithmetic is exact):
  product P1: 60 labor minutes @ $30/h = $30 standard labor, $20 material
  sales (both scenarios): Jan 100 units, Feb 200 units, Mar none
  F-LAB employees: D1 direct  $25/h x 40 h  -> 52,000 / yr = 4,333.33 / month
                   I1 indirect salary 60,000 -> 5,000 / month
                   A1 admin    salary 72,000 -> 6,000 / month
                   U1 (none)   salary 36,000 -> 3,000 / month
  F-STD employees: S1 unclassified $25/h x 40 h (a production worker nobody classed)
"""
import io
import os
import shutil
import sqlite3
import sys
import tempfile

import pytest
from openpyxl import load_workbook

_APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, "app"))
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from app.db.database import DatabaseManager  # noqa: E402
from services import financial_statements as fs  # noqa: E402
from services import operations_forecast as ops  # noqa: E402

LAB = "F-LAB"
STD = "F-STD"
WINDOW = ("2026-01", "2026-03")
DIRECT_MONTH = 25 * 40 * 52 / 12          # 4,333.33
INDIRECT_MONTH = 5000.0
ADMIN_MONTH = 6000.0
UNCLASSIFIED_MONTH = 3000.0
STANDARD = {"2026-01": 3000.0, "2026-02": 6000.0, "2026-03": 0.0}
STANDARD_HOURS = {"2026-01": 100.0, "2026-02": 200.0, "2026-03": 0.0}
PAID_HOURS = 40 * 52 / 12                  # 173.33 per month for a full-time direct employee


def _seed(path: str) -> None:
    conn = sqlite3.connect(path)
    c = conn.cursor()
    c.execute("INSERT INTO forecast (forecast_id, name, is_base) VALUES (?, 'Labor classes', 1)", (LAB,))
    c.execute("INSERT INTO forecast (forecast_id, name, is_base) VALUES (?, 'Standard only', 0)", (STD,))
    c.execute("INSERT INTO customers (customer_id, customer_name) VALUES ('C1', 'Acme')")
    c.execute("""INSERT INTO units (unit_id, unit_name, base_price, bom_id, bom_version, router_id, router_version)
                 VALUES ('P1', 'Widget', 100, 'BOM1', '1.0', 'R1', '1.0')""")
    c.execute("INSERT INTO bom_definitions (bom_id, bom_name, version) VALUES ('BOM1', 'Widget BOM', '1.0')")
    c.execute("""INSERT INTO bom (bom_id, version, bom_line, material_description, qty, unit, unit_price, material_cost, target_cost)
                 VALUES ('BOM1', '1.0', 1, 'Steel', 2, 'ea', 10, 20, 20)""")
    c.execute("INSERT INTO router_definitions (router_id, router_name, version) VALUES ('R1', 'Widget Router', '1.0')")
    c.execute("INSERT INTO labor_rates (rate_id, rate_name, rate_amount, rate_type) VALUES ('L1', 'Assembler', 30, 'hourly')")
    c.execute("INSERT INTO machines (machine_id, machine_name, machine_rate) VALUES ('BENCH', 'Bench', 0)")
    c.execute("""INSERT INTO router_operations (router_id, sequence, machine_id, machine_minutes, labor_minutes, labor_type_id, operation_description)
                 VALUES ('R1', 1, 'BENCH', 0, 60, 'L1', 'Assemble')""")
    n = 0
    for fid in (LAB, STD):
        for period, qty in (("2026-01", 100), ("2026-02", 200)):
            n += 1
            c.execute("""INSERT INTO sales (sale_id, customer_id, unit_id, period, quantity, unit_price, total_revenue, forecast_id)
                         VALUES (?, 'C1', 'P1', ?, ?, 100, ?, ?)""", (f"S{n}", period, qty, qty * 100, fid))
    c.execute("""INSERT INTO payroll_config (config_id, federal_tax_rate, state_tax_rate, social_security_rate, medicare_rate,
                 unemployment_rate, benefits_rate, workers_comp_rate, pay_frequency)
                 VALUES ('CFG', 0, 0, 0, 0, 0, 0, 0, 'monthly')""")
    employees = [
        ("D1", "Dana Direct", "Production", 40, 25, "hourly", "direct", LAB),
        ("I1", "Ivan Indirect", "Maintenance", 40, 60000, "salary", "indirect", LAB),
        ("A1", "Alex Admin", "Finance", 40, 72000, "salary", "admin", LAB),
        ("U1", "Uma Unclassified", "Sales", 40, 36000, "salary", None, LAB),
        ("S1", "Sam Shopfloor", "Production", 40, 25, "hourly", None, STD),
    ]
    for emp_id, name, dept, hours, rate, rate_type, labor_class, fid in employees:
        c.execute("""INSERT INTO payroll (employee_id, employee_name, department, weekly_hours, hourly_rate, rate_type,
                     labor_type, start_date, benefits_eligible, labor_class, forecast_id)
                     VALUES (?, ?, ?, ?, ?, ?, ?, '2025-01-01', 0, ?, ?)""",
                  (emp_id, name, dept, hours, rate, rate_type, dept, labor_class, fid))
    c.execute("""INSERT INTO financial_assumptions (forecast_id, opening_period, opening_cash, dso_days, dio_days, dpo_days,
                 income_tax_rate) VALUES (?, '2025-12', 50000, 30, 30, 30, 0.25)""", (LAB,))
    c.execute("""INSERT INTO financial_assumptions (forecast_id, opening_period, opening_cash, dso_days, dio_days, dpo_days,
                 income_tax_rate) VALUES (?, '2025-12', 50000, 30, 30, 30, 0.25)""", (STD,))
    conn.commit()
    conn.close()


@pytest.fixture(scope="module")
def db():
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, "labor.db")
    mgr = DatabaseManager(database_path=path, data_dir=tmp)
    mgr.initialize()
    _seed(path)
    yield mgr
    mgr.close_all_connections()
    shutil.rmtree(tmp, ignore_errors=True)


def _run(mgr, fn, *args, **kwargs):
    conn = mgr.get_connection()
    try:
        return fn(conn.cursor(), *args, **kwargs)
    finally:
        conn.close()


@pytest.fixture(scope="module")
def lab(db):
    return _run(db, fs.build_statements, [LAB], *WINDOW)


@pytest.fixture(scope="module")
def std(db):
    return _run(db, fs.build_statements, [STD], *WINDOW)


def _rows(stmt, section="income_statement"):
    return {r["period"]: r for r in stmt[section]["by_period"]}


# ---------------------------------------------------------------------------
# payroll_by_period: totals by class
# ---------------------------------------------------------------------------

def test_normalize_labor_class():
    assert fs.normalize_labor_class("Direct") == "direct"
    assert fs.normalize_labor_class(" indirect ") == "indirect"
    assert fs.normalize_labor_class("ADMIN") == "admin"
    assert fs.normalize_labor_class("direct labor") == "direct"
    assert fs.normalize_labor_class("") is None
    assert fs.normalize_labor_class(None) is None
    with pytest.raises(ValueError):
        fs.normalize_labor_class("supervisor")
    assert fs.normalize_labor_class("supervisor", strict=False) is None


def test_payroll_by_period_returns_totals_by_class(db):
    result = _run(db, fs.payroll_by_period, [LAB], list(STANDARD))
    by_class = result["by_labor_class"]
    for p in STANDARD:
        assert by_class["direct"][p] == pytest.approx(DIRECT_MONTH)
        assert by_class["indirect"][p] == pytest.approx(INDIRECT_MONTH)
        assert by_class["admin"][p] == pytest.approx(ADMIN_MONTH)
        assert by_class["unclassified"][p] == pytest.approx(UNCLASSIFIED_MONTH)
        assert sum(by_class[k][p] for k in by_class) == pytest.approx(result["total"][p])
        assert result["paid_hours_by_class"]["direct"][p] == pytest.approx(PAID_HOURS)
    assert [e["employee_id"] for e in result["unclassified_employees"]] == ["U1"]
    classes = {e["employee_id"]: e["labor_class"] for e in result["employees"]}
    assert classes == {"D1": "direct", "I1": "indirect", "A1": "admin", "U1": None}


# ---------------------------------------------------------------------------
# Statements: no double count
# ---------------------------------------------------------------------------

def test_direct_payroll_replaces_the_routing_standard(lab):
    rows = _rows(lab)
    for p, standard in STANDARD.items():
        r = rows[p]
        assert r["cogs_labor"] == pytest.approx(standard, abs=0.01)                 # routing standard, as before
        assert r["labor_variance"] == pytest.approx(DIRECT_MONTH - standard, abs=0.01)
        assert r["cogs_labor"] + r["labor_variance"] == pytest.approx(DIRECT_MONTH, abs=0.01)
        assert r["indirect_labor"] == pytest.approx(INDIRECT_MONTH, abs=0.01)
        assert r["payroll"] == pytest.approx(ADMIN_MONTH + UNCLASSIFIED_MONTH, abs=0.01)
        assert r["total_cogs"] == pytest.approx(
            r["cogs_materials"] + r["cogs_labor"] + r["labor_variance"] + r["cogs_manufacturing"]
            + r["indirect_labor"] + r["factory_overhead"] + r["other_cogs"], abs=0.01)
        assert r["gross_profit"] == pytest.approx(r["revenue"] - r["total_cogs"], abs=0.01)


def test_every_payroll_dollar_is_counted_exactly_once(lab):
    """Reviewer's evidence: labor-related cost was payroll + routing standard. Now it equals payroll."""
    total_payroll = DIRECT_MONTH + INDIRECT_MONTH + ADMIN_MONTH + UNCLASSIFIED_MONTH
    for p, r in _rows(lab).items():
        labor_related = r["cogs_labor"] + r["labor_variance"] + r["indirect_labor"] + r["payroll"]
        assert labor_related == pytest.approx(total_payroll, abs=0.01), p


def test_without_direct_employees_the_statement_is_unchanged(std):
    for p, r in _rows(std).items():
        assert r["cogs_labor"] == pytest.approx(STANDARD[p], abs=0.01)
        assert r["labor_variance"] == 0
        assert r["indirect_labor"] == 0
        assert r["payroll"] == pytest.approx(DIRECT_MONTH, abs=0.01)  # the unclassified production worker: opex, as before


def test_inventory_stays_on_standard_product_cogs(lab, std):
    lab_bs, std_bs = _rows(lab, "balance_sheet"), _rows(std, "balance_sheet")
    for p in STANDARD:
        assert lab_bs[p]["inventory"] == pytest.approx(std_bs[p]["inventory"], abs=0.01)
        product_cogs = _rows(lab)[p]["cogs_materials"] + _rows(lab)[p]["cogs_labor"] + _rows(lab)[p]["cogs_manufacturing"]
        assert lab_bs[p]["inventory"] == pytest.approx(product_cogs * 30 / fs.AVG_DAYS_PER_MONTH, abs=0.01)


def test_balance_check_and_cash_roll_forward_stay_exact(lab):
    for r in lab["balance_sheet"]["by_period"]:
        assert abs(r["balance_check"]) < 0.01
    assert abs(lab["cash_flow"]["reconciliation"]["difference"]) < 0.01
    for r in lab["cash_flow"]["by_period"]:
        assert r["closing_cash"] == pytest.approx(r["opening_cash"] + r["net_cash_flow"], abs=0.01)


def test_totals_and_detail_carry_the_labor_policy(lab):
    totals = lab["income_statement"]["totals"]
    assert totals["labor_variance"] == pytest.approx(sum(DIRECT_MONTH - s for s in STANDARD.values()), abs=0.01)
    assert totals["indirect_labor"] == pytest.approx(3 * INDIRECT_MONTH, abs=0.01)
    policy = lab["detail"]["labor_policy"]
    assert policy["direct_payroll_in_cogs"] is True
    assert [e["employee_name"] for e in policy["unclassified_employees"]] == ["Uma Unclassified"]
    by_class = lab["detail"]["payroll_by_labor_class"]
    assert by_class["direct"]["2026-02"] == pytest.approx(DIRECT_MONTH, abs=0.01)
    emp = {e["employee_id"]: e for e in lab["detail"]["payroll_employees"]}
    assert emp["D1"]["labor_class"] == "direct"


def test_unclassified_staff_with_routing_labor_raise_an_info_warning(std, lab):
    warning = next(w for w in std["warnings"] if w["code"] == "unclassified_payroll")
    assert warning["level"] == "info"
    assert "Sam Shopfloor" in warning["message"]
    assert "Direct" in warning["message"]
    # Classified scenario: one admin-ish unclassified employee still listed, but as info.
    assert any(w["code"] == "unclassified_payroll" for w in lab["warnings"])


def test_carry_forward_window_matches_full_year(db):
    """A March-only report shows the same March lines as a Jan-Mar report."""
    march_only = _run(db, fs.build_statements, [LAB], "2026-03", "2026-03")
    full = _run(db, fs.build_statements, [LAB], *WINDOW)
    a, b = march_only["income_statement"]["by_period"][0], _rows(full)["2026-03"]
    for key in ("cogs_labor", "labor_variance", "indirect_labor", "payroll", "total_cogs", "net_income"):
        assert a[key] == pytest.approx(b[key], abs=0.01), key
    assert march_only["balance_sheet"]["closing"]["cash"] == pytest.approx(full["balance_sheet"]["closing"]["cash"], abs=0.01)


# ---------------------------------------------------------------------------
# Budget vs actual lines
# ---------------------------------------------------------------------------

def test_variance_tracks_the_new_lines(db):
    assert "indirect_labor" in fs.VARIANCE_LINE_ITEMS and "labor_variance" in fs.VARIANCE_LINE_ITEMS
    keys = list(fs.VARIANCE_LINE_ITEMS)
    assert keys.index("cogs_labor") < keys.index("labor_variance") < keys.index("total_cogs")
    assert keys.index("indirect_labor") < keys.index("total_cogs")
    variance = _run(db, fs.build_variance, [LAB], *WINDOW)
    line = next(item for item in variance["line_items"] if item["line_item"] == "labor_variance")
    assert line["by_period"][0]["budget"] == pytest.approx(DIRECT_MONTH - 3000, abs=0.01)
    assert line["direction"] == "expense"


# ---------------------------------------------------------------------------
# COGS forecast and labor forecast tie to the statement
# ---------------------------------------------------------------------------

def test_cogs_forecast_summary_ties_to_the_income_statement(db, lab):
    cogs = _run(db, ops.build_cogs_forecast, [LAB], *WINDOW)
    inc = _rows(lab)
    for row in cogs["summary"]["by_period"]:
        r = inc[row["period"]]
        assert row["indirect_labor"] == pytest.approx(r["indirect_labor"], abs=0.01)
        assert row["labor_variance"] == pytest.approx(r["labor_variance"], abs=0.01)
        assert row["total_cogs"] == pytest.approx(r["total_cogs"], abs=0.01)
        assert row["gross_profit"] == pytest.approx(r["gross_profit"], abs=0.01)
    totals = cogs["summary"]["totals"]
    assert totals["total_cogs"] == pytest.approx(lab["income_statement"]["totals"]["total_cogs"], abs=0.01)
    assert totals["labor_variance"] == pytest.approx(lab["income_statement"]["totals"]["labor_variance"], abs=0.01)


def test_labor_forecast_shows_standard_vs_payroll(db):
    labor = _run(db, ops.build_labor_forecast, [LAB], *WINDOW)
    comparison = {r["period"]: r for r in labor["standard_vs_payroll"]["by_period"]}
    for p in STANDARD:
        row = comparison[p]
        assert row["standard_hours"] == pytest.approx(STANDARD_HOURS[p], abs=0.01)
        assert row["standard_cost"] == pytest.approx(STANDARD[p], abs=0.01)
        assert row["direct_payroll_cost"] == pytest.approx(DIRECT_MONTH, abs=0.01)
        assert row["paid_hours"] == pytest.approx(PAID_HOURS, abs=0.01)
        assert row["labor_variance"] == pytest.approx(DIRECT_MONTH - STANDARD[p], abs=0.01)
        assert row["indirect_payroll_cost"] == pytest.approx(INDIRECT_MONTH, abs=0.01)
    assert comparison["2026-01"]["utilization_pct"] == pytest.approx(100 / PAID_HOURS * 100, abs=0.01)
    assert comparison["2026-03"]["utilization_pct"] == 0
    totals = labor["standard_vs_payroll"]["totals"]
    assert totals["paid_hours"] == pytest.approx(3 * PAID_HOURS, abs=0.01)
    assert totals["utilization_pct"] == pytest.approx(300 / (3 * PAID_HOURS) * 100, abs=0.01)
    assert labor["totals"]["direct_payroll_cost"] == pytest.approx(3 * DIRECT_MONTH, abs=0.01)
    note = labor["payroll_note"]
    assert "keep direct workers out of Payroll" not in note["message"]   # the old generic double-count note
    assert "equals the pay of employees classed Direct" in note["message"]
    assert "Uma Unclassified" in note["message"]
    assert note["direct_payroll_cost"] == pytest.approx(3 * DIRECT_MONTH, abs=0.01)
    assert note["direct_labor_cost"] == pytest.approx(sum(STANDARD.values()), abs=0.01)
    assert [e["employee_name"] for e in note["unclassified_employees"]] == ["Uma Unclassified"]


def test_labor_forecast_note_without_direct_class_names_the_risk(db):
    labor = _run(db, ops.build_labor_forecast, [STD], *WINDOW)
    note = labor["payroll_note"]
    assert note["direct_payroll_cost"] == 0
    assert "Sam Shopfloor" in note["message"]
    assert "Direct" in note["message"]
    rows = labor["standard_vs_payroll"]["by_period"]
    assert all(r["labor_variance"] == 0 and r["utilization_pct"] is None for r in rows)


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def _sheet_labels(ws):
    return [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]


def test_workbook_shows_the_new_lines_and_explains_the_policy(db):
    from services.workbook_export import workbook_bytes
    data = _run(db, workbook_bytes, [LAB], *WINDOW)
    wb = load_workbook(io.BytesIO(data))
    income = _sheet_labels(wb["Income Statement"])
    assert "Indirect labor (factory)" in income
    assert "Direct labor variance (payroll vs standard)" in income
    assert income.index("Direct labor (standard)") < income.index("Direct labor variance (payroll vs standard)")
    margin = _sheet_labels(wb["Gross Margin"])
    assert "Indirect labor (factory)" in margin and "Direct labor variance (payroll vs standard)" in margin
    readme = " ".join(str(v) for row in wb["Read Me"].iter_rows(values_only=True) for v in row if v)
    assert "Labor classes" in readme
    assert "classed Direct" in readme
    labor_sheet = " ".join(str(v) for row in wb["Labor Forecast"].iter_rows(values_only=True) for v in row if v)
    assert "Direct payroll" in labor_sheet and "Utilization" in labor_sheet
