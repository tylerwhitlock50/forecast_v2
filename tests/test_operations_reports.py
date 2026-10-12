"""
COGS / labor forecast, sales detail and Excel workbook endpoints.

Seeds an isolated database with three customers of two types, a rifle and a
pistol with BOM + routing, an accessory with no BOM or routing (a cost
warning), factory overhead, price-list entries, and deduction rules covering
terms (with a customer override), chargebacks, a per-unit allowance,
embedded 11% excise and a commission. Then checks, through the HTTP routes,
that every report ties to /reporting/statements and that the workbook opens,
has every sheet, holds real numbers and agrees with the statements.
"""
import csv
import io
import math
import os
import shutil
import sqlite3
import sys
import tempfile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

_APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, "app"))
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from app.db.database import DatabaseManager  # noqa: E402
from services import workbook_export as wx  # noqa: E402
from services.operations_forecast import DEFAULT_HOURS_PER_FTE  # noqa: E402

F = "F-OPS"
START, END = "2027-01", "2027-03"
PERIODS = ["2027-01", "2027-02", "2027-03"]
WINDOW = {"forecast_ids": F, "start_period": START, "end_period": END}

EXPECTED_SHEETS = [
    "Read Me", "Gross Margin", "Income Statement", "By Product", "By Customer", "By Customer Type",
    "Labor Forecast", "Machine Hours", "Unit Costs", "Sales Detail", "Deduction Rules", "Price List",
    "Cash Flow", "Balance Sheet",
]


def _seed(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    c.execute("INSERT INTO forecast (forecast_id, name, is_base) VALUES (?, 'Budget 2027', 1)", (F,))
    c.execute("INSERT INTO forecast (forecast_id, name) VALUES ('F-ALT', 'Alt')")
    c.executemany("INSERT INTO customers (customer_id, customer_name, customer_type) VALUES (?, ?, ?)", [
        ("C-ACME", "Acme Outdoors", "Dealer"),
        ("C-DLR", "Main Street Guns", "Dealer"),
        ("C-DST", "Big Distributor", "Distributor"),
    ])
    c.executemany("""INSERT INTO units (unit_id, unit_name, base_price, unit_type, bom_id, bom_version, router_id)
                     VALUES (?, ?, ?, ?, ?, '1.0', ?)""", [
        ("RIF", "Hunter Rifle", 1500, "Rifle", "B-RIF", "R-RIF"),
        ("PST", "Compact Pistol", 600, "Pistol", "B-PST", "R-PST"),
        ("ACC", "Sling", 30, "Accessory", None, None),   # no BOM / routing -> cost warning
    ])
    c.executemany("""INSERT INTO bom (bom_id, version, bom_line, material_description, qty, unit, unit_price, material_cost)
                     VALUES (?, '1.0', ?, ?, ?, 'ea', ?, ?)""", [
        ("B-RIF", 1, "Barrel", 1, 210.5, 210.5),
        ("B-RIF", 2, "Stock", 1, 95.25, 95.25),
        ("B-PST", 1, "Slide", 2, 41.1, 82.2),
    ])
    c.executemany("INSERT INTO machines (machine_id, machine_name, machine_rate, available_minutes_per_month) VALUES (?, ?, ?, ?)", [
        ("CNC", "CNC Mill", 115, 900),   # 15 hours available a month
        ("ASM", "Assembly Bench", 10, 6000),
    ])
    c.executemany("INSERT INTO labor_rates (rate_id, rate_name, rate_amount) VALUES (?, ?, ?)", [
        ("MACH", "Machinist", 37.5), ("ASSY", "Assembler", 23.75),
    ])
    c.executemany("""INSERT INTO router_operations (router_id, sequence, machine_id, machine_minutes, labor_minutes, labor_type_id)
                     VALUES (?, ?, ?, ?, ?, ?)""", [
        ("R-RIF", 10, "CNC", 35, 25, "MACH"),
        ("R-RIF", 20, "ASM", 20, 50, "ASSY"),
        ("R-PST", 10, "CNC", 17, 22, "MACH"),
    ])
    sales = [
        # id, customer, unit, period, qty, price
        ("S1", "C-ACME", "RIF", "2027-01", 10, 1110.0),
        ("S2", "C-DST", "RIF", "2027-01", 20, 999.0),
        ("S3", "C-DLR", "RIF", "2027-02", 7, 1050.5),
        ("S4", "C-DST", "PST", "2027-02", 12.5, 420.0),
        ("S5", "C-DLR", "ACC", "2027-03", 40, 25.0),
        ("S6", "C-DST", "RIF", "2027-03", 15, 999.0),
        ("S7", "C-DST", "RIF", "2027-04", 50, 999.0),   # outside the window
    ]
    for sid, cust, unit, period, qty, price in sales:
        c.execute("""INSERT INTO sales (sale_id, customer_id, unit_id, period, quantity, unit_price, total_revenue, forecast_id)
                     VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", (sid, cust, unit, period, qty, price, qty * price, F))
    c.execute("""INSERT INTO sales (sale_id, customer_id, unit_id, period, quantity, unit_price, total_revenue, forecast_id)
                 VALUES ('SX', 'C-DLR', 'RIF', '2027-01', 99, 1000, 99000, 'F-ALT')""")

    rules = [
        # rule_id, forecast_id, type, method, rate, per_unit, customer_id, customer_type, unit_id, unit_type, start
        ("DED-TERMS", None, "terms_allowance", "percent_of_gross", 0.02, 0, None, None, None, None, None),
        ("DED-ACME", None, "terms_allowance", "percent_of_gross", 0.01, 0, "C-ACME", None, None, None, None),
        ("DED-CB", None, "chargebacks", "percent_of_gross", 0.015, 0, None, "Dealer", None, None, None),
        ("DED-RIFLE", F, "other_allowance", "per_unit", 0, 5.0, None, None, "RIF", None, "2027-02"),
        ("DED-FET", None, "excise_tax", "embedded_tax", 0.11, 0, None, None, None, "Rifle", None),
        ("DED-COM", None, "commission", "percent_of_gross", 0.03, 0, None, "Distributor", None, None, None),
        ("DED-ALT", "F-ALT", "chargebacks", "percent_of_gross", 0.5, 0, None, None, None, None, None),
    ]
    for r in rules:
        c.execute("""INSERT INTO sales_deduction_rules (rule_id, forecast_id, deduction_type, method, rate, amount_per_unit,
                     customer_id, customer_type, unit_id, unit_type, start_period, is_active, name)
                     VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?)""", (*r, r[0].title()))
    # Inactive rule whose imported text looks like a formula and carries a control character.
    c.execute("""INSERT INTO sales_deduction_rules (rule_id, deduction_type, method, rate, is_active, name, notes)
                 VALUES ('DED-OFF', 'chargebacks', 'percent_of_gross', 0.9, 0, ?, ?)""",
              ("=HYPERLINK(\"http://example.com\",\"x\")", "bell\x07 note"))
    c.executemany("""INSERT INTO price_list (price_id, forecast_id, unit_id, customer_id, customer_type, method, unit_price, discount_pct)
                     VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", [
        ("PRC-DST-RIF", None, "RIF", None, "Distributor", "fixed", 999.0, None),
        ("PRC-DLR-ALL", None, None, None, "Dealer", "discount_from_base", None, 0.30),
        ("PRC-ACME-RIF", F, "RIF", "C-ACME", None, "fixed", 1050.0, None),
    ])

    c.execute("INSERT INTO expense_categories (category_id, category_name, category_type) VALUES ('FO', 'Plant rent', 'factory_overhead')")
    c.execute("""INSERT INTO expenses (expense_id, expense_name, category_id, amount, frequency, start_date, is_active, forecast_id)
                 VALUES ('X1', 'Plant rent', 'FO', 500, 'monthly', '2027-01-01', 1, NULL)""")
    for i, p in enumerate(PERIODS):
        c.execute("""INSERT INTO expense_allocations (allocation_id, expense_id, period, allocated_amount, allocation_type)
                     VALUES (?, 'X1', ?, 500, 'scheduled')""", (f"A{i}", p))
    c.execute("""INSERT INTO financial_assumptions (forecast_id, opening_period, opening_cash, dso_days, dio_days, dpo_days,
                 income_tax_rate, excise_payable_days) VALUES (?, '2026-12', 50000, 30, 45, 30, 0.21, 15)""", (F,))
    conn.commit()
    conn.close()


@pytest.fixture(scope="module")
def report_db():
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, "ops.db")
    mgr = DatabaseManager(database_path=path, data_dir=tmp)
    mgr.initialize()
    _seed(path)
    yield mgr
    shutil.rmtree(tmp, ignore_errors=True)


@pytest.fixture(scope="module")
def client(report_db):
    import app.db.database as dbmod
    import db.database as dbmod_short
    from api import reporting_routes
    saved = (dbmod.db_manager, dbmod_short.db_manager, reporting_routes.db_manager)
    dbmod.db_manager = report_db
    dbmod_short.db_manager = report_db
    reporting_routes.db_manager = report_db
    app = FastAPI()
    app.include_router(reporting_routes.router)
    yield TestClient(app)
    dbmod.db_manager, dbmod_short.db_manager, reporting_routes.db_manager = saved


def _get(client, path, **params):
    r = client.get(path, params={**WINDOW, **params})
    assert r.status_code == 200, r.text
    return r.json()["data"]


@pytest.fixture(scope="module")
def statements(client):
    return _get(client, "/reporting/statements")


@pytest.fixture(scope="module")
def cogs(client):
    return _get(client, "/reporting/cogs-forecast")


@pytest.fixture(scope="module")
def labor(client):
    return _get(client, "/reporting/labor-forecast")


@pytest.fixture(scope="module")
def detail(client):
    return _get(client, "/reporting/sales-detail")


@pytest.fixture(scope="module")
def workbook_response(client):
    r = client.get("/reporting/workbook", params=WINDOW)
    assert r.status_code == 200, r.text
    return r


@pytest.fixture(scope="module")
def wb(workbook_response):
    return load_workbook(io.BytesIO(workbook_response.content))


def _income(statements, period):
    return next(r for r in statements["income_statement"]["by_period"] if r["period"] == period)


def _rows_by_label(ws):
    """{column A label: [values in the other columns]} for a statement-style sheet."""
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None and row[0] not in out:
            out[row[0]] = list(row[1:])
    return out


# ---------------------------------------------------------------------------
# COGS forecast
# ---------------------------------------------------------------------------

class TestCogsForecastRoute:
    PAIRS = [  # (cogs summary key, income statement key)
        ("gross_sales", "gross_sales"), ("terms_allowance", "terms_allowance"), ("chargebacks", "chargebacks"),
        ("other_allowances", "other_allowances"), ("excise_tax", "excise_tax"),
        ("total_sales_deductions", "total_sales_deductions"), ("net_sales", "revenue"),
        ("materials", "cogs_materials"), ("labor", "cogs_labor"), ("machine", "cogs_manufacturing"),
        ("factory_overhead", "factory_overhead"), ("total_cogs", "total_cogs"), ("gross_profit", "gross_profit"),
        ("units", "units_sold"), ("commissions", "commissions"),
    ]

    def test_monthly_summary_ties_to_statements(self, cogs, statements):
        assert [r["period"] for r in cogs["summary"]["by_period"]] == PERIODS
        for row in cogs["summary"]["by_period"]:
            inc = _income(statements, row["period"])
            for key, stmt_key in self.PAIRS:
                assert row[key] == pytest.approx(inc[stmt_key], abs=0.01), (row["period"], key)

    def test_totals_tie_to_statements(self, cogs, statements):
        totals = cogs["summary"]["totals"]
        inc = statements["income_statement"]["totals"]
        for key, stmt_key in self.PAIRS:
            assert totals[key] == pytest.approx(inc[stmt_key], abs=0.02), key
        assert totals["gross_margin_pct"] == pytest.approx(inc["gross_margin_pct"], abs=0.01)

    def test_deductions_are_applied(self, cogs):
        totals = cogs["summary"]["totals"]
        # Rifles carry 11% embedded FET: rifle gross x 0.11 / 1.11.
        rifle_gross = 10 * 1110 + 20 * 999 + 7 * 1050.5 + 15 * 999
        assert totals["excise_tax"] == pytest.approx(rifle_gross * 0.11 / 1.11, abs=0.01)
        assert totals["net_sales"] == pytest.approx(totals["gross_sales"] - totals["total_sales_deductions"], abs=0.01)
        assert totals["total_sales_deductions"] > 0

    def test_dimensions_are_contribution_before_overhead(self, cogs):
        summary = cogs["summary"]["totals"]
        for dim in ("by_product", "by_customer", "by_customer_type", "by_product_type"):
            assert sum(g["totals"]["contribution_margin"] for g in cogs[dim]) == pytest.approx(
                summary["contribution_margin"], abs=0.05), dim
        assert summary["gross_profit"] == pytest.approx(
            summary["contribution_margin"] - summary["factory_overhead"] - summary["other_cogs"], abs=0.02)
        assert summary["factory_overhead"] == pytest.approx(1500)

    def test_comma_separated_ids_and_other_scenario_excluded(self, client):
        data = _get(client, "/reporting/cogs-forecast", forecast_ids=f"{F},F-ALT")
        assert data["forecast_ids"] == [F, "F-ALT"]
        alone = _get(client, "/reporting/cogs-forecast")
        assert data["summary"]["totals"]["gross_sales"] == pytest.approx(alone["summary"]["totals"]["gross_sales"] + 99000)

    @pytest.mark.parametrize("path", ["/reporting/cogs-forecast", "/reporting/labor-forecast",
                                      "/reporting/sales-detail", "/reporting/workbook"])
    def test_bad_window_and_missing_ids_are_422(self, client, path):
        assert client.get(path, params={**WINDOW, "start_period": "2027-03", "end_period": "2027-01"}).status_code == 422
        assert client.get(path, params={**WINDOW, "start_period": "2027-13"}).status_code == 422
        assert client.get(path, params={"start_period": START, "end_period": END}).status_code == 422
        assert client.get(path, params={**WINDOW, "forecast_ids": " , "}).status_code == 422


# ---------------------------------------------------------------------------
# Labor forecast
# ---------------------------------------------------------------------------

class TestLaborForecastRoute:
    def test_default_hours_per_fte(self, labor):
        assert labor["hours_per_fte"] == DEFAULT_HOURS_PER_FTE

    def test_cost_ties_to_statement_direct_labor(self, labor, statements):
        assert labor["totals"]["cost"] == pytest.approx(statements["income_statement"]["totals"]["cogs_labor"], abs=0.02)
        for row in labor["totals_by_period"]:
            assert row["cost"] == pytest.approx(_income(statements, row["period"])["cogs_labor"], abs=0.02)

    def test_machine_cost_ties_to_statement(self, labor, statements):
        machine_cost = sum(m["totals"]["cost"] for m in labor["by_machine"])
        assert machine_cost == pytest.approx(statements["income_statement"]["totals"]["cogs_manufacturing"], abs=0.02)

    def test_hours_and_fte(self, client):
        data = _get(client, "/reporting/labor-forecast", hours_per_fte=160)
        mach = next(t for t in data["by_labor_type"] if t["labor_type_id"] == "MACH")
        # Jan: 30 rifles x 25 min = 750 min = 12.5 h
        assert mach["by_period"]["2027-01"]["hours"] == pytest.approx(12.5)
        assert mach["by_period"]["2027-01"]["fte"] == pytest.approx(12.5 / 160, abs=0.001)
        assert data["hours_per_fte"] == 160

    @pytest.mark.parametrize("value", ["0", "-5", "nan", "inf", "abc"])
    def test_bad_hours_per_fte_is_422(self, client, value):
        r = client.get("/reporting/labor-forecast", params={**WINDOW, "hours_per_fte": value})
        assert r.status_code == 422
        r = client.get("/reporting/workbook", params={**WINDOW, "hours_per_fte": value})
        assert r.status_code == 422


# ---------------------------------------------------------------------------
# Sales detail
# ---------------------------------------------------------------------------

class TestSalesDetailRoute:
    def test_rows_are_the_window_lines(self, detail):
        assert detail["total_count"] == 6  # S7 is outside the window, SX is another scenario
        assert detail["returned"] == 6 and detail["has_more"] is False
        assert {r["sale_id"] for r in detail["rows"]} == {"S1", "S2", "S3", "S4", "S5", "S6"}
        for row in detail["rows"]:
            assert not any(isinstance(v, (dict, list)) for v in row.values())

    def test_totals_tie_to_statements(self, detail, statements):
        inc = statements["income_statement"]["totals"]
        totals = detail["totals"]
        assert totals["gross_sales"] == pytest.approx(inc["gross_sales"], abs=0.02)
        assert totals["net_sales"] == pytest.approx(inc["revenue"], abs=0.02)
        assert totals["excise_tax"] == pytest.approx(inc["excise_tax"], abs=0.02)
        assert totals["total_sales_deductions"] == pytest.approx(inc["total_sales_deductions"], abs=0.02)
        assert totals["commissions"] == pytest.approx(inc["commissions"], abs=0.02)
        assert totals["product_cogs"] == pytest.approx(
            inc["cogs_materials"] + inc["cogs_labor"] + inc["cogs_manufacturing"], abs=0.02)
        assert totals["quantity"] == pytest.approx(inc["units_sold"])
        assert "unit_price" not in totals

    def test_pagination_keeps_totals_over_all_rows(self, client, detail):
        page = _get(client, "/reporting/sales-detail", limit=2, offset=3)
        assert page["total_count"] == 6 and page["returned"] == 2 and page["has_more"] is True
        assert page["rows"] == detail["rows"][3:5]
        assert page["totals"] == detail["totals"]

    def test_columns_metadata(self, detail):
        columns = {c["key"]: c["label"] for c in detail["columns"]}
        assert columns["net_sales"] == "Net sales"
        assert columns["excise_tax"] == "Excise tax (FET)"

    @pytest.mark.parametrize("params", [{"limit": 0}, {"limit": 20001}, {"offset": -1}])
    def test_bad_paging_is_422(self, client, params):
        assert client.get("/reporting/sales-detail", params={**WINDOW, **params}).status_code == 422


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

class TestWorkbookRoute:
    def test_download_headers(self, workbook_response):
        assert workbook_response.headers["content-type"] == wx.XLSX_MEDIA_TYPE
        assert 'filename="forecast_F-OPS_2027-01_2027-03.xlsx"' in workbook_response.headers["content-disposition"]
        assert workbook_response.headers["content-disposition"].startswith("attachment")

    def test_has_every_sheet_in_order(self, wb):
        assert wb.sheetnames == EXPECTED_SHEETS
        assert wx.SHEET_NAMES == EXPECTED_SHEETS

    def test_gross_margin_ties_to_statements(self, wb, statements):
        ws = wb["Gross Margin"]
        header = [c.value for c in ws[1]]
        assert header == ["Gross margin bridge"] + PERIODS + ["Total"]
        rows = _rows_by_label(ws)
        inc_totals = statements["income_statement"]["totals"]
        assert rows["Net sales"][-1] == pytest.approx(inc_totals["revenue"], abs=0.01)
        assert rows["Gross sales"][-1] == pytest.approx(inc_totals["gross_sales"], abs=0.01)
        assert rows["Excise tax (FET)"][-1] == pytest.approx(inc_totals["excise_tax"], abs=0.01)
        assert rows["Gross profit"][-1] == pytest.approx(inc_totals["gross_profit"], abs=0.01)
        for i, period in enumerate(PERIODS):
            assert rows["Net sales"][i] == pytest.approx(_income(statements, period)["revenue"], abs=0.01)
        # Percentages are stored as fractions with a percent format.
        assert rows["Gross margin %"][-1] == pytest.approx(inc_totals["gross_margin_pct"] / 100, abs=0.0001)
        pct_row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "Gross margin %")
        assert pct_row[-1].number_format == "0.0%"

    def test_income_statement_labels_and_order(self, wb, statements):
        ws = wb["Income Statement"]
        labels = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
        for label in ("Gross sales", "Terms allowance", "Chargebacks allowance", "Other discounts & allowances",
                      "Excise tax (FET)", "Total sales deductions", "Net sales", "Total COGS", "Gross profit",
                      "Sales commissions", "Net income"):
            assert label in labels
        assert labels.index("Gross sales") < labels.index("Excise tax (FET)") < labels.index("Net sales") \
            < labels.index("Total COGS") < labels.index("Gross profit") < labels.index("Sales commissions")
        rows = _rows_by_label(ws)
        assert rows["Net income"][-1] == pytest.approx(statements["income_statement"]["totals"]["net_income"], abs=0.01)

    @pytest.mark.parametrize("sheet", ["Gross Margin", "Income Statement", "Cash Flow", "Balance Sheet",
                                       "Labor Forecast", "Machine Hours"])
    def test_statement_sheets_hold_numbers(self, wb, sheet):
        ws = wb[sheet]
        numbers = 0
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                if cell.value is None:
                    continue
                if sheet in ("Labor Forecast",) and isinstance(cell.value, str):
                    continue  # the "Labor by product" block repeats a header row
                assert isinstance(cell.value, (int, float)), (sheet, cell.coordinate, cell.value)
                assert not isinstance(cell.value, bool)
                numbers += 1
        assert numbers > 0

    def test_dimension_sheets(self, wb, cogs):
        ws = wb["By Product"]
        header = [c.value for c in ws[1]]
        assert header[:3] == ["Product ID", "Product", "Product type"]
        assert "Net sales" in header and "Contribution margin" in header and "Units 2027-01" in header
        net_col = header.index("Net sales")
        body = list(ws.iter_rows(min_row=2, values_only=True))
        products = [r for r in body if r[0] in ("RIF", "PST", "ACC")]
        assert len(products) == 3
        total = next(r for r in body if r[0] == "Total")
        assert total[net_col] == pytest.approx(cogs["summary"]["totals"]["net_sales"], abs=0.01)
        assert sum(r[net_col] for r in products) == pytest.approx(total[net_col], abs=0.05)
        assert ws.freeze_panes == "D2"
        # Fractional units keep their decimals; whole units show none.
        units_col = header.index("Units") + 1
        pistol_row = next(i for i, r in enumerate(body, start=2) if r[0] == "PST")
        assert ws.cell(row=pistol_row, column=units_col).value == 12.5
        assert ws.cell(row=pistol_row, column=units_col).number_format == "#,##0.##"

        types = _rows_by_label(wb["By Customer Type"])
        assert set(types) >= {"Dealer", "Distributor", "Total"}
        customers = [r[1] for r in wb["By Customer"].iter_rows(min_row=2, values_only=True)]
        assert "Acme Outdoors" in customers

    def test_sales_detail_sheet(self, wb, detail, statements):
        ws = wb["Sales Detail"]
        header = [c.value for c in ws[1]]
        body = list(ws.iter_rows(min_row=2, values_only=True))
        lines = [r for r in body if r[0] in PERIODS]
        assert len(lines) == detail["total_count"]
        total = next(r for r in body if r[0] == "Total")
        assert total[header.index("Net sales")] == pytest.approx(statements["income_statement"]["totals"]["revenue"], abs=0.02)
        for r in lines:
            for col in ("Quantity", "Gross sales", "Excise tax (FET)", "Net sales", "Product COGS"):
                assert isinstance(r[header.index(col)], (int, float))
        assert ws.auto_filter.ref.startswith("A1:")

    def test_rule_and_price_sheets_join_names(self, wb):
        ws = wb["Deduction Rules"]
        header = [c.value for c in ws[1]]
        rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        assert "DED-ALT" not in rows  # another scenario's rule
        acme = rows["DED-ACME"]
        assert acme[header.index("Customer")] == "Acme Outdoors"
        assert acme[header.index("Deduction type")] == "Terms allowance"
        fet = rows["DED-FET"]
        assert fet[header.index("Rate")] == pytest.approx(0.11)
        assert fet[header.index("Treatment")].startswith("Sales deduction")
        assert rows["DED-COM"][header.index("Treatment")].startswith("Selling expense")
        assert rows["DED-RIFLE"][header.index("Product")] == "Hunter Rifle"
        assert rows["DED-RIFLE"][header.index("Scenario")] == "Budget 2027 (F-OPS)"
        rate_col = header.index("Rate") + 1
        fet_row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "DED-FET")
        assert fet_row[rate_col - 1].number_format == "0.00%"

        # Inactive rules are listed (marked inactive); formula-like text stays text; control chars are dropped.
        off = rows["DED-OFF"]
        assert off[header.index("Active")] == "No"
        assert off[header.index("Notes")] == "bell note"
        name_cell = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "DED-OFF")[header.index("Name")]
        assert name_cell.value.startswith("=HYPERLINK") and name_cell.data_type == "s"

        ws = wb["Price List"]
        header = [c.value for c in ws[1]]
        rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        assert rows["PRC-ACME-RIF"][header.index("Customer")] == "Acme Outdoors"
        assert rows["PRC-DST-RIF"][header.index("Product")] == "Hunter Rifle"
        assert rows["PRC-DLR-ALL"][header.index("Product")] == "All products"
        assert rows["PRC-DLR-ALL"][header.index("Discount off base")] == pytest.approx(0.30)
        assert rows["PRC-DST-RIF"][header.index("Resulting price")] == pytest.approx(999)

    def test_cash_flow_total_carries_balances(self, wb, statements):
        ws = wb["Cash Flow"]
        assert [c.value for c in ws[1]][-1] == "Total"
        rows = _rows_by_label(ws)
        totals = statements["cash_flow"]["totals"]
        assert rows["Opening cash"][-1] == pytest.approx(totals["opening_cash"], abs=0.01) == pytest.approx(50000)
        assert rows["Closing cash"][-1] == pytest.approx(totals["closing_cash"], abs=0.01)
        assert rows["Increase / (decrease) in excise tax payable"][-1] == pytest.approx(
            totals["change_in_excise_tax_payable"], abs=0.01)

    def test_balance_sheet_has_opening_and_no_total(self, wb, statements):
        ws = wb["Balance Sheet"]
        header = [c.value for c in ws[1]]
        assert header == ["Balance sheet", "Opening (2026-12)"] + PERIODS
        rows = _rows_by_label(ws)
        assert rows["Excise tax payable"][0] == 0
        assert rows["Excise tax payable"][-1] == pytest.approx(
            statements["balance_sheet"]["closing"]["excise_tax_payable"], abs=0.01)
        assert all(abs(v) < 0.01 for v in rows["Balance check (assets - liabilities - equity)"])

    def test_labor_and_machine_sheets(self, wb, labor):
        rows = _rows_by_label(wb["Labor Forecast"])
        assert rows["Machinist"][0] == 37.5  # rate column
        assert "Total direct labor" in rows
        ws = wb["Labor Forecast"]
        total_col = [c.value for c in ws[1]].index("Total")
        hours_total = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "Total direct labor")
        assert hours_total[total_col] == pytest.approx(labor["totals"]["hours"], abs=0.01)
        machine_rows = _rows_by_label(wb["Machine Hours"])
        assert "CNC Mill" in machine_rows

    def test_read_me_explains_and_warns(self, wb):
        ws = wb["Read Me"]
        text = "\n".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None)
        assert "Budget 2027 (F-OPS)" in text
        assert "2027-01 to 2027-03 (3 months)" in text
        assert "Which rule applies" in text
        assert "x 0.11 / 1.11" in text  # the embedded excise formula, via the worked example
        assert "embedded_tax" in text and "percent_of_net" in text
        assert "Contribution margin vs gross profit" in text
        assert "173.33" in text
        assert "ACC (Sling)" in text  # product cost warning expanded per product
        assert "CNC Mill needs more hours" in text  # machine capacity warning from the labor forecast

    def test_header_style_and_freeze(self, wb):
        for name in EXPECTED_SHEETS[1:]:
            ws = wb[name]
            assert ws.freeze_panes, name
            first = ws.cell(row=1, column=1)
            assert first.font.bold and first.fill.fgColor.rgb.endswith("1F3864"), name
        for name in EXPECTED_SHEETS:
            for letter, dim in wb[name].column_dimensions.items():
                if name != "Read Me" and dim.width:
                    assert dim.width <= 50, (name, letter)


# ---------------------------------------------------------------------------
# Service helpers and the statements CSV
# ---------------------------------------------------------------------------

class TestWorkbookService:
    def test_write_workbook(self, report_db, tmp_path):
        conn = report_db.get_connection()
        try:
            target = wx.export_path(str(tmp_path), [F], START, END)
            info = wx.write_workbook(target, conn.cursor(), [F], START, END)
        finally:
            conn.close()
        assert info["path"] == os.path.join(str(tmp_path), "exports", "forecast_F-OPS_2027-01_2027-03.xlsx")
        assert os.path.getsize(info["path"]) == info["bytes"] > 0
        assert info["sheets"] == EXPECTED_SHEETS

    def test_empty_window_still_builds(self, report_db):
        conn = report_db.get_connection()
        try:
            data = wx.workbook_bytes(conn.cursor(), [F], "2030-01", "2030-02")
        finally:
            conn.close()
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == EXPECTED_SHEETS
        assert "No sales" in str(wb["By Product"]["A2"].value)

    def test_filenames_are_safe(self, tmp_path):
        assert wx.workbook_filename(["F 1/2"], START, END) == "forecast_F_1_2_2027-01_2027-03.xlsx"
        path = wx.export_path(str(tmp_path), [F], START, END, filename="../../etc/my report")
        assert path == os.path.join(str(tmp_path), "exports", "my_report.xlsx")
        path = wx.export_path(str(tmp_path), [F], START, END, filename="Budget.XLSX")
        assert os.path.basename(path) == "Budget.xlsx"

    def test_bad_inputs_raise_value_error(self, report_db):
        conn = report_db.get_connection()
        try:
            with pytest.raises(ValueError):
                wx.build_workbook(conn.cursor(), [], START, END)
            with pytest.raises(ValueError):
                wx.build_workbook(conn.cursor(), [F], END, START)
            with pytest.raises(ValueError):
                wx.build_workbook(conn.cursor(), [F], START, END, hours_per_fte=0)
            with pytest.raises(ValueError):
                wx.build_workbook(conn.cursor(), [F], START, END, hours_per_fte=math.nan)
        finally:
            conn.close()


class TestStatementsCsv:
    def test_wide_layout_has_labels_keys_and_deduction_lines(self, client, statements):
        r = client.get("/reporting/statements/export", params=WINDOW)
        assert r.status_code == 200
        rows = list(csv.reader(io.StringIO(r.text)))
        assert rows[0] == ["Income Statement", "line_item"] + PERIODS + ["Total"]
        by_key = {row[1]: row for row in rows if len(row) > 1}
        assert by_key["excise_tax"][0] == "Excise tax (FET)"
        assert float(by_key["excise_tax"][-1]) == pytest.approx(statements["income_statement"]["totals"]["excise_tax"], abs=0.01)
        assert by_key["revenue"][0] == "Net sales"
        assert by_key["gross_sales"][0] == "Gross sales"
        assert by_key["commissions"][0] == "Sales commissions"
        assert by_key["excise_tax_payable"][0] == "Excise tax payable"
        assert by_key["change_in_excise_tax_payable"][0] == "Increase / (decrease) in excise tax payable"
        assert ["Balance Sheet", "line_item"] + PERIODS in rows

    def test_long_layout_unchanged(self, client):
        r = client.get("/reporting/statements/export", params={**WINDOW, "layout": "long"})
        assert r.status_code == 200
        assert r.text.splitlines()[0] == "statement,period,line_item,amount"
        assert "Income Statement,2027-01,excise_tax," in r.text
