"""
Sales Detail cross-foots (F5 / F6 / RM-7) and the workbook Read Me states
the assumptions exactly (RM-6).

* unit_price keeps 4 decimals (a 40%-off-list price of 23.994 is not shown
  as 23.99), so quantity x unit price = gross sales on every line;
* every money column is rounded to the cent per line and each line
  cross-foots: deductions add up to the total, gross - deductions = net,
  materials + labor + machine = product COGS, net - COGS = contribution;
* ``totals`` are the sums of the rounded lines (what SUM gives in Excel);
  ``report_totals`` are computed from the unrounded lines and equal the
  income statement; ``rounding`` is the difference;
* the workbook's Sales Detail sheet shows the 4-decimal price, a Total row
  equal to the SUM of the lines and, when the per-line rounding moves it, a
  rounding row and a total that equals the Income Statement;
* the Read Me shows the income tax rate as stored (25.5%, not 26%) and the
  gross sales rule the engine uses.
"""
import io
import os
import shutil
import sqlite3
import sys
import tempfile
from decimal import Decimal

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

_APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, "app"))
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from app.db.database import DatabaseManager  # noqa: E402

F = "F-RND"
WINDOW = {"forecast_ids": F, "start_period": "2027-01", "end_period": "2027-02"}
MONEY = ("gross_sales", "terms_allowance", "chargebacks", "other_allowances", "excise_tax", "total_sales_deductions",
         "net_sales", "commissions", "materials", "labor", "machine", "product_cogs", "contribution_margin")
PAIRS = [  # (sales detail key, income statement key)
    ("gross_sales", "gross_sales"), ("terms_allowance", "terms_allowance"), ("chargebacks", "chargebacks"),
    ("other_allowances", "other_allowances"), ("excise_tax", "excise_tax"),
    ("total_sales_deductions", "total_sales_deductions"), ("net_sales", "revenue"), ("commissions", "commissions"),
    ("materials", "cogs_materials"), ("labor", "cogs_labor"), ("machine", "cogs_manufacturing"),
    ("quantity", "units_sold"),
]


def _seed(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    c.execute("INSERT INTO forecast (forecast_id, name, is_base) VALUES (?, 'Rounding plan', 1)", (F,))
    c.executemany("INSERT INTO customers (customer_id, customer_name, customer_type) VALUES (?, ?, ?)", [
        ("C-DST", "Distributor Co", "Distributor"), ("C-DLR", "Dealer Co", "Dealer"),
    ])
    c.executemany("""INSERT INTO units (unit_id, unit_name, base_price, unit_type, bom_id, bom_version, router_id)
                     VALUES (?, ?, ?, ?, ?, '1.0', ?)""", [
        ("SRA-A100", "Sling", 39.99, "Accessory", "B-A", "R-A"),
        ("RIF", "Rifle", 1299.0, "Rifle", "B-R", "R-R"),
    ])
    c.executemany("""INSERT INTO bom (bom_id, version, bom_line, material_description, qty, unit, unit_price, material_cost)
                     VALUES (?, '1.0', ?, ?, ?, 'ea', ?, ?)""", [
        ("B-A", 1, "Webbing", 1.3, 3.3333, 1.3 * 3.3333),
        ("B-R", 1, "Barrel", 1, 210.555, 210.555),
    ])
    c.execute("INSERT INTO machines (machine_id, machine_name, machine_rate, available_minutes_per_month) "
              "VALUES ('CNC', 'CNC', 97.3, 100000)")
    c.execute("INSERT INTO labor_rates (rate_id, rate_name, rate_amount) VALUES ('ASSY', 'Assembler', 23.77)")
    c.executemany("""INSERT INTO router_operations (router_id, sequence, machine_id, machine_minutes, labor_minutes, labor_type_id)
                     VALUES (?, 10, 'CNC', ?, ?, 'ASSY')""", [("R-A", 1.7, 2.3), ("R-R", 33.3, 41.7)])
    sales = [
        # 40% off list $39.99 = 23.994 (F5): 630 x 23.994 = 15,116.22, not 630 x 23.99 = 15,113.70
        ("S1", "C-DST", "SRA-A100", "2027-01", 630, 23.994),
        # three lines of 10.004: each rounds to 10.00, the unrounded total is 30.012 -> 30.01
        ("S2", "C-DLR", "SRA-A100", "2027-01", 1, 10.004),
        ("S3", "C-DLR", "SRA-A100", "2027-02", 1, 10.004),
        ("S4", "C-DST", "SRA-A100", "2027-02", 1, 10.004),
        ("S5", "C-DLR", "RIF", "2027-01", 7, 1110.3333),
        ("S6", "C-DST", "RIF", "2027-02", 13, 999.9944),
    ]
    for i in range(20):   # more 4-decimal lines so the per-line rounding visibly moves the totals
        sales.append((f"X{i}", "C-DLR" if i % 2 else "C-DST", "RIF" if i % 3 else "SRA-A100",
                      "2027-01" if i % 2 else "2027-02", 3 + i, 101.0037 + i))
    for sid, cust, unit, period, qty, price in sales:
        c.execute("""INSERT INTO sales (sale_id, customer_id, unit_id, period, quantity, unit_price, total_revenue, forecast_id)
                     VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", (sid, cust, unit, period, qty, price, qty * price, F))
    rules = [
        ("DED-T", "terms_allowance", "percent_of_gross", 0.0175, 0, None, None),
        ("DED-CB", "chargebacks", "percent_of_gross", 0.0133, 0, "Dealer", None),
        ("DED-OA", "other_allowance", "per_unit", 0, 0.333, None, None),
        ("DED-FET", "excise_tax", "embedded_tax", 0.11, 0, None, "Rifle"),
        ("DED-COM", "commission", "percent_of_net", 0.035, 0, None, None),
    ]
    for rule_id, dtype, method, rate, per_unit, ctype, utype in rules:
        c.execute("""INSERT INTO sales_deduction_rules (rule_id, deduction_type, method, rate, amount_per_unit,
                     customer_type, unit_type, is_active) VALUES (?, ?, ?, ?, ?, ?, ?, 1)""",
                  (rule_id, dtype, method, rate, per_unit, ctype, utype))
    c.execute("""INSERT INTO financial_assumptions (forecast_id, opening_period, opening_cash, dso_days, dio_days,
                 dpo_days, income_tax_rate, excise_payable_days) VALUES (?, '2026-12', 50000, 45.5, 60, 30, 0.255, 15)""",
              (F,))
    conn.commit()
    conn.close()


@pytest.fixture(scope="module")
def report_db():
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, "rounding.db")
    mgr = DatabaseManager(database_path=path, data_dir=tmp)
    mgr.initialize()
    _seed(path)
    yield mgr
    shutil.rmtree(tmp, ignore_errors=True)


@pytest.fixture(scope="module")
def client(report_db):
    import app.db.database as dbmod
    import db.database as dbmod_short
    from api import reporting_routes
    saved = (dbmod.db_manager, dbmod_short.db_manager, reporting_routes.db_manager)
    dbmod.db_manager = report_db
    dbmod_short.db_manager = report_db
    reporting_routes.db_manager = report_db
    app = FastAPI()
    app.include_router(reporting_routes.router)
    yield TestClient(app)
    dbmod.db_manager, dbmod_short.db_manager, reporting_routes.db_manager = saved


def _get(client, path, **params):
    r = client.get(path, params={**WINDOW, **params})
    assert r.status_code == 200, r.text
    return r.json()["data"]


@pytest.fixture(scope="module")
def detail(client):
    return _get(client, "/reporting/sales-detail", limit=20000)


@pytest.fixture(scope="module")
def income(client):
    return _get(client, "/reporting/statements")["income_statement"]["totals"]


@pytest.fixture(scope="module")
def wb(client):
    r = client.get("/reporting/workbook", params=WINDOW)
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content))


def _is_cents(value) -> bool:
    return Decimal(repr(float(value))) == Decimal(repr(float(value))).quantize(Decimal("0.01"))


def _cents(value) -> int:
    return int(round(float(value) * 100))


class TestSalesDetailLines:
    def test_unit_price_keeps_four_decimals(self, detail):
        rows = {r["sale_id"]: r for r in detail["rows"]}
        assert rows["S1"]["unit_price"] == 23.994
        assert rows["S1"]["gross_sales"] == 15116.22
        assert rows["S5"]["unit_price"] == 1110.3333
        assert rows["S2"]["unit_price"] == 10.004

    def test_every_money_value_is_a_cent_amount(self, detail):
        for row in detail["rows"]:
            for key in MONEY:
                assert _is_cents(row[key]), (row["sale_id"], key, row[key])

    def test_every_line_cross_foots(self, detail):
        for row in detail["rows"]:
            sid = row["sale_id"]
            assert _cents(row["quantity"] * row["unit_price"]) == _cents(row["gross_sales"]), sid
            deductions = row["terms_allowance"] + row["chargebacks"] + row["other_allowances"] + row["excise_tax"]
            assert _cents(deductions) == _cents(row["total_sales_deductions"]), sid
            assert _cents(row["gross_sales"] - row["total_sales_deductions"]) == _cents(row["net_sales"]), sid
            assert _cents(row["materials"] + row["labor"] + row["machine"]) == _cents(row["product_cogs"]), sid
            assert _cents(row["net_sales"] - row["product_cogs"]) == _cents(row["contribution_margin"]), sid


class TestSalesDetailTotals:
    def test_totals_are_the_sums_of_the_rounded_lines(self, detail):
        assert detail["total_count"] == len(detail["rows"]) == 26
        for key in MONEY:
            assert _cents(detail["totals"][key]) == sum(_cents(r[key]) for r in detail["rows"]), key
        assert detail["totals"]["quantity"] == pytest.approx(sum(r["quantity"] for r in detail["rows"]))
        assert "unit_price" not in detail["totals"]

    def test_report_totals_equal_the_income_statement(self, detail, income):
        report = detail["report_totals"]
        for key, stmt_key in PAIRS:
            assert report[key] == pytest.approx(income[stmt_key], abs=0.001), key
        assert report["product_cogs"] == pytest.approx(
            income["cogs_materials"] + income["cogs_labor"] + income["cogs_manufacturing"], abs=0.011)

    def test_rounding_bridges_the_line_sums_to_the_statements(self, detail):
        rounding = detail["rounding"]
        for key in MONEY:
            assert rounding[key] == pytest.approx(detail["report_totals"][key] - detail["totals"][key], abs=0.001), key
        # The seed is built so the per-line rounding moves gross sales (3 x 10.004 alone is +0.01).
        assert any(abs(rounding[k]) >= 0.01 for k in MONEY)

    def test_a_page_keeps_the_window_totals(self, client, detail):
        page = _get(client, "/reporting/sales-detail", limit=3, offset=2)
        assert page["totals"] == detail["totals"]
        assert page["report_totals"] == detail["report_totals"]
        assert page["rounding"] == detail["rounding"]


class TestWorkbookSalesDetail:
    def _sheet(self, wb):
        ws = wb["Sales Detail"]
        header = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2))
        return ws, header, rows

    def test_unit_price_cell_keeps_four_decimals(self, wb):
        ws, header, rows = self._sheet(wb)
        col = header.index("Unit price")
        s1 = next(r for r in rows if r[header.index("Sale ID")].value == "S1")
        assert s1[col].value == 23.994
        assert s1[col].number_format == "#,##0.00##"
        assert s1[header.index("Gross sales")].value == 15116.22

    def test_total_row_is_the_sum_of_the_lines_and_rounding_reaches_the_statement(self, wb, detail, income):
        ws, header, rows = self._sheet(wb)
        lines = [r for r in rows if r[0].value in ("2027-01", "2027-02")]
        assert len(lines) == 26
        by_label = {r[0].value: r for r in rows if isinstance(r[0].value, str) and not r[0].value.startswith("2027")}
        total = by_label["Total"]
        for key, label in (("gross_sales", "Gross sales"), ("net_sales", "Net sales"),
                           ("total_sales_deductions", "Total sales deductions"), ("excise_tax", "Excise tax (FET)")):
            col = header.index(label)
            assert _cents(total[col].value) == sum(_cents(r[col].value) for r in lines), label
        rounding = next(r for label, r in by_label.items() if label.startswith("Rounding"))
        stmt = next(r for label, r in by_label.items() if label.startswith("Total per Income Statement"))
        for key, stmt_key in (("gross_sales", "gross_sales"), ("net_sales", "revenue"),
                              ("total_sales_deductions", "total_sales_deductions"), ("excise_tax", "excise_tax"),
                              ("commissions", "commissions")):
            col = header.index(next(h for k, h in (("gross_sales", "Gross sales"), ("net_sales", "Net sales"),
                                                     ("total_sales_deductions", "Total sales deductions"),
                                                     ("excise_tax", "Excise tax (FET)"),
                                                     ("commissions", "Sales commissions")) if k == key))
            assert stmt[col].value == pytest.approx(income[stmt_key], abs=0.001), key
            assert _cents(total[col].value) + _cents(rounding[col].value or 0) == _cents(stmt[col].value), key


class TestReadMe:
    def _text(self, wb):
        return {row[0]: row[1] for row in wb["Read Me"].iter_rows(values_only=True) if row and row[0]}

    def test_income_tax_rate_is_shown_as_stored(self, wb):
        text = self._text(wb)["Working capital"]
        assert "income tax rate of 25.5%" in text
        assert "26%" not in text
        assert "DSO 45.5" in text

    def test_gross_sales_rule_matches_the_engine(self, wb):
        text = self._text(wb)["Gross sales"]
        assert "Quantity x unit price" in text
        assert "only when the line has no price" in text
        assert "when present" not in text
