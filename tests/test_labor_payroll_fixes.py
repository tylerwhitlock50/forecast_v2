"""
Labor / payroll verified findings (review round v12), one test per finding:

* LAB-1 / PAYROLL-BRIDGE-CADENCE - one fixed pay calendar
  (payroll_config.pay_anchor_date): the same month has the same pay dates,
  paid amount and month-end accrued payroll whatever window the bridge is
  asked for, and the pay-run forecast uses the same paydays.
* LAB-3 - a raise applies from next_review_date on both sides (statements
  and pay runs), day-weighted inside the month / pay period.
* LAB-2 / LAB-5 - the labor policy message, the partial-coverage warning and
  the unclassified warning describe the months with and without direct pay,
  and their figures add up to the income statement's direct labor.
* LAB-4 / LABOR-CENT-MISMATCH - the labor forecast's standard cost and
  variance are the statement's figures to the cent (sample import data).
* LAB-7 / LOCK-4 - POST /payroll/bulk-update refuses a locked scenario's
  employee with 409 and the readable message; nothing is written.
* LAB-8 - POST /payroll/config validates pay_frequency and keeps omitted fields.
* LAB-9 - decimal weekly hours, end_date before start_date, negative pay.
* MCP-COGS-BRIDGE-COLUMNS - forecast_cogs_forecast's monthly bridge adds up.
"""
import csv
import io
import json
import os
import shutil
import sqlite3
import sys
import tempfile
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

_REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
_APP_DIR = os.path.join(_REPO, "app")
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from app.db.database import DatabaseManager  # noqa: E402
from services import financial_statements as fs  # noqa: E402

BASE = "F001"
LOCKED = "is approved and locked: branch it to make changes, or unlock it with a reason."
ZERO_BURDEN = {
    "federal_tax_rate": 0, "state_tax_rate": 0, "social_security_rate": 0, "medicare_rate": 0,
    "unemployment_rate": 0, "benefits_rate": 0, "workers_comp_rate": 0,
}
HEADERS = {"host": "localhost:8000", "accept": "application/json, text/event-stream",
           "content-type": "application/json"}
SAMPLE_DIR = os.path.join(_REPO, "examples", "sample_import_2027")


# ---------------------------------------------------------------------------
# Fixtures: the real app (every router, the governance middleware, MCP) on a scratch database
# ---------------------------------------------------------------------------

def _patch_db_manager(mgr):
    originals = {}
    for name, module in list(sys.modules.items()):
        if module is None:
            continue
        if name in ("db.database", "app.db.database") or name.startswith(("api.", "app.api.", "utils.")):
            if hasattr(module, "db_manager"):
                originals[name] = module.db_manager
                module.db_manager = mgr
    return originals


def _app(seed):
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, "forecast.db")
    mgr = DatabaseManager(database_path=path, data_dir=tmp)
    mgr.initialize()
    seed(path)
    from app.main import app as fastapi_app
    originals = _patch_db_manager(mgr)
    return tmp, mgr, fastapi_app, originals


def _teardown(tmp, mgr, originals):
    for name, original in originals.items():
        sys.modules[name].db_manager = original
    mgr.close_all_connections()
    shutil.rmtree(tmp, ignore_errors=True)


def _seed_small(path: str) -> None:
    """F001 with routing labor of $30 a unit (60 min @ $30/h): 100 units a month, Jan-Dec 2027."""
    conn = sqlite3.connect(path)
    c = conn.cursor()
    c.execute("INSERT INTO forecast (forecast_id, name, is_base) VALUES ('F001', '2027 Budget', 1)")
    c.execute("INSERT INTO customers (customer_id, customer_name, customer_type) VALUES ('C1', 'Acme', 'Dealer')")
    c.execute("""INSERT INTO units (unit_id, unit_name, base_price, bom_id, bom_version, router_id, router_version)
                 VALUES ('P1', 'Widget', 100, 'BOM1', '1.0', 'R1', '1.0')""")
    c.execute("INSERT INTO bom_definitions (bom_id, bom_name, version) VALUES ('BOM1', 'Widget BOM', '1.0')")
    c.execute("""INSERT INTO bom (bom_id, version, bom_line, material_description, qty, unit, unit_price, material_cost)
                 VALUES ('BOM1', '1.0', 1, 'Steel', 2, 'ea', 10, 20)""")
    c.execute("INSERT INTO router_definitions (router_id, router_name, version) VALUES ('R1', 'Widget Router', '1.0')")
    c.execute("INSERT INTO labor_rates (rate_id, rate_name, rate_amount, rate_type) VALUES ('L1', 'Assembler', 30, 'hourly')")
    c.execute("INSERT INTO machines (machine_id, machine_name, machine_rate) VALUES ('BENCH', 'Bench', 0)")
    c.execute("""INSERT INTO router_operations (router_id, sequence, machine_id, machine_minutes, labor_minutes, labor_type_id)
                 VALUES ('R1', 1, 'BENCH', 0, 60, 'L1')""")
    for i in range(1, 13):
        c.execute("INSERT INTO sales VALUES (?, 'C1', 'P1', ?, 100, 100, 10000, 'F001')", (f"S{i}", f"2027-{i:02d}"))
    c.execute("""INSERT INTO payroll_config (config_id, federal_tax_rate, state_tax_rate, social_security_rate, medicare_rate,
                 unemployment_rate, benefits_rate, workers_comp_rate, pay_frequency)
                 VALUES ('CFG', 0, 0, 0, 0, 0, 0, 0, 'biweekly')""")
    c.execute("""INSERT INTO financial_assumptions (forecast_id, opening_period, opening_cash, dso_days, dio_days, dpo_days,
                 income_tax_rate) VALUES ('F001', '2026-12', 50000, 30, 30, 30, 0.25)""")
    conn.commit()
    conn.close()


@pytest.fixture()
def app_env():
    tmp, mgr, fastapi_app, originals = _app(_seed_small)
    try:
        with TestClient(fastapi_app) as client:
            yield client, mgr
    finally:
        _teardown(tmp, mgr, originals)


def _data(res):
    assert res.status_code == 200, res.text
    return res.json()["data"]


def _config(client, **fields):
    _data(client.post("/payroll/config", json={**ZERO_BURDEN, **fields}))


def _hire(client, **fields):
    body = {"employee_name": "Erin", "department": "Production", "weekly_hours": 40, "hourly_rate": 26,
            "rate_type": "hourly", "labor_type": "Production", "start_date": "2020-01-01",
            "benefits_eligible": False, "expected_raise": 0}
    body.update(fields)
    return _data(client.post("/payroll/employees", json=body))["employee_id"]


def _bridge(client, start, end, **params):
    return _data(client.get("/payroll/statement-bridge", params={"forecast_id": BASE, "start_period": start,
                                                                 "end_period": end, **params}))


def _rows(bridge):
    return {r["period"]: r for r in bridge["by_period"]}


def _statements(client, start="2027-01", end="2027-12"):
    return _data(client.get("/reporting/statements", params={"forecast_ids": BASE, "start_period": start,
                                                             "end_period": end}))


def _mcp(client, name, arguments):
    res = client.post("/mcp", headers=HEADERS, follow_redirects=False, json={
        "jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": name, "arguments": arguments}})
    assert res.status_code == 200, res.text
    result = res.json()["result"]
    text = "".join(b.get("text", "") for b in result.get("content", []) if b.get("type") == "text")
    assert not result.get("isError"), text
    return json.loads(text)


# ---------------------------------------------------------------------------
# LAB-1 / PAYROLL-BRIDGE-CADENCE: one fixed pay calendar
# ---------------------------------------------------------------------------

def test_biweekly_pay_dates_do_not_depend_on_the_start_date():
    def march(start, anchor=None):
        return [d for d in fs.pay_run_dates("biweekly", start, 260, anchor) if (d.year, d.month) == (2027, 3)]

    starts = (date(2027, 1, 1), date(2027, 2, 1), date(2027, 3, 1), date(2026, 10, 12), date(2019, 5, 5))
    assert {tuple(march(s)) for s in starts} == {(date(2027, 3, 12), date(2027, 3, 26))}
    # The default calendar (no anchor configured) pays on 1 Jan 2027: 27 runs that year.
    runs_2027 = [d for d in fs.pay_run_dates("biweekly", date(2027, 1, 1), 30) if d.year == 2027]
    assert runs_2027[0] == date(2027, 1, 1) and len(runs_2027) == 27
    # Another payday moves every date by a week, whatever the start.
    assert march(date(2027, 2, 1), date(2027, 1, 8)) == [date(2027, 3, 5), date(2027, 3, 19)]
    assert march(date(2026, 1, 1), date(2027, 1, 8)) == [date(2027, 3, 5), date(2027, 3, 19)]
    # Weekly is every Friday; semimonthly / monthly are calendar days.
    assert fs.pay_run_dates("weekly", date(2027, 3, 1), 2) == [date(2027, 3, 5), date(2027, 3, 12)]


def test_bridge_month_is_the_same_in_every_window(app_env):
    client, _mgr = app_env
    _config(client, pay_frequency="biweekly")
    _hire(client, employee_name="Alan Admin", hourly_rate=104000, rate_type="salary", labor_class="admin",
          start_date="2022-01-01")
    _hire(client, employee_name="Dana Direct", hourly_rate=28, labor_class="direct", start_date="2027-03-17")
    full = _bridge(client, "2027-01", "2027-12")
    rows = _rows(full)
    assert [p for p, r in rows.items() if r["pay_runs"] == 3] == ["2027-01", "2027-07", "2027-12"]
    assert full["pay_anchor_date"] == fs.DEFAULT_PAY_ANCHOR_DATE.isoformat()
    for start, end in (("2027-02", "2027-12"), ("2027-03", "2027-03"), ("2027-04", "2027-06"), ("2027-02", "2027-04")):
        part = _bridge(client, start, end)
        before = fs.shift_period(start, -1)
        assert part["opening_accrued_payroll"] == pytest.approx(rows[before]["accrued_payroll_balance"], abs=0.01)
        for r in part["by_period"]:
            same = rows[r["period"]]
            for key in ("pay_dates", "pay_runs", "paid", "accrued", "earned", "accrued_payroll_balance"):
                assert r[key] == same[key], (start, end, r["period"], key)
    # The balance moves by pay earned less pay paid, month by month.
    previous = full["opening_accrued_payroll"]
    for r in full["by_period"]:
        assert r["accrued_payroll_balance"] == pytest.approx(previous + r["earned"] - r["paid"], abs=0.02)
        previous = r["accrued_payroll_balance"]
    # 31 Dec 2027 is a payday: nothing is owed at the year end.
    assert rows["2027-12"]["pay_dates"][-1] == "2027-12-31" and full["closing_accrued_payroll"] == 0


def test_pay_run_forecast_uses_the_bridge_paydays(app_env):
    client, _mgr = app_env
    _config(client, pay_frequency="biweekly")
    _hire(client)
    year = date.today().year
    bridge_dates = {d for r in _bridge(client, "2027-01", "2028-12")["by_period"] for d in r["pay_dates"]}
    bridge_dates |= {d for r in _bridge(client, f"{year}-01", f"{year + 2}-12")["by_period"] for d in r["pay_dates"]}
    for start in ("2027-03-01", "2027-02-17", None):      # None: from today
        params = {"forecast_id": BASE, "periods": 30, "include_employees": "false"}
        if start:
            params["start_date"] = start
        dates = [r["date"] for r in _data(client.get("/payroll/forecast", params=params))["forecast"]]
        assert len(dates) == 30 and set(dates) <= bridge_dates, start
    first = _data(client.get("/payroll/forecast", params={"forecast_id": BASE, "start_date": "2027-03-01",
                                                          "periods": 1}))["forecast"][0]
    assert first["date"] == "2027-03-12"


def test_pay_anchor_date_is_configurable_and_validated(app_env):
    client, _mgr = app_env
    _config(client, pay_frequency="biweekly", pay_anchor_date="2027-01-08")
    config = _data(client.get("/payroll/config"))["config"]
    assert config["pay_anchor_date"] == "2027-01-08" and config["pay_frequency"] == "biweekly"
    _hire(client)
    rows = _rows(_bridge(client, "2027-01", "2027-12"))
    assert rows["2027-01"]["pay_dates"] == ["2027-01-08", "2027-01-22"]
    assert sum(r["pay_runs"] for r in rows.values()) == 26
    # A rate change keeps the calendar.
    _data(client.post("/payroll/config", json={"benefits_rate": 0.2}))
    assert _data(client.get("/payroll/config"))["config"]["pay_anchor_date"] == "2027-01-08"
    tuesday = client.post("/payroll/config", json={"pay_anchor_date": "2027-01-05"})
    assert tuesday.status_code == 422 and "Friday" in tuesday.json()["detail"] and "2027-01-08" in tuesday.json()["detail"]
    assert client.post("/payroll/config", json={"pay_anchor_date": "soon"}).status_code == 422
    assert _data(client.get("/payroll/config"))["config"]["pay_anchor_date"] == "2027-01-08"


def test_mcp_config_tool_sets_the_pay_anchor(app_env):
    client, _mgr = app_env
    _mcp(client, "forecast_update_payroll_config", {"pay_anchor_date": "2027-01-15", "pay_frequency": "biweekly"})
    config = _mcp(client, "forecast_get_payroll_config", {})["config"]
    assert config["pay_anchor_date"] == "2027-01-15"
    _hire(client)
    bridge = _mcp(client, "forecast_payroll_statement_bridge", {"forecast_id": BASE, "start_period": "2027-01",
                                                               "end_period": "2027-01", "include_pay_dates": True})
    assert bridge["by_period"][0]["pay_dates"] == ["2027-01-01", "2027-01-15", "2027-01-29"]


def test_migration_adds_pay_anchor_date_to_an_older_database():
    tmp = tempfile.mkdtemp()
    try:
        path = os.path.join(tmp, "old.db")
        conn = sqlite3.connect(path)
        conn.execute("""CREATE TABLE payroll_config (config_id TEXT PRIMARY KEY, federal_tax_rate REAL,
                        state_tax_rate REAL, social_security_rate REAL, medicare_rate REAL, unemployment_rate REAL,
                        benefits_rate REAL, workers_comp_rate REAL, created_at TIMESTAMP)""")
        conn.execute("INSERT INTO payroll_config (config_id, benefits_rate) VALUES ('OLD', 0.3)")
        conn.commit()
        conn.close()
        mgr = DatabaseManager(database_path=path, data_dir=tmp)
        mgr.initialize_schema()
        mgr.initialize_schema()   # idempotent
        conn = sqlite3.connect(path)
        columns = {row[1] for row in conn.execute("PRAGMA table_info(payroll_config)")}
        assert {"pay_frequency", "pay_anchor_date"} <= columns
        config = fs.load_payroll_config(conn.cursor())
        conn.close()
        assert config["pay_anchor_date"] == fs.DEFAULT_PAY_ANCHOR_DATE.isoformat() and config["benefits_rate"] == 0.3
        mgr.close_all_connections()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------------------
# LAB-3: one raise rule, day-weighted on both sides
# ---------------------------------------------------------------------------

RAY = dict(employee_name="Ray FlatRaise", hourly_rate=24, labor_class="direct", next_review_date="2027-09-16",
           expected_raise=2.0)
OLD_MONTH = 24 * 40 * 52 / 12      # 4,160.00
NEW_MONTH = 26 * 40 * 52 / 12      # 4,506.67


def test_mid_month_raise_is_day_weighted_in_the_statements(app_env):
    client, mgr = app_env
    _hire(client, **RAY)
    conn = mgr.get_connection()
    try:
        result = fs.payroll_by_period(conn.cursor(), [BASE], ["2027-08", "2027-09", "2027-10"])
    finally:
        conn.close()
    total = result["total"]
    assert total["2027-08"] == pytest.approx(OLD_MONTH, abs=0.01)
    assert total["2027-09"] == pytest.approx(OLD_MONTH * 15 / 30 + NEW_MONTH * 15 / 30, abs=0.01)   # 4,333.33
    assert total["2027-10"] == pytest.approx(NEW_MONTH, abs=0.01)


@pytest.mark.parametrize("frequency", ["semimonthly", "monthly"])
def test_a_year_of_pay_runs_equals_the_accrual_with_a_mid_month_raise(app_env, frequency):
    client, _mgr = app_env
    _config(client, pay_frequency=frequency)
    _hire(client, **RAY)
    bridge = _bridge(client, "2027-01", "2027-12")
    for r in bridge["by_period"]:
        assert r["difference"] == pytest.approx(0, abs=0.01), r["period"]
    assert bridge["totals"]["accrued"] == pytest.approx(bridge["totals"]["paid"], abs=0.02)
    assert bridge["closing_accrued_payroll"] == pytest.approx(0, abs=0.01)
    assert _rows(bridge)["2027-09"]["paid"] == pytest.approx(OLD_MONTH * 15 / 30 + NEW_MONTH * 15 / 30, abs=0.01)
    assert "timing only" not in bridge["note"]
    runs = _data(client.get("/payroll/forecast", params={"forecast_id": BASE, "start_date": "2027-01-01",
                                                         "include_employees": "false"}))
    accrued = sum(r["accrued"] for r in bridge["by_period"])
    assert runs["totals"]["total_cost"] == pytest.approx(accrued, abs=0.05)


def test_biweekly_run_across_a_raise_pays_each_day_at_its_rate(app_env):
    client, _mgr = app_env
    _config(client, pay_frequency="biweekly")
    _hire(client, **RAY)
    runs = _data(client.get("/payroll/forecast", params={"forecast_id": BASE, "start_date": "2027-09-10",
                                                         "periods": 2, "include_employees": "false"}))["forecast"]
    # 10 Sep run pays 28 Aug - 10 Sep (old rate); 24 Sep run pays 11-24 Sep: 5 days old, 9 days from 16 Sep.
    assert [r["date"] for r in runs] == ["2027-09-10", "2027-09-24"]
    assert runs[0]["gross_pay"] == pytest.approx(24 * 80, abs=0.01)
    assert runs[1]["gross_pay"] == pytest.approx(24 * 80 * 5 / 14 + 26 * 80 * 9 / 14, abs=0.01)


# ---------------------------------------------------------------------------
# LAB-2 / LAB-5: month-by-month labor policy and unclassified warning
# ---------------------------------------------------------------------------

def _money(value):
    return f"${value:,.2f}"


def test_policy_message_names_standard_only_months_and_adds_up(app_env):
    client, mgr = app_env
    _config(client, pay_frequency="monthly")
    _hire(client, employee_name="Mia MidYear", hourly_rate=25, labor_class="direct", start_date="2027-03-01")
    _hire(client, employee_name="Nora NoClass", hourly_rate=36000, rate_type="salary", start_date="2020-01-01")
    st = _statements(client)
    totals = st["income_statement"]["totals"]
    direct_labor = round(totals["cogs_labor"] + totals["labor_variance"], 2)
    policy = st["detail"]["labor_policy"]
    message = policy["message"]
    assert policy["months_with_direct_pay"][0] == "2027-03" and policy["standard_only_months"] == ["2027-01", "2027-02"]
    assert policy["cogs_direct_labor"] == direct_labor
    assert "so it equals the pay" not in message
    assert f"Direct labor in COGS is {_money(direct_labor)}." in message
    assert f"({_money(policy['direct_payroll'])})" in message                       # pay of employees classed Direct
    assert f"routing standard alone ({_money(direct_labor - policy['direct_payroll'])})" in message
    assert "2027-03..2027-12" in message and "2027-01..2027-02" in message
    # Jan + Feb carry the standard alone: 2 x 100 units x $30.
    assert policy["standard_only_labor"] == pytest.approx(6000, abs=0.01)
    partial = next(w for w in st["warnings"] if w["code"] == "partial_direct_payroll")
    assert partial["periods"] == ["2027-01", "2027-02"] and "Gross margin moves" in partial["message"]
    unclassified = next(w for w in st["warnings"] if w["code"] == "unclassified_payroll")
    assert "Nora NoClass" in unclassified["message"]
    assert "in months without direct pay (2027-01..2027-02)" in unclassified["message"]
    assert "instead of COGS" in unclassified["message"]
    # The labor forecast and the workbook say the same, with the same figures.
    labor = _data(client.get("/reporting/labor-forecast", params={"forecast_ids": BASE, "start_period": "2027-01",
                                                                  "end_period": "2027-12"}))
    assert labor["payroll_note"]["message"].startswith(message)
    from services.workbook_export import workbook_bytes
    conn = mgr.get_connection()
    try:
        wb = load_workbook(io.BytesIO(workbook_bytes(conn.cursor(), [BASE], "2027-01", "2027-12")))
    finally:
        conn.close()
    readme = [row for row in wb["Read Me"].iter_rows(values_only=True) if row and row[0] == "Payroll vs direct labor"]
    assert readme and str(readme[0][1]).startswith(message)


def test_unclassified_warning_follows_the_direct_payroll_in_cogs(app_env):
    client, _mgr = app_env
    _config(client, pay_frequency="monthly")
    nora = _hire(client, employee_name="Nora NoClass", hourly_rate=25, start_date="2020-01-01")
    # No one classed Direct: the routing standard stands alone, so an unclassified production worker counts twice.
    st = _statements(client)
    warning = next(w for w in st["warnings"] if w["code"] == "unclassified_payroll")
    assert "counted twice" in warning["message"]
    assert "counted twice" in st["detail"]["labor_policy"]["message"]
    assert warning["direct_payroll_in_cogs"] is False
    # Direct pay in every month: direct payroll replaced the standard, so Nora is counted once, in opex.
    _hire(client, employee_name="Dana Direct", hourly_rate=25, labor_class="direct", start_date="2020-01-01")
    st = _statements(client)
    warning = next(w for w in st["warnings"] if w["code"] == "unclassified_payroll")
    assert "counted twice" not in warning["message"]
    assert "counted once, in Payroll (operating expenses) instead of COGS" in warning["message"]
    assert warning["direct_payroll_in_cogs"] is True and warning["standard_only_months"] == []
    policy = st["detail"]["labor_policy"]
    assert "equals the pay of employees classed Direct" in policy["message"]
    assert "counted twice (routing standard" not in policy["message"]
    assert "instead of COGS" in policy["message"]
    assert not any(w["code"] == "partial_direct_payroll" for w in st["warnings"])
    totals = st["income_statement"]["totals"]
    assert f"({_money(round(totals['cogs_labor'] + totals['labor_variance'], 2))})" in policy["message"]
    assert nora


# ---------------------------------------------------------------------------
# LAB-4 / LABOR-CENT-MISMATCH: the labor forecast is the statement, to the cent
# ---------------------------------------------------------------------------

def _csv(name):
    with open(os.path.join(SAMPLE_DIR, name), newline="", encoding="utf-8-sig") as f:
        return [dict(r) for r in csv.DictReader(f)]


def _load_sample(client):
    fid = _data(client.post("/forecast/scenario", json={"name": "Sample 2027"}))["forecast_id"]
    for table, name in (("customers", "customers.csv"), ("units", "products.csv"), ("machines", "machines.csv"),
                        ("labor_rates", "labor_rates.csv")):
        _data(client.post("/import/records", json={"table": table, "rows": _csv(name)}))
    _data(client.post("/import/routings", json={"rows": _csv("routings.csv")}))
    sheet = load_workbook(os.path.join(SAMPLE_DIR, "boms.xlsx"), data_only=True).active
    raw = list(sheet.iter_rows(values_only=True))
    _data(client.post("/import/boms", json={"rows": [dict(zip(raw[1], r)) for r in raw[2:] if any(v is not None for v in r)]}))
    _data(client.post("/forecast/smart_import_sales", json={"forecast_id": fid, "rows": _csv("sales_forecast_2027.csv")}))
    return fid


def test_labor_forecast_ties_to_the_income_statement_to_the_cent(app_env):
    client, mgr = app_env
    fid = _load_sample(client)
    # The reviewer's setup: default burden, benefits-eligible direct staff (hires, a salaried worker, a
    # raise), so the variance has a month that lands on a half cent.
    _data(client.post("/payroll/config", json={
        "federal_tax_rate": 0.22, "state_tax_rate": 0.06, "social_security_rate": 0.062, "medicare_rate": 0.0145,
        "unemployment_rate": 0.006, "benefits_rate": 0.25, "workers_comp_rate": 0.015, "pay_frequency": "biweekly"}))
    for emp in (dict(employee_name="Dana DirectHourly", hourly_rate=28.0, start_date="2025-01-06"),
                dict(employee_name="Sam DirectSalaried", hourly_rate=62000, rate_type="salary", start_date="2024-06-01"),
                dict(employee_name="Mia MidMonth", hourly_rate=25.0, start_date="2027-03-17"),
                dict(employee_name="Ray FlatRaise", hourly_rate=24.0, start_date="2021-01-01",
                     next_review_date="2027-09-15", expected_raise=2.0),
                dict(employee_name="Scott ScenarioDirect", hourly_rate=27.0, start_date="2027-01-04", forecast_id=fid)):
        _hire(client, labor_class="direct", benefits_eligible=True, **emp)
    params = {"forecast_ids": fid, "start_period": "2027-01", "end_period": "2027-12"}
    labor = _data(client.get("/reporting/labor-forecast", params=params))
    st = _data(client.get("/reporting/statements", params=params))
    income = {r["period"]: r for r in st["income_statement"]["by_period"]}
    standard = st["income_statement"]["totals"]["cogs_labor"]
    assert labor["totals"]["cost"] == labor["payroll_note"]["direct_labor_cost"] == standard
    assert labor["standard_vs_payroll"]["totals"]["standard_cost"] == standard
    assert labor["totals"]["labor_variance"] == st["income_statement"]["totals"]["labor_variance"]
    for row in labor["standard_vs_payroll"]["by_period"]:
        assert row["standard_cost"] == income[row["period"]]["cogs_labor"], row["period"]
        assert row["labor_variance"] == income[row["period"]]["labor_variance"], row["period"]
    for row in labor["totals_by_period"]:
        assert row["cost"] == income[row["period"]]["cogs_labor"], row["period"]
    # The workbook's Labor Forecast sheet shows one figure for the standard.
    from services.workbook_export import workbook_bytes
    conn = mgr.get_connection()
    try:
        wb = load_workbook(io.BytesIO(workbook_bytes(conn.cursor(), [fid], "2027-01", "2027-12")))
    finally:
        conn.close()
    ws = wb["Labor Forecast"]
    total_col = ws.max_column
    found = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label in ("Total direct labor", "Direct labor (standard)") and isinstance(ws.cell(row=r, column=total_col).value, float):
            found.setdefault(label, []).append(ws.cell(row=r, column=total_col).value)
    assert standard in found["Direct labor (standard)"]
    assert standard in found["Total direct labor"]


# ---------------------------------------------------------------------------
# LAB-7 / LOCK-4: bulk update of a locked scenario's employee
# ---------------------------------------------------------------------------

def test_bulk_update_refuses_a_locked_scenarios_employee_and_writes_nothing(app_env):
    client, mgr = app_env
    _config(client, pay_frequency="monthly")
    shared = _hire(client, employee_name="Shared Admin", hourly_rate=25, labor_class="admin")
    scoped = _hire(client, employee_name="Scenario Worker", hourly_rate=30, forecast_id=BASE)
    approved = client.post(f"/forecast/scenario/{BASE}/approve", json={"approved_by": "Dana CFO"})
    assert approved.status_code == 200, approved.text

    res = client.post("/payroll/bulk-update", json=[{"employee_id": scoped, "labor_class": "admin"}])
    assert res.status_code == 409
    detail = res.json()["detail"]
    assert "Scenario '2027 Budget' " + LOCKED in detail and "SCENARIO_LOCKED" not in detail
    assert "nothing was changed" in detail
    # The same body as every other lock refusal (LAB7-AUDIT).
    assert res.json()["code"] == "SCENARIO_LOCKED" and res.json()["forecast_id"] == BASE

    mixed = client.post("/payroll/bulk-update", json=[{"employee_id": shared, "weekly_hours": 30},
                                                      {"employee_id": scoped, "weekly_hours": 20}])
    assert mixed.status_code == 409
    # Moving a company-wide employee into the locked scenario is refused too.
    moved = client.post("/payroll/bulk-update", json=[{"employee_id": shared, "forecast_id": BASE}])
    assert moved.status_code == 409
    conn = sqlite3.connect(mgr.database_path)
    try:
        rows = dict(conn.execute("SELECT employee_id, weekly_hours FROM payroll").fetchall())
        audit = conn.execute("SELECT status, summary, forecast_id FROM audit_log "
                             "WHERE path = '/payroll/bulk-update'").fetchall()
        assert conn.execute("SELECT forecast_id FROM payroll WHERE employee_id = ?", (shared,)).fetchone()[0] is None
    finally:
        conn.close()
    assert rows[shared] == 40 and rows[scoped] == 40
    # Every refusal (the mixed company-wide + scenario batch included) is filed under the locked scenario.
    assert len(audit) == 3 and all(status == 409 and summary.startswith("Refused (409)") and fid == BASE
                                   for status, summary, fid in audit)
    # A company-wide employee alone still updates.
    ok = _data(client.post("/payroll/bulk-update", json=[{"employee_id": shared, "weekly_hours": 37.5}]))
    assert ok["updated_count"] == 1 and not ok["errors"]


# ---------------------------------------------------------------------------
# LAB-8: POST /payroll/config
# ---------------------------------------------------------------------------

def test_config_rejects_an_unknown_frequency_and_keeps_omitted_fields(app_env):
    client, _mgr = app_env
    _data(client.post("/payroll/config", json={"pay_frequency": "semimonthly", "benefits_rate": 0.30,
                                               "workers_comp_rate": 0.05}))
    for bad in ("fortnightly", "every two weeks", ""):
        res = client.post("/payroll/config", json={"pay_frequency": bad, "benefits_rate": 0.1})
        assert res.status_code == 422, bad
    config = _data(client.get("/payroll/config"))["config"]
    assert config["pay_frequency"] == "semimonthly" and config["benefits_rate"] == 0.30
    # A partial update changes only what it names.
    saved = _data(client.post("/payroll/config", json={"pay_frequency": "weekly"}))
    assert saved["config"]["pay_frequency"] == "weekly"
    config = _data(client.get("/payroll/config"))["config"]
    assert (config["benefits_rate"], config["workers_comp_rate"]) == (0.30, 0.05)
    # The UI's spelling is read as the frequency it names, never as another one.
    _data(client.post("/payroll/config", json={"pay_frequency": "Semi-monthly"}))
    assert _data(client.get("/payroll/config"))["config"]["pay_frequency"] == "semimonthly"
    for body in ({"benefits_rate": -2}, {"medicare_rate": 1.5}, {}):
        assert client.post("/payroll/config", json=body).status_code == 422, body


# ---------------------------------------------------------------------------
# LAB-9: employee input guards
# ---------------------------------------------------------------------------

def test_employee_inputs_accept_decimal_hours_and_refuse_impossible_values(app_env):
    client, mgr = app_env
    emp_id = _hire(client, employee_name="Part Timer", weekly_hours=37.5, hourly_rate=20, start_date="2027-01-01")
    emp = _data(client.get(f"/payroll/employees/{emp_id}"))["employee"]
    assert emp["weekly_hours"] == 37.5 and emp["annual_gross"] == pytest.approx(20 * 37.5 * 52, abs=0.01)
    conn = mgr.get_connection()
    try:
        jan = fs.payroll_by_period(conn.cursor(), [BASE], ["2027-01"])["total"]["2027-01"]
    finally:
        conn.close()
    assert jan == pytest.approx(20 * 37.5 * 52 / 12, abs=0.01)
    body = {"employee_name": "Part Timer", "department": "Production", "weekly_hours": 32.25, "hourly_rate": 20,
            "rate_type": "hourly", "labor_type": "Production", "start_date": "2027-01-01"}
    _data(client.put(f"/payroll/employees/{emp_id}", json=body))
    assert _data(client.get(f"/payroll/employees/{emp_id}"))["employee"]["weekly_hours"] == 32.25

    base = {"employee_name": "A", "department": "Ops", "weekly_hours": 40, "hourly_rate": 20, "rate_type": "hourly",
            "labor_type": "Assembler", "start_date": "2030-06-01"}
    reversed_dates = client.post("/payroll/employees", json={**base, "end_date": "2030-01-31"})
    assert reversed_dates.status_code == 422 and "before start_date" in reversed_dates.json()["detail"]
    assert client.put(f"/payroll/employees/{emp_id}", json={**body, "end_date": "2026-12-31"}).status_code == 422
    for bad in ({"hourly_rate": -20}, {"weekly_hours": -10}, {"weekly_hours": 200}):
        assert client.post("/payroll/employees", json={**base, **bad}).status_code == 422, bad
    res = _data(client.post("/payroll/bulk-update", json=[{"employee_id": emp_id, "end_date": "2026-06-30"},
                                                          {"employee_id": emp_id, "hourly_rate": -5}]))
    assert res["updated_count"] == 0 and len(res["errors"]) == 2
    assert _data(client.get(f"/payroll/employees/{emp_id}"))["employee"]["end_date"] is None


def test_mcp_create_employee_takes_decimal_hours(app_env):
    client, _mgr = app_env
    created = _mcp(client, "forecast_create_employee", {
        "employee_name": "Office", "department": "Admin", "weekly_hours": 37.5, "hourly_rate": 30,
        "labor_type": "Admin", "start_date": "2027-01-01", "labor_class": "admin"})
    emp = _mcp(client, "forecast_get_employee", {"employee_id": created["employee_id"]})["employee"]
    assert emp["weekly_hours"] == 37.5


# ---------------------------------------------------------------------------
# MCP-COGS-BRIDGE-COLUMNS
# ---------------------------------------------------------------------------

def test_mcp_cogs_bridge_months_add_up(app_env):
    client, _mgr = app_env
    _config(client, pay_frequency="monthly")
    _hire(client, employee_name="Direct Dan", hourly_rate=30, labor_class="direct")
    _hire(client, employee_name="Indirect Ivy", hourly_rate=25, labor_class="indirect")
    cogs = _mcp(client, "forecast_cogs_forecast", {"forecast_ids": [BASE]})
    columns = cogs["by_period"]["columns"]
    assert {"labor_variance", "indirect_labor"} <= set(columns)
    assert columns.index("product_cogs") < columns.index("labor_variance") < columns.index("indirect_labor") \
        < columns.index("factory_overhead")
    rows = [dict(zip(columns, r)) for r in cogs["by_period"]["rows"]]
    income = {r["period"]: r for r in _statements(client)["income_statement"]["by_period"]}
    for row in rows:
        parts = row["product_cogs"] + row["labor_variance"] + row["indirect_labor"] + row["factory_overhead"] \
            + row["other_cogs"]
        assert parts == pytest.approx(row["total_cogs"], abs=0.02), row["period"]
        assert row["labor_variance"] == pytest.approx(income[row["period"]]["labor_variance"], abs=0.01)
        assert row["indirect_labor"] == pytest.approx(income[row["period"]]["indirect_labor"], abs=0.01)
        assert row["indirect_labor"] == pytest.approx(25 * 40 * 52 / 12, abs=0.01)
    for key in ("labor_variance", "indirect_labor"):
        assert sum(r[key] for r in rows) == pytest.approx(cogs["totals"][key], abs=0.05)
