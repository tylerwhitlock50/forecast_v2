"""
End-to-end MCP workflow test: build a small firearms model from nothing
ONLY through MCP tools (scenario, customers / products / machines / labor
rates, routings from a CSV file, BOMs from an XLSX file, prices, deduction
rules, the 2027 sales forecast), then check that the COGS forecast, labor
forecast and statements agree, the deduction maths (11% embedded FET,
precedence, exemptions), that dry runs write nothing, that unknown
arguments are rejected, and that the workbook export lands in the data
folder.

Drives the real FastAPI app over JSON-RPC at /mcp against an isolated
SQLite database in a temporary data folder, like tests/test_mcp_server.py.
"""
import csv
import json
import os
import shutil
import sys
import tempfile

import pytest
from fastapi.testclient import TestClient

_APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, "app"))
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from app.db.database import DatabaseManager  # noqa: E402

HEADERS = {
    "host": "localhost:8000",
    "accept": "application/json, text/event-stream",
    "content-type": "application/json",
}
YEAR = [f"2027-{m:02d}" for m in range(1, 13)]


def _patch_db_manager(mgr):
    originals = {}
    for name, module in list(sys.modules.items()):
        if module is None:
            continue
        if name in ("db.database", "app.db.database") or name.startswith("api.") or name.startswith("app.api."):
            if hasattr(module, "db_manager"):
                originals[name] = module.db_manager
                module.db_manager = mgr
    return originals


@pytest.fixture()
def env():
    tmp = tempfile.mkdtemp()
    data_dir = os.path.join(tmp, "data")
    os.makedirs(data_dir)
    mgr = DatabaseManager(database_path=os.path.join(data_dir, "workflow.db"), data_dir=data_dir)
    mgr.initialize()

    from app.main import app as fastapi_app
    originals = _patch_db_manager(mgr)
    try:
        with TestClient(fastapi_app) as client:
            yield client, mgr, data_dir, tmp
    finally:
        for name, original in originals.items():
            sys.modules[name].db_manager = original
        mgr.close_all_connections()
        shutil.rmtree(tmp, ignore_errors=True)


def raw_call(client, name, arguments=None):
    res = client.post("/mcp", headers=HEADERS, follow_redirects=False, json={
        "jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": name, "arguments": arguments or {}},
    })
    assert res.status_code == 200, res.text
    result = res.json()["result"]
    text = "".join(b.get("text", "") for b in result.get("content", []) if b.get("type") == "text")
    return result, text


def ok(client, name, arguments=None):
    result, text = raw_call(client, name, arguments)
    assert not result.get("isError"), text
    if result.get("structuredContent"):
        payload = result["structuredContent"]
        return payload.get("result", payload)
    return json.loads(text)


def err(client, name, arguments=None):
    result, text = raw_call(client, name, arguments)
    assert result.get("isError"), text
    return text


def count(mgr, sql, params=()):
    conn = mgr.get_connection()
    try:
        return conn.execute(sql, params).fetchone()[0]
    finally:
        mgr.close_connection(conn)


def approx(a, b, tol=0.02):
    return abs(float(a) - float(b)) <= tol


# ---------------------------------------------------------------------------
# Model inputs
# ---------------------------------------------------------------------------

CUSTOMERS = [
    {"Customer ID": "C-DLR", "Name": "Main Street Guns", "Channel": "Dealer", "Region": "Midwest"},
    {"Customer ID": "C-DST", "Name": "Big Sky Distributors", "Channel": "Distributor", "Region": "Mountain"},
    {"Customer ID": "C-GOV", "Name": "County Sheriff", "Channel": "Government", "Region": "Midwest"},
]
PRODUCTS = [
    {"SKU": "RFL-308", "Product Name": "Ridgeline .308 Rifle", "Product Type": "Rifle", "MSRP": "$1,499.00"},
    {"SKU": "PST-9C", "Product Name": "Compact 9mm Pistol", "Product Type": "Pistol", "MSRP": 599},
]
MACHINES = [
    {"machine_id": "CNC-01", "machine_name": "Haas CNC", "machine_rate": 85, "available_hours_per_month": 320},
    {"machine_id": "ASSY-01", "machine_name": "Assembly bench", "machine_rate": 0, "available_hours_per_month": 640},
]
LABOR = [
    {"rate_id": "MACHINIST", "rate_name": "CNC Machinist", "rate_amount": 32},
    {"rate_id": "ASSEMBLER", "rate_name": "Assembler", "rate_amount": 26},
]
ROUTING_CSV = [
    ["Routing", "Op", "Work Center", "Run Minutes", "Labor Min", "Labor Type", "Description", "Product"],
    ["R-RFL308", "10", "CNC-01", "60", "20", "MACHINIST", "Machine receiver", "RFL-308"],
    ["R-RFL308", "20", "Assembly bench", "0", "30", "Assembler", "Assemble and test fire", "RFL-308"],
    ["R-PST9C", "10", "CNC-01", "30", "10", "MACHINIST", "Machine slide", "PST-9C"],
    ["R-PST9C", "20", "ASSY-01", "0", "20", "ASSEMBLER", "Assemble", "PST-9C"],
]
BOM_ROWS = [
    {"Parent": "BOM-RFL308", "Rev": 1, "Component": "BRL-308", "Description": "Barrel blank", "Qty Per": 1, "Std Cost": 185, "Product": "RFL-308"},
    {"Parent": "BOM-RFL308", "Rev": 1, "Component": "RCV-700", "Description": "Receiver forging", "Qty Per": 1, "Std Cost": 140, "Product": "RFL-308"},
    {"Parent": "BOM-RFL308", "Rev": 1, "Component": "STK-SYN", "Description": "Synthetic stock", "Qty Per": 1, "Std Cost": 95, "Product": "RFL-308"},
    {"Parent": "BOM-RFL308", "Rev": 1, "Component": "SCR-10", "Description": "Screws", "Qty Per": 8, "Std Cost": 0.25, "Product": "RFL-308"},
    {"Parent": "BOM-PST9C", "Rev": 1, "Component": "SLD-9", "Description": "Slide", "Qty Per": 1, "Std Cost": 60, "Product": "PST-9C"},
    {"Parent": "BOM-PST9C", "Rev": 1, "Component": "FRM-9", "Description": "Frame", "Qty Per": 1, "Std Cost": 35, "Product": "PST-9C"},
    {"Parent": "BOM-PST9C", "Rev": 1, "Component": "KIT-9", "Description": "Small parts kit", "Qty Per": 1, "Std Cost": 25, "Product": "PST-9C"},
]
PRICES = [
    {"product": "RFL-308", "customer_type": "Dealer", "price": 1110, "start": "2027-01", "end": "2027-12"},
    {"product": "RFL-308", "customer_type": "Distributor", "price": 999},
    {"customer_type": "Government", "discount": "20%"},
    {"product": "PST-9C", "customer": "Big Sky Distributors", "price": 520},
]
DEDUCTIONS = [
    {"deduction_type": "terms_allowance", "name": "2% 10 net 30", "customer_type": "Dealer", "rate": 0.02},
    {"deduction_type": "chargebacks", "rate": 0.01},
    {"deduction_type": "excise_tax", "name": "FET 11%", "method": "embedded_tax", "rate": 0.11},
    {"deduction_type": "excise_tax", "name": "FET pistols 10%", "method": "embedded_tax", "rate": 0.10, "unit_type": "Pistol"},
    {"deduction_type": "excise_tax", "name": "Government exempt", "method": "embedded_tax", "rate": 0, "customer_type": "Government"},
    {"deduction_type": "commission", "name": "Distributor rep", "method": "percent_of_net", "rate": 0.03, "customer_type": "Distributor"},
]


def _sales_rows():
    period_forms = ["Jan-27", "2027-02", "3/2027", "2027/04", "May 2027", "2027-06", "7/1/2027", "2027-08",
                    "2027-09", "Oct-2027", "2027-11", "December 2027"]
    rows = [{"customer_id": "C-DLR", "unit_id": "RFL-308", "period": p, "quantity": 10} for p in period_forms]
    rows += [{"customer_id": "C-DST", "unit_id": "PST-9C", "period": p, "quantity": 50} for p in YEAR[:3]]
    rows.append({"customer_id": "C-GOV", "unit_id": "RFL-308", "period": "2027-06", "quantity": 5})
    return rows


def _write_inputs(data_dir):
    folder = os.path.join(data_dir, "imports")
    os.makedirs(folder, exist_ok=True)
    with open(os.path.join(folder, "routings.csv"), "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(ROUTING_CSV)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM export"
    ws.append(["Indented BOM report"])            # title row above the header
    headers = list(BOM_ROWS[0].keys())
    ws.append(headers)
    for row in BOM_ROWS:
        ws.append([row[h] for h in headers])
    wb.save(os.path.join(folder, "boms.xlsx"))


def _build_model(client, mgr, data_dir):
    scenario = ok(client, "forecast_create_scenario", {"name": "Budget 2027", "is_base": True})
    fid = scenario.get("forecast_id") or scenario.get("scenario", {}).get("forecast_id")
    assert fid
    _write_inputs(data_dir)

    for table, rows in (("customers", CUSTOMERS), ("products", PRODUCTS), ("machines", MACHINES),
                        ("labor_rates", LABOR)):
        res = ok(client, "forecast_import_records", {"table": table, "rows": rows})
        assert res["status"] == "ok" and res["created"] == len(rows), res
    assert count(mgr, "SELECT customer_type FROM customers WHERE customer_id = 'C-DST'") == "Distributor"
    assert count(mgr, "SELECT base_price FROM units WHERE unit_id = 'RFL-308'") == 1499

    routings = ok(client, "forecast_import_routings", {"file_path": "imports/routings.csv"})
    assert routings["status"] == "ok" and routings["created"] == 4, routings
    assert routings["source"]["row_count"] == 4
    rfl_route = next(r for r in routings["routings"] if r["router_id"] == "R-RFL308")
    assert approx(rfl_route["labor_cost_per_unit"], 20 / 60 * 32 + 30 / 60 * 26)
    assert rfl_route["used_by_products"] == ["RFL-308"]

    preview = ok(client, "forecast_preview_file", {"file_path": "imports/boms.xlsx"})
    assert preview["header_row"] == 2 and preview["row_count"] == 7
    assert preview["likely_imports"][0]["kind"] == "boms"

    dry = ok(client, "forecast_import_boms", {"file_path": "imports/boms.xlsx", "dry_run": True})
    assert dry["dry_run"] is True and dry["created"] == 7 and "next" in dry
    assert count(mgr, "SELECT COUNT(*) FROM bom") == 0, "a dry run must not write"
    boms = ok(client, "forecast_import_boms", {"file_path": "imports/boms.xlsx"})
    by_bom = {b["bom_id"]: b for b in boms["boms"]}
    assert approx(by_bom["BOM-RFL308"]["material_cost_per_unit"], 422)
    assert approx(by_bom["BOM-PST9C"]["material_cost_per_unit"], 120)
    assert by_bom["BOM-RFL308"]["version"] == "1.0" and by_bom["BOM-RFL308"]["used_by_products"] == ["RFL-308"]

    prices = ok(client, "forecast_import_prices", {"rows": PRICES})
    assert prices["status"] == "ok" and prices["created"] == 4, prices
    deductions = ok(client, "forecast_import_deductions", {"rows": DEDUCTIONS[:2]})
    assert deductions["created"] == 2
    upserted = ok(client, "forecast_upsert_deduction_rules", {"rules": DEDUCTIONS[2:]})
    assert len(upserted["created"]) == 4 and "errors" not in upserted

    sales = ok(client, "forecast_import_sales", {"forecast_id": fid, "rows": _sales_rows()})
    assert sales["imported"] == 16 and not sales["errors"], sales
    assert sales["price_sources"] == {"type_product": 12, "customer_product": 3, "type_all": 1}
    return fid


# Expected figures (see the inputs above).
RFL_COST = 422 + (20 / 60 * 32 + 30 / 60 * 26) + 85       # material + labor + machine per rifle
PST_COST = 120 + (10 / 60 * 32 + 20 / 60 * 26) + 30 / 60 * 85
DEALER_GROSS = 12 * 10 * 1110
DIST_GROSS = 3 * 50 * 520
GOV_GROSS = 5 * 1499 * 0.8
GROSS = DEALER_GROSS + DIST_GROSS + GOV_GROSS
EXCISE = DEALER_GROSS * 0.11 / 1.11 + DIST_GROSS * 0.10 / 1.10      # government is exempt
TERMS = DEALER_GROSS * 0.02
CHARGEBACKS = GROSS * 0.01
NET = GROSS - EXCISE - TERMS - CHARGEBACKS
COMMISSION = 0.03 * (DIST_GROSS - DIST_GROSS * 0.01 - DIST_GROSS * 0.10 / 1.10)
PRODUCT_COGS = 125 * RFL_COST + 150 * PST_COST
LABOR_COST = 125 * (20 / 60 * 32 + 30 / 60 * 26) + 150 * (10 / 60 * 32 + 20 / 60 * 26)
LABOR_HOURS = 125 * 50 / 60 + 150 * 30 / 60


def test_build_model_through_mcp_and_reports_agree(env):
    client, mgr, data_dir, _tmp = env
    fid = _build_model(client, mgr, data_dir)

    # --- COGS forecast (window defaults to the scenario's 2027 sales window) ---------------
    result, text = raw_call(client, "forecast_cogs_forecast", {"forecast_ids": [fid]})
    assert not result.get("isError"), text
    assert "structuredContent" not in result and "\n" not in text   # one compact text block
    cogs = json.loads(text)
    assert (cogs["start_period"], cogs["end_period"]) == ("2027-01", "2027-12")
    totals = cogs["totals"]
    assert approx(totals["gross_sales"], GROSS)
    assert approx(totals["excise_tax"], EXCISE)
    assert approx(totals["terms_allowance"], TERMS)
    assert approx(totals["chargebacks"], CHARGEBACKS)
    assert approx(totals["net_sales"], NET)
    assert approx(totals["commissions"], COMMISSION)
    assert approx(totals["product_cogs"], PRODUCT_COGS)
    assert approx(totals["labor"], LABOR_COST)
    assert approx(totals["gross_sales"] - totals["total_sales_deductions"], totals["net_sales"])
    bridge = cogs["by_period"]
    assert bridge["columns"][0] == "period" and len(bridge["rows"]) == 12
    jan = dict(zip(bridge["columns"], bridge["rows"][0]))
    assert approx(jan["excise_tax"], 11100 * 0.11 / 1.11 + 26000 * 0.10 / 1.10)
    assert approx(sum(dict(zip(bridge["columns"], r))["gross_profit"] for r in bridge["rows"]), totals["gross_profit"], 0.06)

    # --- statements tie to the COGS forecast -------------------------------------------------
    stmts = ok(client, "forecast_financial_statements", {
        "forecast_ids": [fid], "start_period": "2027-01", "end_period": "2027-12", "sections": ["income_statement"],
    })
    inc = stmts["income_statement"]["totals"]
    assert approx(inc["gross_sales"], totals["gross_sales"], 0.01)
    assert approx(inc["revenue"], totals["net_sales"], 0.01)
    assert approx(inc["excise_tax"], totals["excise_tax"], 0.01)
    assert approx(inc["total_cogs"], totals["total_cogs"], 0.01)
    assert approx(inc["gross_profit"], totals["gross_profit"], 0.01)
    assert approx(inc["commissions"], totals["commissions"], 0.01)

    # --- labor forecast ties to the statement's direct labor ---------------------------------
    labor = ok(client, "forecast_labor_forecast", {"forecast_ids": [fid], "include_products": True})
    assert labor["periods"] == YEAR
    assert approx(labor["totals"]["cost"], inc["cogs_labor"], 0.01)
    assert approx(labor["totals"]["hours"], LABOR_HOURS)
    types = {t["labor_type_id"]: t for t in labor["by_labor_type"]}
    assert approx(types["MACHINIST"]["totals"]["hours"], 125 * 20 / 60 + 150 * 10 / 60)
    assert approx(types["ASSEMBLER"]["totals"]["hours"], 125 * 30 / 60 + 150 * 20 / 60)
    assert len(types["MACHINIST"]["hours"]) == 12
    assert approx(types["ASSEMBLER"]["fte"][0], (10 * 0.5 + 50 * 20 / 60) / 173.33, 0.001)
    cnc = next(m for m in labor["by_machine"] if m["machine_id"] == "CNC-01")
    assert approx(cnc["hours"][0], 10 * 1 + 50 * 0.5)                  # January: 10 rifles + 50 pistols
    assert approx(cnc["utilization_pct"][0], 35 / 320 * 100)
    machine_cost = sum(m["totals"]["cost"] for m in labor["by_machine"])
    assert approx(machine_cost, inc["cogs_manufacturing"], 0.01)
    assert labor["by_product"]["columns"][0] == "unit_id"

    # --- groups, unit costs, sales detail ----------------------------------------------------
    by_type = ok(client, "forecast_cogs_forecast", {"forecast_ids": [fid], "group_by": "customer_type",
                                                    "include_periods": True})["by_customer_type"]
    gov = dict(zip(by_type["table"]["columns"], next(r for r in by_type["table"]["rows"] if r[0] == "Government")))
    assert approx(gov["excise_tax"], 0) and approx(gov["gross_sales"], GOV_GROSS)
    assert len(by_type["by_period"][0]["net_sales"]) == 12
    top1 = ok(client, "forecast_cogs_forecast", {"forecast_ids": [fid], "group_by": "customer", "top_n": 1})["by_customer"]
    assert top1["total_groups"] == 3 and len(top1["table"]["rows"]) == 1 and top1["others"]["groups"] == 2
    unit_costs = ok(client, "forecast_cogs_forecast", {"forecast_ids": [fid], "group_by": "unit_costs"})["unit_costs"]
    rfl = dict(zip(unit_costs["columns"], next(r for r in unit_costs["rows"] if r[0] == "RFL-308")))
    assert approx(rfl["total_per_unit"], RFL_COST)

    detail = ok(client, "forecast_sales_detail", {"forecast_ids": [fid], "limit": 5})
    assert detail["total_count"] == 16 and detail["returned"] == 5 and detail["next_offset"] == 5
    assert approx(detail["totals"]["net_sales"], NET)
    gov_rows = ok(client, "forecast_sales_detail", {
        "forecast_ids": [fid], "customer_type": "government",
        "columns": ["period", "customer_id", "gross_sales", "excise_tax", "rules_applied"],
    })
    assert gov_rows["total_count"] == 1 and gov_rows["rows"]["columns"][-1] == "rules_applied"
    assert "excise_tax=" in gov_rows["rows"]["rows"][0][-1]
    assert "Unknown column" in err(client, "forecast_sales_detail", {"forecast_ids": [fid], "columns": ["nope"]})


def test_deduction_maths_precedence_and_pricing(env):
    client, mgr, data_dir, _tmp = env
    fid = _build_model(client, mgr, data_dir)

    # 11% FET embedded in a $1,110 dealer invoice is $110; terms 2%, chargebacks 1%.
    sim = ok(client, "forecast_simulate_deductions", {"gross": 1110, "customer_id": "C-DLR", "unit_id": "RFL-308"})
    amounts = {d["deduction_type"]: d["amount"] for d in sim["deductions"]}
    assert amounts["excise_tax"] == 110 and amounts["terms_allowance"] == 22.2 and amounts["chargebacks"] == 11.1
    assert approx(sim["net_sales"], 1110 - 110 - 22.2 - 11.1, 0.001)

    # Most specific wins: Government (customer type, 4) beats Pistol (product type, 1); 0% = exemption.
    gov = ok(client, "forecast_simulate_deductions", {"gross": 1000, "customer_type": "Government", "unit_type": "Pistol"})
    excise = next(d for d in gov["deductions"] if d["deduction_type"] == "excise_tax")
    assert excise["amount"] == 0 and excise["rule"]["name"] == "Government exempt"
    assert len(excise["overridden_rules"]) == 2

    rules = ok(client, "forecast_list_deduction_rules", {"forecast_id": fid})
    assert rules["count"] == 6 and rules["by_type"]["excise_tax"] == 3
    preview = ok(client, "forecast_deduction_preview", {"forecast_id": fid})
    assert approx(preview["totals"]["excise_tax"], EXCISE) and approx(preview["totals"]["net_sales"], NET)
    assert approx(preview["totals"]["deductions_pct_of_gross"], (GROSS - NET) / GROSS * 100, 0.01)

    # A rule copied from the listing (read-only keys included) can be sent back as an update.
    listed = next(r for r in rules["rules"] if r["deduction_type"] == "chargebacks")
    assert ok(client, "forecast_upsert_deduction_rules", {"rules": [listed]})["updated"][0]["rate"] == 0.01

    # Partial update by rule_id; percent given as a whole number is rejected with a hint.
    terms = next(r for r in rules["rules"] if r["deduction_type"] == "terms_allowance")
    upd = ok(client, "forecast_upsert_deduction_rules", {"rules": [{"rule_id": terms["rule_id"], "rate": 0.025}]})
    assert upd["updated"][0]["rate"] == 0.025 and upd["updated"][0]["customer_type"] == "Dealer"
    bad = err(client, "forecast_upsert_deduction_rules", {"rules": [{"deduction_type": "chargebacks", "rate": 2}]})
    assert "0.02" in bad
    assert "unknown field" in err(client, "forecast_upsert_deduction_rules", {"rules": [{"deduction_typ": "x"}]})
    assert "not found" in err(client, "forecast_upsert_deduction_rules", {"rules": [{"rule_id": "DED-NOPE", "rate": 0.01}]})
    ok(client, "forecast_delete_deduction_rule", {"rule_id": terms["rule_id"]})
    assert ok(client, "forecast_list_deduction_rules", {})["count"] == 5

    # Prices: resolve, list, change one, dry-run apply writes nothing, then apply.
    res = ok(client, "forecast_resolve_price", {"unit_id": "RFL-308", "customer_id": "C-GOV", "period": "2027-06"})
    assert res["price"] == 1199.2 and res["source"] == "type_all"
    prices = ok(client, "forecast_list_prices", {"unit_id": "RFL-308"})
    dealer = next(p for p in prices["prices"] if p.get("customer_type") == "Dealer")
    ok(client, "forecast_upsert_prices", {"prices": [{"price_id": dealer["price_id"], "unit_price": 1150}]})
    before = count(mgr, "SELECT SUM(total_revenue) FROM sales")
    dry = ok(client, "forecast_apply_price_list", {"forecast_id": fid, "customer_types": ["Dealer"], "dry_run": True})
    assert dry["changed"] == 12 and approx(dry["revenue_change"], 12 * 10 * 40)
    assert count(mgr, "SELECT SUM(total_revenue) FROM sales") == before, "a dry run must not write"
    # A misspelled filter is an error, not a re-price of the whole scenario.
    assert "unknown argument" in err(client, "forecast_apply_price_list",
                                     {"forecast_id": fid, "customer_type": ["Dealer"], "dry_run": True}).lower()
    applied = ok(client, "forecast_apply_price_list", {"forecast_id": fid, "customer_types": ["Dealer"]})
    assert applied["changed"] == 12
    assert approx(count(mgr, "SELECT SUM(total_revenue) FROM sales"), before + 4800)
    ok(client, "forecast_delete_price", {"price_id": dealer["price_id"]})
    assert ok(client, "forecast_resolve_price", {"unit_id": "RFL-308", "customer_id": "C-DLR"})["source"] == "base_price"


def test_import_guards_templates_and_dry_runs(env, monkeypatch):
    client, mgr, data_dir, tmp = env
    templates = ok(client, "forecast_list_import_templates")["templates"]
    assert [t["kind"] for t in templates][:2] == ["customers", "products"]
    assert next(t for t in templates if t["kind"] == "boms")["tool"] == "forecast_import_boms"
    one = ok(client, "forecast_list_import_templates", {"kind": "deductions"})["templates"][0]
    assert "excise_tax" in one["example_csv"] and any(c.get("aliases") for c in one["columns"])

    # Dry run of master data writes nothing; the real run does.
    dry = ok(client, "forecast_import_records", {"table": "customers", "rows": CUSTOMERS, "dry_run": True})
    assert dry["created"] == 3 and count(mgr, "SELECT COUNT(*) FROM customers") == 0

    # Rows and file paths.
    with open(os.path.join(tmp, "outside.csv"), "w") as fh:
        fh.write("customer_name,customer_type\nX,Dealer\n")
    assert "outside the data folder" in err(client, "forecast_import_records",
                                            {"table": "customers", "file_path": "../outside.csv"})
    assert "or file_path" in err(client, "forecast_import_records", {"table": "customers"})
    assert "not found" in err(client, "forecast_import_boms", {"file_path": "imports/missing.xlsx"})
    # Absolute / '~' paths outside the data folder only with FORECAST_ALLOW_ABSOLUTE_PATHS=1 (review MCP-6).
    monkeypatch.delenv("FORECAST_ALLOW_ABSOLUTE_PATHS", raising=False)
    outside = os.path.join(tmp, "outside.csv")
    assert "FORECAST_ALLOW_ABSOLUTE_PATHS=1" in err(client, "forecast_import_records",
                                                    {"table": "customers", "file_path": outside, "dry_run": True})
    assert "outside the data folder" in err(client, "forecast_preview_file", {"file_path": "/"})
    assert "home-folder" in err(client, "forecast_preview_file", {"file_path": "~"})
    inside = os.path.join(data_dir, "inside.csv")
    with open(inside, "w") as fh:
        fh.write("customer_name,customer_type\nY,Dealer\n")
    assert ok(client, "forecast_preview_file", {"file_path": inside})["row_count"] == 1   # absolute, but inside
    monkeypatch.setenv("FORECAST_ALLOW_ABSOLUTE_PATHS", "1")
    abs_ok = ok(client, "forecast_import_records", {"table": "customers", "file_path": outside, "dry_run": True})
    assert abs_ok["created"] == 1 and abs_ok["source"]["row_count"] == 1
    monkeypatch.delenv("FORECAST_ALLOW_ABSOLUTE_PATHS")
    listing = ok(client, "forecast_preview_file", {})
    assert listing["folder"] == os.path.realpath(data_dir)

    # Row errors come back with 1-based row numbers and a failed status.
    bad = ok(client, "forecast_import_routings", {"rows": [
        {"router_id": "R1", "machine_id": "NOPE", "labor_minutes": 5, "labor_type_id": "X"}]})
    assert bad["status"] == "failed" and bad["errors"][0]["row"] == 1
    assert "Unknown machine" in bad["errors"][0]["error"]

    # Unknown arguments are rejected everywhere in the new tools.
    assert "unknown argument" in err(client, "forecast_import_boms", {"rows": BOM_ROWS, "dryrun": True}).lower()
    assert count(mgr, "SELECT COUNT(*) FROM bom") == 0


def test_export_workbook_writes_to_data_exports(env):
    client, mgr, data_dir, _tmp = env
    fid = _build_model(client, mgr, data_dir)
    saved = ok(client, "forecast_export_workbook", {"forecast_ids": [fid]})
    expected = os.path.join(os.path.realpath(data_dir), "exports", f"forecast_{fid}_2027-01_2027-12.xlsx")
    assert os.path.realpath(saved["path"]) == expected and os.path.getsize(expected) == saved["bytes"]
    assert "host_path_hint" not in saved            # only when the data folder is the Docker volume /data
    for sheet in ("Read Me", "Gross Margin", "Income Statement", "By Customer Type", "Labor Forecast",
                  "Sales Detail", "Deduction Rules", "Price List"):
        assert sheet in saved["sheets"]
    from openpyxl import load_workbook
    wb = load_workbook(expected, read_only=True)
    assert "Labor Forecast" in wb.sheetnames
    wb.close()

    named = ok(client, "forecast_export_workbook", {"forecast_ids": [fid], "start_period": "2027-01",
                                                    "end_period": "2027-03", "filename": "../Q1 budget"})
    assert os.path.dirname(named["path"]) == os.path.join(os.path.realpath(data_dir), "exports")
    assert named["filename"] == "Q1_budget.xlsx"
    assert "start_period must be on or before end_period" in err(
        client, "forecast_export_workbook", {"forecast_ids": [fid], "start_period": "2027-05", "end_period": "2027-01"})
    assert "Scenario 'F999' not found" in err(client, "forecast_export_workbook", {"forecast_ids": ["F999"]})
    assert "Scenario 'F999' not found" in err(client, "forecast_cogs_forecast", {
        "forecast_ids": ["F999"], "start_period": "2027-01", "end_period": "2027-02"})
    empty = ok(client, "forecast_cogs_forecast", {"forecast_ids": [fid], "start_period": "2028-01", "end_period": "2028-03"})
    assert empty["sales_lines"] == 0 and "2027-01..2027-12" in empty["hint"]


def test_docker_data_folder_gets_a_host_path_hint(env, monkeypatch):
    client, mgr, data_dir, _tmp = env
    fid = _build_model(client, mgr, data_dir)
    import mcp_server.tools.forecasts as forecasts
    monkeypatch.setattr(forecasts, "DOCKER_DATA_DIR", os.path.realpath(data_dir))
    saved = ok(client, "forecast_export_workbook", {"forecast_ids": [fid], "filename": "budget"})
    assert saved["host_path_hint"] == "./data/exports/budget.xlsx"
    assert os.path.exists(os.path.join(data_dir, "exports", "budget.xlsx"))


def test_new_tools_are_registered_with_annotations_and_strict_schemas(env):
    client, _mgr, _data_dir, _tmp = env
    res = client.post("/mcp", headers=HEADERS, json={"jsonrpc": "2.0", "id": 1, "method": "tools/list", "params": {}})
    tools = {t["name"]: t for t in res.json()["result"]["tools"]}
    # destructiveHint: anything that overwrites or deletes stored data (review MCP-10); idempotentHint: the
    # same call twice leaves the same state (the upserts match existing rules by scope, review MCP-1).
    expected = {
        "forecast_import_boms": "destructive", "forecast_import_routings": "destructive",
        "forecast_import_records": "upsert", "forecast_import_prices": "destructive",
        "forecast_import_deductions": "destructive", "forecast_list_import_templates": "read",
        "forecast_preview_file": "read", "forecast_list_deduction_rules": "read",
        "forecast_upsert_deduction_rules": "upsert", "forecast_delete_deduction_rule": "destructive",
        "forecast_simulate_deductions": "read", "forecast_deduction_preview": "read", "forecast_list_prices": "read",
        "forecast_upsert_prices": "upsert", "forecast_delete_price": "destructive", "forecast_resolve_price": "read",
        "forecast_apply_price_list": "upsert", "forecast_cogs_forecast": "read", "forecast_labor_forecast": "read",
        "forecast_sales_detail": "read", "forecast_export_workbook": "upsert",
        "forecast_import_sales": "upsert", "forecast_adjust_sales": "destructive",
        "forecast_adjust_material_costs": "destructive", "forecast_upsert_sales": "destructive",
        "forecast_create_scenario": "write", "forecast_save_database": "write",
    }
    for name, kind in expected.items():
        tool = tools[name]
        ann = tool["annotations"]
        assert tool["inputSchema"].get("additionalProperties") is False, name
        assert "outputSchema" not in tool, name
        assert ann["readOnlyHint"] is (kind == "read"), name
        assert ann["destructiveHint"] is (kind in ("destructive", "upsert")), name
        assert ann["idempotentHint"] is (kind in ("read", "upsert")), name
    assert tools["forecast_cogs_forecast"]["inputSchema"]["properties"]["group_by"]["enum"] == [
        "summary", "product", "customer", "customer_type", "product_type", "unit_costs"]


# ---------------------------------------------------------------------------
# Helpers without the transport
# ---------------------------------------------------------------------------

def test_read_table_file_handles_csv_dialects_xlsx_sheets_and_header_rows(env, monkeypatch):
    _client, _mgr, data_dir, _tmp = env
    from mcp.server.mcpserver.exceptions import ToolError
    from mcp_server.tools.imports import read_table_file

    with open(os.path.join(data_dir, "semi.csv"), "wb") as fh:   # Excel-on-Windows export: ';' and cp1252
        fh.write("Kunde;Name;Typ\nC1;Caf\xe9 Guns;Dealer\n;;\nC2;Range;Retail\n".encode("cp1252"))
    rows, info = read_table_file("semi.csv")
    assert info["columns"] == ["Kunde", "Name", "Typ"] and info["row_count"] == 2
    assert rows[0]["Name"] == "Caf\xe9 Guns"

    from datetime import datetime
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Notes"
    wb.active.append(["nothing here"])
    ws = wb.create_sheet("Forecast")
    ws.append(["Customer", "SKU", datetime(2027, 1, 1), datetime(2027, 2, 1)])
    ws.append(["C1", "RFL-308", 10, 12.0])
    wb.save(os.path.join(data_dir, "wide.xlsx"))
    rows, info = read_table_file("data/wide.xlsx", sheet="forecast")      # host-style prefix, any-case sheet
    assert info["sheet"] == "Forecast" and info["sheets"] == ["Notes", "Forecast"]
    assert info["columns"] == ["Customer", "SKU", "2027-01", "2027-02"] and rows[0]["2027-02"] == 12.0
    with pytest.raises(ToolError, match="Sheet 'Plan' not found"):
        read_table_file("wide.xlsx", sheet="Plan")
    with pytest.raises(ToolError, match="only .csv"):
        open(os.path.join(data_dir, "old.xls"), "wb").close()
        read_table_file("old.xls")
    rows, info = read_table_file("semi.csv", header_row=2)
    assert info["header_row"] == 2 and info["columns"][0] == "C1"


def test_strict_arguments_fallback_rejects_unknown_keys_without_echoing_input():
    from mcp.server.mcpserver import MCPServer
    from pydantic import ValidationError
    from mcp_server.tools._common import _forbid_unknown_arguments

    server = MCPServer("t")

    @server.tool(name="demo", structured_output=False)
    async def demo(customer_ids: list = None, dry_run: bool = False) -> str:  # pragma: no cover - not called
        return "x"

    tool = server._tool_manager.get_tool("demo")
    _forbid_unknown_arguments(tool)
    assert tool.parameters["additionalProperties"] is False
    assert tool.fn_metadata.validate_arguments({"customer_ids": ["C1"]})["customer_ids"] == ["C1"]
    with pytest.raises(ValidationError) as info:
        tool.fn_metadata.validate_arguments({"customer_type": "Dealer", "rows": ["secret"] * 50})
    text = str(info.value)
    assert "unknown argument(s) customer_type, rows" in text and "customer_ids, dry_run" in text
    assert "secret" not in text
